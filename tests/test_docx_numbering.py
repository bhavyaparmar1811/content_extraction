"""Word auto-numbering, content controls and list levels in the DOCX parser.

Real SOPs number their headings through ``w:numPr`` on the Heading styles, so
"6.1" never appears in the heading text. These tests build that structure by
hand and check the parser, AST builder and exporter reproduce Word's numbers.
"""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.config.settings import Settings
from app.schemas.document import ElementType
from app.services.export.migration_exporter import MigrationExporter
from app.services.hierarchy.ast_builder import ASTBuilder
from app.services.parser.docx_parser import DocxParser
from app.services.parser.ooxml import NumberingResolver, element_text, iter_body_blocks

HEADING_NUM = "901"
LIST_NUM = "902"
LETTER_NUM = "903"
RESTART_NUM = "904"


@pytest.fixture
def settings(tmp_path):
    s = Settings(project_root=tmp_path)
    s.resolve_paths(tmp_path)
    s.ensure_directories()
    return s


# ── Builders ───────────────────────────────────────────────────────────

def _lvl(ilvl: int, fmt: str, text: str, start: int = 1) -> str:
    return (f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/></w:lvl>')


def _add_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstracts = {
        "801": [_lvl(0, "decimal", "%1"), _lvl(1, "decimal", "%1.%2"), _lvl(2, "decimal", "%1.%2.%3")],
        "802": [_lvl(0, "bullet", "•"), _lvl(1, "bullet", "o")],
        "803": [_lvl(0, "lowerLetter", "%1)"), _lvl(1, "lowerRoman", "%2.")],
    }
    first_num = numbering.find(qn("w:num"))
    for aid, levels in abstracts.items():
        an = parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}">{"".join(levels)}</w:abstractNum>')
        first_num.addprevious(an)  # abstractNum elements must precede num elements
    for num_id, aid, extra in [
        (HEADING_NUM, "801", ""),
        (LIST_NUM, "802", ""),
        (LETTER_NUM, "803", ""),
        (RESTART_NUM, "803", '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="4"/></w:lvlOverride>'),
    ]:
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{aid}"/>{extra}</w:num>'
        ))


def _number_style(doc, style_name: str, ilvl: int):
    ppr = doc.styles[style_name].element.get_or_add_pPr()
    ppr.insert(0, parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{HEADING_NUM}"/></w:numPr>'
    ))


def _number_paragraph(paragraph, num_id: str, ilvl: int = 0):
    paragraph._p.get_or_add_pPr().insert(0, parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))


def _wrap_in_sdt(paragraph, toc: bool = False):
    gallery = ('<w:docPartObj><w:docPartGallery w:val="Table of Contents"/></w:docPartObj>' if toc else "")
    sdt = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr>{gallery}</w:sdtPr><w:sdtContent/></w:sdt>')
    paragraph._p.addprevious(sdt)
    sdt.find(qn("w:sdtContent")).append(paragraph._p)


def _numbered_doc(path: Path) -> Path:
    doc = Document()
    _add_numbering(doc)
    for level in (1, 2, 3):
        _number_style(doc, f"Heading {level}", level - 1)

    doc.add_paragraph("Cover sheet text")
    _wrap_in_sdt(doc.add_paragraph("1PURPOSE3"), toc=True)
    doc.add_heading("PURPOSE", level=1)                  # 1
    doc.add_paragraph("Why this exists.")
    doc.add_heading("PROCESS", level=1)                  # 2
    doc.add_heading("General", level=2)                  # 2.1
    doc.add_heading("Detail", level=3)                   # 2.1.1
    doc.add_heading("", level=2)                         # 2.2: empty, still counted
    doc.add_heading("Special", level=2)                  # 2.3
    unnumbered = doc.add_heading("Unnumbered aside", level=2)
    _number_paragraph(unnumbered, "0")                   # numId 0 switches numbering off
    doc.add_heading("REFERENCES", level=1)               # 3

    for text, ilvl in [("first bullet", 0), ("nested bullet", 1), ("second bullet", 0)]:
        _number_paragraph(doc.add_paragraph(text), LIST_NUM, ilvl)
    doc.add_paragraph("Between lists.")
    for text, ilvl in [("alpha", 0), ("alpha sub", 1), ("beta", 0)]:
        _number_paragraph(doc.add_paragraph(text), LETTER_NUM, ilvl)
    _number_paragraph(doc.add_paragraph("restarted"), RESTART_NUM, 0)

    _wrap_in_sdt(doc.add_paragraph("Inside a content control"))
    doc.save(path)
    return path


def _parse(settings, path):
    return DocxParser(settings=settings).parse(str(path), document_id="numbering_test")


def _elements(raw):
    return [el for page in raw.pages for el in page.elements]


# ── Resolver ───────────────────────────────────────────────────────────

def test_resolver_reproduces_word_numbers(tmp_path):
    doc = Document(str(_numbered_doc(tmp_path / "n.docx")))
    resolver = NumberingResolver(doc)
    numbers = {}
    for block in iter_body_blocks(doc.element.body):
        info = resolver.advance(block)
        text = element_text(block)
        if text and info is not None:
            numbers[text] = (info.number, info.is_bullet, info.ilvl)

    assert numbers["PURPOSE"][0] == "1"
    assert numbers["PROCESS"][0] == "2"
    assert numbers["General"][0] == "2.1"
    assert numbers["Detail"][0] == "2.1.1"
    assert numbers["Special"][0] == "2.3"  # the empty Heading 2 took 2.2
    assert "Unnumbered aside" not in numbers
    assert numbers["REFERENCES"][0] == "3"
    assert numbers["first bullet"] == (None, True, 0)
    assert numbers["nested bullet"] == (None, True, 1)
    assert numbers["alpha"][0] == "a)"
    assert numbers["alpha sub"][0] == "i"
    assert numbers["beta"][0] == "b)"
    assert numbers["restarted"][0] == "d)"  # startOverride 4 on its own numId


def test_toc_content_control_is_skipped_but_others_are_read(tmp_path):
    doc = Document(str(_numbered_doc(tmp_path / "n.docx")))
    texts = [element_text(b) for b in iter_body_blocks(doc.element.body)]
    assert "1PURPOSE3" not in texts
    assert "Inside a content control" in texts


def test_style_chain_numbering_is_inherited(tmp_path):
    doc = Document()
    _add_numbering(doc)
    _number_style(doc, "Heading 1", 0)
    styles = doc.styles.element
    styles.append(parse_xml(
        f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="ChapterHeading">'
        '<w:name w:val="Chapter Heading"/><w:basedOn w:val="Heading1"/></w:style>'
    ))
    doc.add_paragraph("Chapter", style="Chapter Heading")
    doc.add_paragraph("Second chapter", style="Chapter Heading")
    resolver = NumberingResolver(doc)
    infos = [resolver.advance(p._p) for p in doc.paragraphs]
    assert [i.number for i in infos] == ["1", "2"]
    assert resolver.outline_level(doc.paragraphs[0]._p) == 0  # inherited from Heading 1


def test_element_text_reads_inline_content_controls():
    p = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:t>Version </w:t></w:r>'
        '<w:sdt><w:sdtContent><w:r><w:t>5</w:t></w:r></w:sdtContent></w:sdt>'
        '<w:r><w:t>.</w:t></w:r>'
        '<w:sdt><w:sdtContent><w:r><w:t>0</w:t></w:r></w:sdtContent></w:sdt></w:p>'
    )
    assert element_text(p) == "Version 5.0"


# ── Parser, AST and exporter ───────────────────────────────────────────

def test_parser_tags_headings_and_lists(settings, tmp_path):
    elements = _elements(_parse(settings, _numbered_doc(tmp_path / "n.docx")))
    headings = {el.content: el for el in elements if el.element_type == ElementType.HEADING}
    assert headings["Detail"].metadata["numbering"] == "2.1.1"
    assert headings["Detail"].metadata["numbering_source"] == "numPr"
    assert "numbering" not in headings["Unnumbered aside"].metadata

    by_text = {el.content: el for el in elements}
    assert by_text["nested bullet"].element_type == ElementType.LIST_ITEM
    assert by_text["nested bullet"].metadata["list_level"] == 1
    assert by_text["beta"].element_type == ElementType.NUMBERED_STEP
    assert by_text["beta"].metadata["list_number"] == "b)"
    assert by_text["Inside a content control"].element_type == ElementType.PARAGRAPH
    assert "1PURPOSE3" not in by_text
    indices = [el.metadata["paragraph_index"] for el in elements if "paragraph_index" in el.metadata]
    assert indices == sorted(indices) and len(set(indices)) == len(indices)


def test_ast_lists_carry_levels(settings, tmp_path):
    ast = ASTBuilder(settings=settings).build(_parse(settings, _numbered_doc(tmp_path / "n.docx")))

    def walk(node):
        yield node
        for child in getattr(node, "children", []) or []:
            yield from walk(child)

    lists = [n for n in walk(ast) if getattr(n, "node_type", None) == "list"]
    bullets = next(l for l in lists if l.items[0].text == "first bullet")
    assert bullets.list_type == "unordered"
    assert bullets.nesting_depth == 1
    assert [i.metadata["list_level"] for i in bullets.items] == [0, 1, 0]
    letters = next(l for l in lists if l.items[0].text == "alpha")
    assert letters.list_type == "ordered"
    assert letters.items[0].source_location.paragraph_index is not None


def test_exporter_builds_sections_from_word_numbers(settings, tmp_path):
    raw = _parse(settings, _numbered_doc(tmp_path / "n.docx"))
    out = MigrationExporter.export(document_id="n", ast=ASTBuilder(settings=settings).build(raw))

    assert [(s.section_number, s.title) for s in out.sections] == [
        ("0", "0 PREAMBLE"),
        ("1", "1 PURPOSE"),
        ("2", "2 PROCESS"),
        ("3", "3 REFERENCES"),
    ]
    process = out.sections[2]
    sub_headings = [e.text for e in process.elements if e.element_type == "heading"]
    assert sub_headings == ["2.1 General", "2.1.1 Detail", "2.3 Special", "Unnumbered aside"]


def test_unnumbered_heading_1_gets_running_number(settings, tmp_path):
    doc = Document()
    doc.add_heading("PURPOSE", level=1)
    doc.add_paragraph("Body.")
    doc.add_heading("SCOPE", level=1)
    doc.save(tmp_path / "plain.docx")
    raw = _parse(settings, tmp_path / "plain.docx")
    out = MigrationExporter.export(document_id="p", ast=ASTBuilder(settings=settings).build(raw))
    assert [s.title for s in out.sections] == ["1 PURPOSE", "2 SCOPE"]
