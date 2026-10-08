"""Tests for the v2 template callout palette (Phase 3).

Synthetic templates cover the detection rules; the real sample template, when
present, pins the four palette colours.
"""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, RGBColor
from PIL import Image

from app.schemas.v2 import CalloutKind, CalloutLayout
from app.services.migration_v2.template.callout_palette import extract_callout_palette, kind_for_label

ROOT = Path(__file__).resolve().parent.parent
BLUE = RGBColor(0x00, 0x75, 0xFF)
SAMPLE_TEMPLATE = "Template Main GP Docs.docx"
SAMPLE_DIRS = [
    ROOT / "tests" / "fixtures" / "migration_v2" / "template",
    ROOT / "data" / "samples" / "template",
    ROOT / "documents" / "Templates",
]


# ── Helpers ────────────────────────────────────────────────────────────

def _png(path: Path, color: tuple[int, int, int]) -> Path:
    Image.new("RGB", (16, 16), color).save(path)
    return path


def _shade(cell, fill: str):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _box_row(table, row_idx: int, fill: str, icon: Path | None, text: str, blue: bool = False):
    row = table.rows[row_idx]
    if icon is not None:
        row.cells[0].paragraphs[0].add_run().add_picture(str(icon), width=Inches(0.22))
    run = row.cells[-1].paragraphs[0].add_run(text)
    if blue:
        run.font.color.rgb = BLUE
    for cell in row.cells:
        _shade(cell, fill)


def _legend(doc, rows: list[tuple[str, Path, str]]):
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    table.rows[0].cells[0].text = "Infographics"
    table.rows[0].cells[1].text = "Description"
    for i, (fill, icon, text) in enumerate(rows, start=1):
        _box_row(table, i, fill, icon, text)
    return table


def _prototype(doc, fill: str, icon: Path, text: str):
    _box_row(doc.add_table(rows=1, cols=2), 0, fill, icon, text, blue=True)


@pytest.fixture
def icons(tmp_path):
    return {
        "intro": _png(tmp_path / "intro.png", (30, 90, 200)),
        "expl": _png(tmp_path / "expl.png", (40, 160, 60)),
        "attn": _png(tmp_path / "attn.png", (200, 40, 40)),
        "key": _png(tmp_path / "key.png", (220, 200, 40)),
    }


def _standard_template(path: Path, icons) -> Path:
    doc = Document()
    doc.add_heading("PURPOSE", level=1)
    _legend(doc, [
        ("D2F2F7", icons["intro"], "Introduction/Executive Summary – short description."),
        ("00E47C", icons["expl"], "Explanation – additional information to the topic."),
        ("F5CDB9", icons["attn"], "Attention – notice taken of something important."),
        ("FBF9AA", icons["key"], "Key take way - what you should remember."),
    ])
    doc.add_heading("PROCESS", level=1)
    _prototype(doc, "F5CDB9", icons["attn"], "Attention section (what you should know). It is pink.")
    _prototype(doc, "D2F2F7", icons["intro"], "Executive Summary/Introduction (short description).")
    doc.save(path)
    return path


# ── Label mapping ──────────────────────────────────────────────────────

@pytest.mark.parametrize("label, kind", [
    ("Introduction/Executive Summary – short description", CalloutKind.INTRODUCTION),
    ("Executive Summary/Introduction (short description of subject)", CalloutKind.INTRODUCTION),
    ("Explanation – additional information", CalloutKind.EXPLANATION),
    ("Attention – notice taken of something", CalloutKind.ATTENTION),
    ("Warning: do not mix up", CalloutKind.ATTENTION),
    ("Key take way - what you should remember", CalloutKind.KEY_TAKEAWAY),
    ("Key-take-away section", CalloutKind.KEY_TAKEAWAY),
    ("Roles and responsibilities", None),
])
def test_kind_for_label(label, kind):
    assert kind_for_label(label) == kind


# ── Synthetic templates ────────────────────────────────────────────────

def test_legend_and_prototypes_build_palette(tmp_path, icons):
    result = extract_callout_palette(_standard_template(tmp_path / "t.docx", icons))

    assert result.legend_table_index == 0
    assert [s.kind for s in result.palette] == list(CalloutKind)
    by_kind = {s.kind: s for s in result.palette}
    assert by_kind[CalloutKind.ATTENTION].fill_hex == "F5CDB9"
    assert by_kind[CalloutKind.ATTENTION].prototype_table_index == 1
    assert by_kind[CalloutKind.INTRODUCTION].prototype_table_index == 2
    assert by_kind[CalloutKind.KEY_TAKEAWAY].label == "Key take way"
    assert all(s.layout == CalloutLayout.ICON_TEXT_TWO_CELL for s in result.palette)
    assert len({s.icon_key for s in result.palette}) == 4

    assert [(p.kind, p.section_heading) for p in result.prototypes] == [
        (CalloutKind.ATTENTION, "PROCESS"),
        (CalloutKind.INTRODUCTION, "PROCESS"),
    ]
    assert "explanation has no prototype" in " ".join(result.issues)


def test_prototype_kind_comes_from_fill_not_text(tmp_path, icons):
    """A pink box whose instruction mentions 'introduction' is still attention."""
    doc = Document()
    _legend(doc, [("F5CDB9", icons["attn"], "Attention – important.")])
    _prototype(doc, "F5CDB9", icons["attn"], "Introduction of the hazards you must not mix up.")
    doc.save(tmp_path / "t.docx")

    result = extract_callout_palette(tmp_path / "t.docx")
    assert [p.kind for p in result.prototypes] == [CalloutKind.ATTENTION]


def test_mixed_fill_rows_are_not_callouts(tmp_path, icons):
    """RACI-style rows: shaded cells next to unshaded ones."""
    doc = Document()
    raci = doc.add_table(rows=2, cols=3)
    raci.rows[0].cells[0].text = "Process Step"
    raci.rows[1].cells[0].text = "6.1"
    _shade(raci.rows[1].cells[0], "D2F2F7")
    _shade(raci.rows[1].cells[1], "D2F2F7")
    doc.save(tmp_path / "t.docx")

    result = extract_callout_palette(tmp_path / "t.docx")
    assert result.palette == [] and result.prototypes == []


def test_single_cell_box(tmp_path):
    doc = Document()
    _box_row(doc.add_table(rows=1, cols=1), 0, "FBF9AA", None, "Key take-away: remember the deadline.")
    doc.save(tmp_path / "t.docx")

    result = extract_callout_palette(tmp_path / "t.docx")
    (style,) = result.palette
    assert style.kind == CalloutKind.KEY_TAKEAWAY
    assert style.layout == CalloutLayout.SINGLE_CELL
    assert style.prototype_table_index == 0


def test_unknown_fill_is_reported(tmp_path, icons):
    doc = Document()
    _prototype(doc, "ABCDEF", icons["attn"], "Something with no recognisable label.")
    doc.save(tmp_path / "t.docx")

    result = extract_callout_palette(tmp_path / "t.docx")
    assert result.palette == []
    assert any("ABCDEF" in issue for issue in result.issues)


def test_fill_override_wins(tmp_path, icons):
    doc = Document()
    _prototype(doc, "ABCDEF", icons["attn"], "Something with no recognisable label.")
    doc.save(tmp_path / "t.docx")

    result = extract_callout_palette(tmp_path / "t.docx", fill_overrides={"#abcdef": "explanation"})
    (style,) = result.palette
    assert (style.kind, style.fill_hex, style.prototype_table_index) == (CalloutKind.EXPLANATION, "ABCDEF", 0)


def test_highlighted_optional_heading(tmp_path):
    doc = Document()
    heading = doc.add_heading("Distribution of controlled prints ", level=1)
    heading.add_run("(optional)").font.highlight_color = 7  # WD_COLOR_INDEX.YELLOW
    doc.add_paragraph().add_run("Check this wording").font.highlight_color = 7
    doc.save(tmp_path / "t.docx")

    result = extract_callout_palette(tmp_path / "t.docx")
    assert result.optional_section_headings == ["Distribution of controlled prints (optional)"]
    assert len(result.highlights) == 2
    assert any("Check this wording" in issue for issue in result.issues)


# ── Real sample template ───────────────────────────────────────────────

def _sample_template() -> Path | None:
    return next((d / SAMPLE_TEMPLATE for d in SAMPLE_DIRS if (d / SAMPLE_TEMPLATE).exists()), None)


@pytest.mark.skipif(_sample_template() is None, reason="sample template not available")
def test_sample_template_palette():
    result = extract_callout_palette(_sample_template())

    assert {s.kind: s.fill_hex for s in result.palette} == {
        CalloutKind.INTRODUCTION: "D2F2F7",
        CalloutKind.EXPLANATION: "00E47C",
        CalloutKind.ATTENTION: "F5CDB9",
        CalloutKind.KEY_TAKEAWAY: "FBF9AA",
    }
    assert all(s.icon_key and s.prototype_table_index is not None for s in result.palette)
    assert len(result.prototypes) == 4  # the RACI matrix's shaded rows are not callouts
    assert result.optional_section_headings == ["distribution of controlled prints/copies (optional)"]
    assert result.issues == []
