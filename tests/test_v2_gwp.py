"""Phase 4: GWP ingestion, rule extraction (mocked LLM), review, approval and selection."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.auth import require_user
from app.config.settings import Settings, get_settings
from app.schemas.v2 import (
    ContentType,
    GwpRule,
    GwpRuleSet,
    RuleCategory,
    RuleCheck,
    RuleOrigin,
    RuleStatus,
    SourceDocument,
    SourceSection,
    SourceUnit,
    UnitType,
)
from app.services.migration_v2.gwp import service
from app.services.migration_v2.gwp.baseline import BASELINE_IDS, baseline_rules
from app.services.migration_v2.gwp.checks import check_params_problems
from app.services.migration_v2.gwp.extractor import (
    CandidateRule,
    CandidateSkip,
    GwpExtractionOutput,
    GwpRuleExtractor,
    build_batches,
    build_rule_set,
)
from app.services.migration_v2.gwp.ingest import parse_guide, rule_input_sections
from app.services.llm.prompts.v2.gwp_extractor import PROMPT_VERSION
from app.stores.gwp_store import GwpStore


# ── Fixtures ──────────────────────────────────────────────────────────


def _unit(section: str, n: int, seq: int, text: str, boilerplate: bool = False) -> SourceUnit:
    return SourceUnit(
        unit_id=f"{section}-U{n:03d}", content_hash=f"h{seq}", section_id=section, seq=seq,
        unit_type=UnitType.PARAGRAPH, text=text, is_boilerplate=boilerplate,
    )


def _doc() -> SourceDocument:
    """A guide with a cover page, TOC, two rule sections, references and history."""
    sections = [
        ("SRC-0", "0", "PREAMBLE", ["Title | Good Writing"]),
        ("SRC-S01", None, "Table of Content", ["1 PURPOSE .... 3"]),
        ("SRC-1", "1", "PURPOSE", ["This Guidance outlines writing."]),
        ("SRC-6.1", "6.1", "LANGUAGE AND WORDING", [
            'Do not use "please".', 'Prefer "confirm" or "verify" over "ensure".', "Write short sentences.",
        ]),
        ("SRC-6.2", "6.2", "INFOGRAPHICS", ["Use the Attention infographic with Light Red colour (HEX #F5CDB9)."]),
        ("SRC-9", "9", "REFERENCES", ["1 | BI-VQD-1 | Manual"]),
        ("SRC-10", "10", "DOCUMENT HISTORY", ["3.0 | Revision"]),
    ]
    seq = 0
    out = []
    for order, (sid, number, heading, texts) in enumerate(sections):
        units = []
        for n, text in enumerate(texts, start=1):
            units.append(_unit(sid, n, seq, text, boilerplate=heading == "DOCUMENT HISTORY"))
            seq += 1
        out.append(SourceSection(section_id=sid, number=number, heading=heading, level=1, section_order=order, units=units))
    return SourceDocument(document_id="GWP", source_file="gwp.pdf", file_type="pdf", sections=out)


def _output() -> GwpExtractionOutput:
    return GwpExtractionOutput(
        rules=[
            CandidateRule(
                category=RuleCategory.STYLE, text='Do not use "please".', severity="high", check="deterministic",
                params_json='{"kind": "forbidden_terms", "terms": ["please"]}', source_unit_ids=["SRC-6.1-U001"],
            ),
            CandidateRule(
                category=RuleCategory.FORMATTING, text="Shade Attention callouts Light Red (#F5CDB9).",
                check="deterministic", params_json='{"kind": "callout_palette", "colors": {"attention": "F5CDB9"}}',
                source_unit_ids=["SRC-6.2-U001"],
            ),
            CandidateRule(  # cited before the "please" rule in guide order, so it becomes STY-001
                category=RuleCategory.STYLE, text="Describe the purpose briefly.", source_unit_ids=["SRC-1-U001"],
            ),
            CandidateRule(  # bad params: downgraded to semantic
                category=RuleCategory.STYLE, text='Prefer "confirm" or "verify" over "ensure".', check="deterministic",
                params_json='{"kind": "forbidden_terms"}', source_unit_ids=["SRC-6.1-U002"],
            ),
            CandidateRule(  # cites nothing that exists: dropped
                category=RuleCategory.STRUCTURAL, text="Invented rule.", source_unit_ids=["SRC-99-U001"],
            ),
            CandidateRule(  # a guide PRES rule must not take a baseline ID
                category=RuleCategory.PRESERVATION, text="Keep technical detail needed to execute the procedure.",
                source_unit_ids=["SRC-6.1-U003"],
            ),
        ],
        skipped=[CandidateSkip(source_unit_ids=["SRC-6.1-U003", "SRC-NOPE"], reason="general advice")],
    )


class FakeChain:
    def __init__(self, outputs):
        self.outputs = list(outputs)
        self.calls = []

    async def ainvoke(self, messages):
        self.calls.append(messages)
        return self.outputs.pop(0) if self.outputs else GwpExtractionOutput()


# ── Checks and baseline ───────────────────────────────────────────────


def _det(params: dict) -> GwpRule:
    return GwpRule(rule_id="STY-001", category=RuleCategory.STYLE, text="x", check=RuleCheck.DETERMINISTIC, params=params)


@pytest.mark.parametrize("params, ok", [
    ({"kind": "forbidden_terms", "terms": ["please"], "prefer": ["confirm"]}, True),
    ({"kind": "forbidden_terms", "terms": []}, False),
    ({"kind": "forbidden_terms"}, False),
    ({"kind": "max_sentence_words", "max_words": 25}, True),
    ({"kind": "max_sentence_words", "max_words": 0}, False),
    ({"kind": "passive_ratio", "max_ratio": 0.1}, True),
    ({"kind": "passive_ratio", "max_ratio": 10}, False),
    ({"kind": "readability", "flesch_reading_ease_min": 30}, True),
    ({"kind": "readability"}, False),
    ({"kind": "callout_palette", "colors": {"attention": "F5CDB9"}}, True),
    ({"kind": "callout_palette", "colors": {"banner": "F5CDB9"}}, False),
    ({"kind": "protected_values", "fact_kinds": ["number"]}, True),
    ({"kind": "protected_values", "fact_kinds": ["mood"]}, False),
    ({"kind": "modality", "extra": 1}, False),
    ({"kind": "modality", "note": "how the reviewer read it"}, True),
    ({"kind": "readability", "note": "a note alone is not a threshold"}, False),
    ({"kind": "spellcheck"}, False),
])
def test_check_params(params, ok):
    assert (check_params_problems(_det(params)) == []) is ok


def test_semantic_rules_need_no_check_params():
    rule = GwpRule(rule_id="STY-001", category=RuleCategory.STYLE, text="x", params={"anything": 1})
    assert check_params_problems(rule) == []


def test_baseline_rules_are_valid_approved_preservation_rules():
    rules = baseline_rules()
    assert {r.rule_id for r in rules} == BASELINE_IDS
    for rule in rules:
        assert rule.category == RuleCategory.PRESERVATION
        assert rule.status == RuleStatus.APPROVED and rule.origin == RuleOrigin.BASELINE
        assert check_params_problems(rule) == []
    # baseline_rules() hands out copies
    rules[0].text = "changed"
    assert baseline_rules()[0].text != "changed"


# ── Input selection and batching ──────────────────────────────────────


def test_rule_input_skips_cover_toc_references_history():
    sections = [s.heading for s, _ in rule_input_sections(_doc())]
    assert sections == ["PURPOSE", "LANGUAGE AND WORDING", "INFOGRAPHICS"]


def test_batches_keep_whole_sections_in_order():
    doc = _doc()
    one = build_batches(doc)
    assert len(one) == 1
    assert one[0].unit_ids == ["SRC-1-U001", "SRC-6.1-U001", "SRC-6.1-U002", "SRC-6.1-U003", "SRC-6.2-U001"]
    small = build_batches(doc, max_chars=60)
    assert len(small) == 3
    assert [b.unit_ids[0] for b in small] == ["SRC-1-U001", "SRC-6.1-U001", "SRC-6.2-U001"]
    assert "## 6.1 LANGUAGE AND WORDING" in small[1].parts[0]
    assert '[SRC-6.1-U001] (paragraph) Do not use "please".' in small[1].parts[0]


# ── Post-processing ───────────────────────────────────────────────────


def _built():
    doc = _doc()
    ids = build_batches(doc)[0].unit_ids
    return build_rule_set(doc, [_output()], "GWP", 1, ids, model="fake/model")


def test_build_rule_set_numbers_validates_and_reports():
    rule_set, report = _built()
    by_id = {r.rule_id: r for r in rule_set.rules}

    # Baseline first, approved; guide rules are candidates.
    assert [r.rule_id for r in rule_set.rules[: len(BASELINE_IDS)]] == sorted(BASELINE_IDS)
    guide = [r for r in rule_set.rules if r.origin == RuleOrigin.GUIDE]
    assert all(r.status == RuleStatus.CANDIDATE for r in guide)

    # Numbered per category in guide order.
    assert by_id["STY-001"].text == "Describe the purpose briefly."
    assert by_id["STY-002"].text == 'Do not use "please".'
    assert by_id["STY-002"].params == {"kind": "forbidden_terms", "terms": ["please"]}
    assert by_id["FMT-001"].check == RuleCheck.DETERMINISTIC
    # The guide's PRES rule skips the reserved baseline IDs.
    assert by_id["PRES-007"].source_unit_ids == ["SRC-6.1-U003"]

    # Invalid deterministic params are downgraded and reported.
    assert by_id["STY-003"].check == RuleCheck.SEMANTIC
    assert report.downgraded == {"STY-003": ["missing param 'terms'"]}

    # A candidate with no valid source is dropped, not kept.
    assert "Invented rule." not in {r.text for r in rule_set.rules}
    assert report.dropped[0].text == "Invented rule." and "SRC-99-U001" in report.dropped[0].reason

    # Skips keep only known IDs; coverage counts cited and skipped units.
    assert report.skipped[0].source_unit_ids == ["SRC-6.1-U003"]
    assert report.uncovered_unit_ids == []
    assert report.prompt_version == PROMPT_VERSION and report.model == "fake/model"


def test_duplicates_across_batches_merge_citations():
    doc = _doc()
    a = GwpExtractionOutput(rules=[CandidateRule(category="STY", text="Write short sentences.", source_unit_ids=["SRC-1-U001"])])
    b = GwpExtractionOutput(rules=[CandidateRule(category="STY", text="Write short sentences", source_unit_ids=["SRC-6.1-U003"])])
    ids = build_batches(doc)[0].unit_ids
    rule_set, report = build_rule_set(doc, [a, b], "GWP", 1, ids)
    guide = [r for r in rule_set.rules if r.origin == RuleOrigin.GUIDE]
    assert len(guide) == 1 and guide[0].source_unit_ids == ["SRC-1-U001", "SRC-6.1-U003"]
    assert set(report.uncovered_unit_ids) == {"SRC-6.1-U001", "SRC-6.1-U002", "SRC-6.2-U001"}


def test_bad_params_json_is_reported():
    doc = _doc()
    out = GwpExtractionOutput(rules=[CandidateRule(
        category="STY", text="x", check="deterministic", params_json="{not json", source_unit_ids=["SRC-1-U001"],
    )])
    rule_set, report = build_rule_set(doc, [out], "GWP", 1, ["SRC-1-U001"])
    assert rule_set.rules[-1].check == RuleCheck.SEMANTIC and rule_set.rules[-1].params == {}
    assert "params_json is not valid JSON" in report.downgraded["STY-001"][0]


async def test_extractor_calls_llm_per_batch():
    doc = _doc()
    chain = FakeChain([_output(), GwpExtractionOutput(), GwpExtractionOutput()])
    rule_set, report = await GwpRuleExtractor(chain, model_name="fake", max_chars=60).extract(doc, "GWP", 2, "Guide")
    assert len(chain.calls) == 3
    system, human = chain.calls[1]
    assert "STR (placement" in system.content and "forbidden_terms" in system.content
    assert "ordered_procedure" in system.content  # content types listed
    assert "PART 2 of 3" in human.content and "[SRC-6.1-U001]" in human.content
    assert "[SRC-1-U001]" not in human.content
    assert rule_set.version == 2 and report.version == 2


# ── Review, approval and selection ────────────────────────────────────


def test_validate_rule_set():
    rule_set, _ = _built()
    doc = _doc()
    assert service.validate_rule_set(rule_set, "GWP", 1, doc) == []
    assert "expected GWP v2" in service.validate_rule_set(rule_set, "GWP", 2, doc)[0]

    bad = rule_set.model_copy(update={"rules": rule_set.rules + [
        GwpRule(rule_id="STR-009", category="STR", text="no source"),
        GwpRule(rule_id="STR-010", category="STR", text="ghost", source_unit_ids=["SRC-X-U001"]),
        GwpRule(rule_id="STR-011", category="STR", text="manual", origin="manual"),
        GwpRule(rule_id="PRES-099", category="PRES", text="fake baseline", origin="baseline"),
        GwpRule(rule_id="STY-099", category="STY", text="bad", check="deterministic", params={"kind": "nope"},
                source_unit_ids=["SRC-1-U001"]),
    ]})
    problems = service.validate_rule_set(bad, "GWP", 1, doc)
    assert any(p.startswith("STR-009: a guide rule must cite") for p in problems)
    assert any(p.startswith("STR-010: unknown source units") for p in problems)
    assert not any(p.startswith("STR-011") for p in problems)
    assert any(p.startswith("PRES-099: only built-in") for p in problems)
    assert any(p.startswith("STY-099: unknown check kind") for p in problems)


def test_with_baseline_restores_deleted_builtins():
    rule_set, _ = _built()
    trimmed = rule_set.model_copy(update={"rules": [r for r in rule_set.rules if r.rule_id != "PRES-004"]})
    restored = service.with_baseline(trimmed)
    assert "PRES-004" in {r.rule_id for r in restored.rules}


def test_approve_rules():
    rule_set, _ = _built()
    rejected = rule_set.model_copy(update={"rules": [
        r.model_copy(update={"status": RuleStatus.REJECTED}) if r.rule_id == "STY-001" else r for r in rule_set.rules
    ]})
    some, unknown = service.approve_rules(rejected, ["STY-002", "NOPE-1"])
    assert unknown == ["NOPE-1"]
    status = {r.rule_id: r.status for r in some.rules}
    assert status["STY-002"] == RuleStatus.APPROVED and status["FMT-001"] == RuleStatus.CANDIDATE
    assert service.review_status(some) == "in_review"

    everything, _ = service.approve_rules(rejected)
    status = {r.rule_id: r.status for r in everything.rules}
    assert status["STY-001"] == RuleStatus.REJECTED  # rejected stays rejected
    assert service.review_status(everything) == "approved"
    assert service.rule_counts(everything)["candidate_rules"] == 0


def test_select_rules():
    rule_set, _ = _built()
    approved, _ = service.approve_rules(rule_set, ["STY-002", "FMT-001"])
    style = service.select_rules(approved, [RuleCategory.STYLE])
    assert [r.rule_id for r in style] == ["STY-002"]  # candidates excluded

    scoped = approved.model_copy(update={"rules": approved.rules + [GwpRule(
        rule_id="STR-001", category="STR", text="Roles only", status="approved",
        applies_to_content_types=[ContentType.RESPONSIBILITY], source_unit_ids=["SRC-1-U001"],
    )]})
    assert [r.rule_id for r in service.select_rules(scoped, ["STR"], [ContentType.RESPONSIBILITY])] == ["STR-001"]
    assert service.select_rules(scoped, ["STR"], [ContentType.ORDERED_PROCEDURE]) == []
    # Rules without content types apply everywhere.
    assert [r.rule_id for r in service.select_rules(scoped, ["STY"], [ContentType.ORDERED_PROCEDURE])] == ["STY-002"]

    # No approved guide: only the baseline preservation rules.
    assert {r.rule_id for r in service.select_rules(None, list(RuleCategory))} == BASELINE_IDS
    assert service.select_rules(None, [RuleCategory.STYLE]) == []


# ── Store ─────────────────────────────────────────────────────────────


def test_store_versions_status_active_and_delete(tmp_path):
    store = GwpStore(tmp_path / "gwp.db", Settings(project_root=tmp_path))
    first = store.create_record("G", "Guide", "pdf")
    second = store.create_record("G", "Guide", "pdf")
    assert (first["version"], second["version"]) == (1, 2)
    assert store.get_latest("G")["id"] == second["id"]
    assert store.get_active() is None

    with pytest.raises(ValueError):
        store.update_fields(first["id"], status="done")
    with pytest.raises(ValueError):
        store.update_fields(first["id"], guide_id="other")

    rules = tmp_path / "rules.json"
    rules.write_text("{}")
    store.update_fields(first["id"], status="approved", rules_path=str(rules))
    assert store.get_active()["id"] == first["id"]
    assert [r["version"] for r in store.get_records(status="approved")] == [1]

    assert store.delete_record(first["id"]) and not rules.exists()
    assert store.get_active() is None
    assert not store.delete_record(first["id"])


# ── Ingest and API, on a synthetic .docx guide ────────────────────────


def _guide_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("PURPOSE", level=1)
    doc.add_paragraph("This Guidance outlines how to write procedure documents.")
    doc.add_heading("LANGUAGE AND WORDING", level=1)
    doc.add_paragraph('Do not use "please".', style="List Bullet")
    doc.add_paragraph('Prefer "confirm" or "verify" over "ensure".', style="List Bullet")
    doc.add_paragraph("Write short, straightforward sentences, ideally not longer than two lines.")
    doc.add_heading("REFERENCES", level=1)
    doc.add_paragraph("BI-VQD-08234-D Quality Manual")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def test_parse_guide_docx(tmp_path):
    settings = Settings(project_root=tmp_path)
    settings.resolve_paths(tmp_path)
    settings.ensure_directories()
    doc = parse_guide(_guide_docx(tmp_path / "guide.docx"), settings, "G")
    selected = rule_input_sections(doc)
    headings = [s.heading for s, _ in selected]
    assert "LANGUAGE AND WORDING" in headings and "REFERENCES" not in headings
    texts = [u.text for _, units in selected for u in units]
    assert 'Do not use "please".' in texts


@pytest.fixture
def api(tmp_path, monkeypatch):
    from app.main import app

    settings = Settings(project_root=tmp_path)
    settings.resolve_paths(tmp_path)
    settings.ensure_directories()
    store = GwpStore(tmp_path / "gwp.db", settings)
    app.dependency_overrides[get_settings] = lambda: settings
    app.dependency_overrides[require_user] = lambda: {"userId": "reviewer-1"}
    with TestClient(app) as client:
        previous = getattr(app.state, "gwp_store", None)
        app.state.gwp_store = store
        try:
            yield client, store, settings
        finally:
            app.state.gwp_store = previous
            app.dependency_overrides.pop(get_settings, None)
            app.dependency_overrides.pop(require_user, None)


def _fake_extractor(monkeypatch, outputs_for_doc):
    """Patch the factory so /extract uses a fake chain built from the parsed guide."""
    def from_factory(chain_factory, **kwargs):
        return _DocAwareExtractor(outputs_for_doc)
    monkeypatch.setattr(GwpRuleExtractor, "from_factory", classmethod(lambda cls, cf, **kw: from_factory(cf, **kw)))


class _DocAwareExtractor(GwpRuleExtractor):
    def __init__(self, outputs_for_doc):
        super().__init__(chain=None, model_name="fake/model")
        self.outputs_for_doc = outputs_for_doc

    async def extract(self, doc, guide_id, version, guide_title=""):
        self.chain = FakeChain([self.outputs_for_doc(doc)])
        return await super().extract(doc, guide_id, version, guide_title)


def test_api_upload_extract_review_approve(api, tmp_path, monkeypatch):
    client, store, settings = api

    def outputs_for(doc):
        unit = next(u for u in doc.iter_units() if u.text == 'Do not use "please".')
        return GwpExtractionOutput(rules=[CandidateRule(
            category="STY", text='Do not use "please".', check="deterministic",
            params_json='{"kind": "forbidden_terms", "terms": ["please"]}', source_unit_ids=[unit.unit_id],
        )])

    _fake_extractor(monkeypatch, outputs_for)

    bad = client.post("/api/v1/gwp", files={"file": ("guide.txt", b"x", "text/plain")})
    assert bad.status_code == 400

    path = _guide_docx(tmp_path / "in" / "My Guide.docx")
    with path.open("rb") as fh:
        up = client.post("/api/v1/gwp", files={"file": ("My Guide.docx", fh, "application/octet-stream")})
    assert up.status_code == 200, up.text
    record = up.json()
    assert (record["guide_id"], record["version"], record["status"]) == ("My_Guide", 1, "parsed")
    assert record["total_units"] > 0 and Path(record["units_path"]).exists()
    assert client.get("/api/v1/gwp/My_Guide/units").json()["document_id"] == "My_Guide"
    assert client.get("/api/v1/gwp/My_Guide/rules").status_code == 404

    ex = client.post("/api/v1/gwp/My_Guide/extract")
    assert ex.status_code == 200, ex.text
    assert ex.json()["status"] == "in_review" and ex.json()["candidate_rules"] == 1
    assert ex.json()["prompt_version"] == PROMPT_VERSION
    rules = client.get("/api/v1/gwp/My_Guide/rules").json()
    assert Path(settings.gwp_dir / "My_Guide_v1_rules.json").exists()
    report = client.get("/api/v1/gwp/My_Guide/report").json()
    assert report["uncovered_unit_ids"]  # the fake cited one unit only
    assert client.get("/api/v1/gwp/active/rules").status_code == 404

    # A reviewer edit that cites a unit the guide does not have is refused.
    edited = json.loads(json.dumps(rules))
    edited["rules"].append({"rule_id": "STR-001", "category": "STR", "text": "x", "source_unit_ids": ["SRC-404-U001"]})
    put = client.put("/api/v1/gwp/My_Guide/rules", json=edited)
    assert put.status_code == 422
    assert put.json()["error"]["code"] == "GWP_RULES_INVALID" and "unknown source units" in put.text

    # A valid edit: reword the rule and add a manual one.
    edited = json.loads(json.dumps(rules))
    sty = next(r for r in edited["rules"] if r["rule_id"] == "STY-001")
    sty["text"] = 'Never write "please".'
    edited["rules"].append({"rule_id": "STR-001", "category": "STR", "text": "One topic per section.", "origin": "manual"})
    put = client.put("/api/v1/gwp/My_Guide/rules", json=edited)
    assert put.status_code == 200, put.text
    assert put.json()["candidate_rules"] == 2

    assert client.post("/api/v1/gwp/My_Guide/approve", json={"rule_ids": ["NOPE"]}).status_code == 422
    part = client.post("/api/v1/gwp/My_Guide/approve", json={"rule_ids": ["STY-001"]})
    assert part.json()["status"] == "in_review"
    done = client.post("/api/v1/gwp/My_Guide/approve")
    assert done.status_code == 200, done.text
    assert done.json()["status"] == "approved" and done.json()["candidate_rules"] == 0

    active = client.get("/api/v1/gwp/active/rules").json()
    texts = {r["rule_id"]: r["text"] for r in active["rules"]}
    assert texts["STY-001"] == 'Never write "please".' and "PRES-004" in texts
    assert [r.rule_id for r in service.select_rules(service.load_active_rule_set(store), ["STY"])] == ["STY-001"]

    # A second upload is version 2; the approved v1 stays active.
    with path.open("rb") as fh:
        v2 = client.post("/api/v1/gwp", files={"file": ("My Guide.docx", fh, "application/octet-stream")}).json()
    assert v2["version"] == 2
    assert client.get("/api/v1/gwp/My_Guide").json()["version"] == 2
    assert client.get("/api/v1/gwp/My_Guide", params={"version": 1}).json()["status"] == "approved"
    assert client.get("/api/v1/gwp/active/rules").status_code == 200

    gone = client.delete("/api/v1/gwp/My_Guide", params={"version": 1})
    assert gone.status_code == 200 and not Path(record["units_path"]).exists()
    assert client.get("/api/v1/gwp/active/rules").status_code == 404
    assert client.get("/api/v1/gwp/nope").status_code == 404
