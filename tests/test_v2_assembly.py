"""Phase 12: the whole-document model (numbers, references), reconciliation, the LLM audit, the final export and
traceability.
"""

from __future__ import annotations

import csv
import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml.ns import qn
from langchain_core.messages import HumanMessage, SystemMessage

from app.schemas.v2 import (
    ArtifactKind,
    CrossReference,
    Gate,
    IssueCategory,
    IssueSource,
    JobStatus,
    MigrationJob,
    NumberKind,
    QualityReport,
    RefKind,
    RefStatus,
    Severity,
    ValidationIssue,
)
from app.services.llm.prompts.v2 import critic as critic_prompt
from app.services.llm.prompts.v2 import drafter as drafter_prompt
from app.services.migration_v2.assembly import assemble, present_sections
from app.services.migration_v2.assembly.xref_resolver import relative_issues, resolve_refs, show
from app.services.migration_v2.audit import AuditedChainFactory, describe
from app.services.migration_v2.drafting.drafter import draft_all
from app.services.migration_v2.export import ExportBlocked, build_traceability, export_word, traceability_csv
from app.services.migration_v2.quality.reconcile import (
    ReconcileFinding,
    ReconcileOutput,
    ReconcilePatch,
    reconcile,
    reconcile_checks,
)
from app.services.migration_v2.quality.validator import validate_drafts
from app.services.migration_v2.render import render_document
from app.services.migration_v2.render.renderer import assembled_ref_text, default_ref_text
from app.services.migration_v2.render.xml import add_runs, plain_text, ref_field, w_el
from tests.test_v2_drafter import FakeChain, FakeFactory, _claims, drafter
from tests.test_v2_renderer import _texts, scenario, tpl  # noqa: F401  (fixture)
from tests.test_v2_slot_planner import template


async def _drafts(rules=None):
    d, doc, slp = drafter(rules)
    return d, doc, slp, d.build_drafts(await draft_all(d, None), {})


# ── Numbers ───────────────────────────────────────────────────────────


async def test_numbers_follow_the_chapters_that_remain():
    d, doc, _, drafts = await _drafts()
    a = assemble(template(), drafts, doc, d.section_plan, job_id="J")
    chapters = {e.source_id: e.target_number for e in a.number_map.entries if e.kind == NumberKind.SECTION
                and e.source_id == e.target_section_id}
    # The template has no chapter 4: ROLES (TGT-5) is the 4th chapter, PROCESS the 5th.
    assert chapters["TGT-5"] == "4" and chapters["TGT-6"] == "5"
    headings = {e.source_id: e.target_number for e in a.number_map.entries if e.kind == NumberKind.HEADING}
    assert headings == {"SRC-6.1": "5.1", "SRC-6.2": "5.2"}
    assert a.document.section_order == present_sections(template(), drafts) and a.document.draft_versions["TGT-6"] == 1

    # An optional chapter with nothing in it is removed: the chapters after it move up.
    tpl_optional = template().model_copy(deep=True)
    for s in tpl_optional.sections:
        if s.section_id == "TGT-9":
            s.optional_marker = True
    b = assemble(tpl_optional, drafts, doc, d.section_plan, job_id="J")
    assert b.number_map.sections_removed == ["TGT-9"] and "TGT-9" not in b.document.section_order
    assert next(e.target_number for e in b.number_map.entries if e.source_id == "TGT-10") == "6"


async def test_references_keep_the_source_phrase_with_the_new_number():
    d, doc, _, drafts = await _drafts()
    a = assemble(template(), drafts, doc, d.section_plan, job_id="J")
    refs = {r.target: r for r in a.document.refs}
    chapter = refs["SRC-5"]  # "see section 5": the source chapter 5 is now chapter 4
    assert (chapter.text, chapter.number, chapter.bookmark, chapter.status) == ("section 4", "4", "_Ref_TGT_5", RefStatus.RESOLVED)
    assert chapter.text[chapter.number_at:].startswith("4")
    heading = refs["SRC-6.2"]
    assert (heading.source_phrase, heading.text, heading.bookmark) == ("described in 6.2", "described in 5.2", "_Ref_SRC_6_2")
    assert a.issues == []

    shown = assembled_ref_text(a.document, default_ref_text(doc))
    claim = next(c for c in _claims(drafts, "TGT-6-CONTENT") if "{{ref:SRC-5}}" in c.text)
    assert ref_field("_Ref_TGT_5", "4") in shown("SRC-5", claim)
    assert assembled_ref_text(a.document, default_ref_text(doc), fields=False)("SRC-5", claim) == "section 4"


def test_show_replaces_only_the_numbers():
    assert show("Chapter 7, no. 1", "7", "6", "1", "1") == ("Chapter 6, no. 1", 8)
    assert show("Chapter 7, no. 7", "7", "6", "7", "9") == ("Chapter 6, no. 9", 8)  # the chapter, then the entry
    assert show("see step 4", None, None, "4", "5") == ("see step 5", None)
    assert show("chapter 6.12.1", "6.12.1", "6.13.1") == ("chapter 6.13.1", 8)
    assert show("section 6.1", "6.12", "7") == ("section 6.1", None)  # "6.12" is not in "6.1"


async def test_merged_and_unresolved_targets_are_issues():
    d, doc, _, drafts = await _drafts()
    a = assemble(template(), drafts, doc, d.section_plan, job_id="J")
    merged = a.number_map.model_copy(deep=True)
    entry = merged.lookup("SRC-6.2")
    entry.exact, entry.note = False, "6.2 has no heading of its own in the new document; it is now part of 5"
    refs, issues = resolve_refs(drafts, merged, doc)
    assert next(r for r in refs if r.target == "SRC-6.2").status == RefStatus.MERGED
    assert [(i.severity, i.gate) for i in issues] == [(Severity.MEDIUM, None)] and "Check that the reference" in issues[0].message

    missing = a.number_map.model_copy(update={"entries": [e for e in a.number_map.entries if e.source_id != "SRC-5"]})
    refs, issues = resolve_refs(drafts, missing, doc)
    lost = next(r for r in refs if r.target == "SRC-5")
    assert lost.status == RefStatus.UNRESOLVED and lost.text == "section 5" and lost.bookmark is None
    assert [(i.severity, i.gate) for i in issues] == [(Severity.HIGH, Gate.BROKEN_CROSS_REFERENCE)]


async def test_relative_references_keep_their_side():
    d, doc, _, drafts = await _drafts()
    doc.cross_references.append(CrossReference(ref_id="XREF-R", from_unit_id="SRC-6.1-U003", raw_text="as described above",
                                               ref_kind=RefKind.RELATIVE, direction="before"))
    assert relative_issues(drafts, template(), doc) == []
    process = next(s for x in drafts for s in x.slots if s.slot_id == "TGT-6-CONTENT")
    bullet = next(c for c in process.claims if c.source_unit_ids == ["SRC-6.1-U002"])
    process.claims.remove(bullet)
    process.claims.append(bullet)  # the passage just above it in the source now comes after it
    issues = relative_issues(drafts, template(), doc)
    assert len(issues) == 1 and "'as described above'" in issues[0].message and issues[0].severity == Severity.MEDIUM


def test_ref_fields_become_word_fields_and_read_as_plain_text():
    text = "see Chapter " + ref_field("_Ref_TGT_7", "7") + ", no. 1"
    p = w_el("p")
    add_runs(p, text)
    assert plain_text(text) == "see Chapter 7, no. 1"
    assert [i.text for i in p.iter(qn("w:instrText"))] == [" REF _Ref_TGT_7 \\w \\h "]
    assert "".join(t.text for t in p.iter(qn("w:t"))) == "see Chapter 7, no. 1"
    assert [f.get(qn("w:fldCharType")) for f in p.iter(qn("w:fldChar"))] == ["begin", "separate", "end"]


def test_unnumbered_template_headings_get_plain_numbers(tpl, tmp_path):  # noqa: F811
    source, drafts, plan = scenario(tmp_path)
    a = assemble(tpl, drafts, source, None, job_id="J")
    report = render_document(tpl, drafts, source, tmp_path / "out.docx", job_id="J", slot_plan=plan,
                             assembled=a.document, number_map=a.number_map)
    doc = Document(str(tmp_path / "out.docx"))
    assert report.problems == [] and not [i for i in doc.element.body.iter(qn("w:instrText")) if " REF " in (i.text or "")]
    # The scenario cites "see section 4.2", which the SOP does not have: shown as written, a blocking issue.
    assert "(see section 4.2)" in _texts(doc) and a.issues[0].gate == Gate.BROKEN_CROSS_REFERENCE


# ── Reconciliation ────────────────────────────────────────────────────


async def test_reconcile_checks_abbreviations_roles_and_numbering():
    d, doc, _, drafts = await _drafts()
    a = assemble(template(), drafts, doc, d.section_plan, job_id="J")
    baseline = [i.message for i in reconcile_checks(drafts, doc, a.number_map)]
    assert baseline == ["[reconcile:abbreviation] 1 abbreviation(s) are neither in the abbreviations table nor spelled "
                        "out in the text: SOP (C-TGT-1-001)."]  # QA is in the abbreviations table
    claim = _claims(drafts, "TGT-1-WHAT")[0]
    claim.text = "This SOP describes how CC records are assessed by the Deviation Manager."
    later = _claims(drafts, "TGT-1-INTENTION")[0]
    later.text = "Quality Assessment (QA) and Change Control (CC) keep deviations under control."
    messages = [i.message for i in reconcile_checks(drafts, doc, a.number_map)]
    assert any("'CC' is first used in C-TGT-1-001 but spelled out only later, in C-TGT-1-002" in m for m in messages)
    assert any("'QA' is spelled out as 'Quality Assessment'" in m and "'Quality Assurance'" in m for m in messages)
    assert any("not in the roles table: Deviation Manager (C-TGT-1-001)" in m for m in messages)
    assert all(i.severity == Severity.LOW for i in reconcile_checks(drafts, doc, a.number_map))

    clash = a.number_map.model_copy(deep=True)
    clash.lookup("SRC-6.2").target_number = "5.1"
    issue = next(i for i in reconcile_checks(drafts, doc, clash) if i.gate)
    assert issue.gate == Gate.SEQUENCE_VIOLATION and "5.1" in issue.message


async def test_reconcile_patches_only_reworded_claims_that_still_validate():
    d, doc, slp, drafts = await _drafts()
    what = _claims(drafts, "TGT-1-WHAT")[0]
    what.text = "This SOP explains how deviations are recorded and assessed."  # as a GWP rewrite would
    deadline = next(c for c in _claims(drafts, "TGT-6-CONTENT") if "5 days" in c.text)
    deadline.text = deadline.text.replace("must record", "must log")
    verbatim = _claims(drafts, "TGT-1-INTENTION")[0]

    def respond(prompt, n):
        assert "(paragraph, reworded)" in prompt and "(paragraph, verbatim)" in prompt
        return ReconcileOutput(findings=[
            ReconcileFinding(problem="terminology", claim_ids=[what.claim_id, verbatim.claim_id], note="two names",
                             patch=ReconcilePatch(claim_id=what.claim_id, reason="one name for deviations",
                                                  text="This SOP describes how deviations are recorded and assessed by QA.")),
            ReconcileFinding(problem="contradiction", claim_ids=[deadline.claim_id, verbatim.claim_id], note="deadline",
                             patch=ReconcilePatch(claim_id=deadline.claim_id, reason="align",
                                                  text=deadline.text.replace("5 days", "6 days"))),
            ReconcileFinding(problem="duplicate", claim_ids=[verbatim.claim_id], note="said twice",
                             patch=ReconcilePatch(claim_id=verbatim.claim_id, reason="x", text="Shorter.")),
        ])

    def validate(candidate):
        return validate_drafts(candidate, slp, doc, template(), d.facts).issues

    run = await reconcile(drafts, template(), doc, validate, FakeFactory(ReconcileOutput=FakeChain(respond)))
    assert [a["claim_id"] for a in run.applied] == [what.claim_id]
    patched = run.patched["TGT-1"]
    assert patched.origin.value == "reconcile" and _claims([patched], "TGT-1-WHAT")[0].text.endswith("assessed by QA.")
    assert "verbatim" not in refused_reason(run, what.claim_id)
    refused = {r["claim_id"]: r["refused"] for r in run.rejected}
    assert "fails validation" in refused[deadline.claim_id]  # a changed number is a numerical_change gate
    assert "verbatim" in refused[verbatim.claim_id]
    kinds = {i.message.split("]")[0] for i in run.issues}
    assert kinds == {"[reconcile:patched", "[reconcile:contradiction", "[reconcile:duplicate"}
    assert all(i.source == IssueSource.CRITIC for i in run.issues)
    contradiction = next(i for i in run.issues if "contradiction" in i.message)
    assert contradiction.severity == Severity.MEDIUM and "(patch refused" in contradiction.message


def refused_reason(run, claim_id: str) -> str:
    return next((r["refused"] for r in run.rejected if r["claim_id"] == claim_id), "")


async def test_reconcile_without_an_llm_is_skipped():
    _, doc, _, drafts = await _drafts()
    run = await reconcile(drafts, template(), doc, lambda c: [], None)
    assert run.skipped == "no LLM configured" and not run.issues and not run.patched


# ── Audit ─────────────────────────────────────────────────────────────


def test_describe_recognises_the_prompt_and_its_inputs():
    first = describe([SystemMessage(content=drafter_prompt.REWRITE_SYSTEM_PROMPT),
                      HumanMessage(content="### PURPOSE.what\n[SRC-1-U001] (paragraph) text\n[SRC-6.1.2-U003] x")])
    assert first["task"] == "drafter_rewrite" and first["prompt_version"] == drafter_prompt.PROMPT_VERSION
    assert first["unit_ids"] == ["SRC-1-U001", "SRC-6.1.2-U003"] and not first["retry"]
    second = describe([SystemMessage(content=critic_prompt.SYSTEM_PROMPT),
                       HumanMessage(content="[C-TGT-6-012] (paragraph) x\n\nYour previous answer could not be read.")])
    assert second["task"] == "critic" and second["claim_ids"] == ["C-TGT-6-012"] and second["retry"]


async def test_audited_factory_records_every_call_and_failure():
    calls = {"n": 0}

    def respond(prompt, n):
        calls["n"] += 1
        if n == 1:
            return RuntimeError("rate limited")
        return ReconcileOutput()

    factory = AuditedChainFactory(FakeFactory(ReconcileOutput=FakeChain(respond)))
    chain = factory.create_structured_planner(ReconcileOutput, include_raw=True)
    messages = [SystemMessage(content=drafter_prompt.SPLIT_SYSTEM_PROMPT), HumanMessage(content="R1: passage [SRC-2-U001]: x")]
    with pytest.raises(RuntimeError):
        await chain.ainvoke(messages)
    await chain.ainvoke(messages)
    assert [c["status"] for c in factory.calls] == ["error", "ok"]
    ok = factory.calls[1]
    assert ok["task"] == "drafter_split" and ok["model"] == "azure/fake" and ok["unit_ids"] == ["SRC-2-U001"]
    assert ok["input_tokens"] == 500 and ok["latency_ms"] >= 0 and factory.calls[0]["error"] == "rate limited"


async def test_a_job_writes_one_event_per_llm_call(tmp_path):
    from tests.test_v2_quality import flag_first, gwp_env

    env, orch, _ = gwp_env(tmp_path, flag_first(rounds=1))
    job = await orch.wait(env.create(orch, gwp_id="GWP").job_id)
    calls = [e for e in env.store.list_events(job.job_id) if e["event"] == "llm_call"]
    tasks = {e["detail"]["task"] for e in calls}
    assert {"drafter_rewrite", "critic", "repair"} <= tasks
    assert all(e["detail"]["prompt_version"] and e["detail"]["stage"] for e in calls)
    assert any(e["detail"]["unit_ids"] for e in calls if e["detail"]["task"] == "drafter_rewrite")


# ── Export and traceability ───────────────────────────────────────────


def _export_env(tmp_path, tpl, gap=True):  # noqa: F811
    from app.config.settings import Settings
    from app.services.migration_v2.artifacts import ArtifactWriter
    from app.stores.migration_store import MigrationStore
    from app.stores.sop_store import SopStore

    settings = Settings(project_root=tmp_path)
    settings.resolve_paths(tmp_path)
    settings.ensure_directories()
    store = MigrationStore(tmp_path / "m.db", settings)
    writer = ArtifactWriter(tmp_path / "artifacts", store)
    sops = SopStore(tmp_path / "sop.db", settings)
    sop = sops.upsert_record(job_id="x", document_uid="SOP-T")
    job = store.create_job(MigrationJob(job_id="MIG-X", sop_record_id=sop["id"], template_id="tpl", template_version=1))
    source, drafts, plan = scenario(tmp_path, gap_geography=gap)
    writer.write(job.job_id, ArtifactKind.SOURCE_MODEL, source)
    writer.write(job.job_id, ArtifactKind.TEMPLATE_MODEL, tpl)
    writer.write(job.job_id, ArtifactKind.SLOT_PLAN, plan)
    for d in drafts:
        writer.write(job.job_id, ArtifactKind.SECTION_DRAFT, d, scope=d.target_section_id)
    gap_issue = ValidationIssue(issue_id="ISS-gap", severity=Severity.HIGH, category=IssueCategory.STRUCTURE,
                                gate=Gate.MISSING_SLOT, slot_id="TGT-2-GEOGRAPHY", target_section_id="TGT-2",
                                message="Required slot 'geography' has no source content.")
    writer.write(job.job_id, ArtifactKind.QUALITY_REPORT, QualityReport(job_id=job.job_id, version=1, issues=[gap_issue]))
    return SimpleNamespace(store=store, writer=writer, sops=sops, sop=sop, job=store.get_job(job.job_id), gap=gap_issue)


def test_export_is_refused_until_the_job_completed(tpl, tmp_path):  # noqa: F811
    e = _export_env(tmp_path, tpl)
    e.store.set_status(e.job.job_id, JobStatus.HUMAN_REVIEW_REQUIRED, "system")
    with pytest.raises(ExportBlocked) as blocked:
        export_word(e.store.get_job(e.job.job_id), e.store, e.writer)
    assert blocked.value.code == "EXPORT_BLOCKED" and blocked.value.details == [e.gap.message]

    # Completed, but the gap was never accepted: the final render refuses.
    e.store.set_status(e.job.job_id, JobStatus.COMPLETED_WITH_WARNINGS, "system")
    with pytest.raises(ExportBlocked) as blocked:
        export_word(e.store.get_job(e.job.job_id), e.store, e.writer)
    assert blocked.value.code == "GAPS_UNRESOLVED" and blocked.value.details == ["TGT-2-GEOGRAPHY"]


def test_export_removes_accepted_gaps_records_them_and_marks_the_sop(tpl, tmp_path):  # noqa: F811
    e = _export_env(tmp_path, tpl)
    accepted = e.gap.model_copy(update={"resolved": True, "resolution_note": "reviewer-1: N/A"})
    e.writer.write(e.job.job_id, ArtifactKind.QUALITY_REPORT, QualityReport(job_id=e.job.job_id, version=2, issues=[accepted]))
    e.store.set_status(e.job.job_id, JobStatus.COMPLETED_WITH_WARNINGS, "system")
    job = e.store.get_job(e.job.job_id)

    result = export_word(job, e.store, e.writer, e.sops, "reviewer-1")
    assert not result.reused and result.report.ok and result.report.mode.value == "final"
    assert result.docx.scope == "final" and result.docx.created_by == "reviewer-1"
    doc = Document(io.BytesIO(e.writer.read_bytes(result.docx)))
    assert "Source content not found" not in _texts(doc) and "World-wide" not in _texts(doc)
    assert not [s for s in doc.element.body.iter(qn("w:sdt"))
                if (s.find(f"{qn('w:sdtPr')}/{qn('w:tag')}") is not None
                    and s.find(f"{qn('w:sdtPr')}/{qn('w:tag')}").get(qn("w:val")).startswith("CC_"))]
    events = e.store.list_events(job.job_id)
    assert [x["detail"] for x in events if x["event"] == "gap_removed"] == [
        {"slot_id": "TGT-2-GEOGRAPHY", "issue_id": "ISS-gap", "note": "reviewer-1: N/A"}]
    export = next(x for x in events if x["event"] == "export")
    assert export["actor"] == "reviewer-1" and export["detail"]["accepted_gaps"] == ["TGT-2-GEOGRAPHY"]
    job = e.store.get_job(job.job_id)
    assert {a.kind for a in job.artifacts if a.scope == "final"} == {
        ArtifactKind.DOCX, ArtifactKind.RENDER_REPORT, ArtifactKind.NUMBER_MAP, ArtifactKind.ASSEMBLED_DRAFT}
    assert job.latest_artifact(ArtifactKind.TRACEABILITY) is not None
    assert e.sops.get_record_by_id(e.sop["id"])["migrated_path"] == result.docx.path

    again = export_word(job, e.store, e.writer, e.sops, "reviewer-1")
    assert again.reused and again.docx.version == result.docx.version  # nothing changed: the same file


def test_traceability_maps_every_claim_to_its_source(tpl, tmp_path):  # noqa: F811
    source, drafts, _ = scenario(tmp_path)
    a = assemble(tpl, drafts, source, None, job_id="J")
    trace = build_traceability("J", drafts, source, tpl, a.document)
    rows = {(r.claim_id, r.unit_id): r for r in trace.rows}
    step = rows[("C-5-002", "SRC-6-U002")]
    assert (step.target_section_id, step.target_number, step.slot_id, step.kind) == ("TGT-5", "5", "TGT-5-CONTENT", "step")
    assert step.source_section_heading == "PROCEDURE" and step.unit_type == "procedure_step"
    assert rows[("C-5-008", "SRC-6.1-U001")].target_number == "5.1"  # under its sub-heading
    assert rows[("C-5-001", "SRC-6-U001")].text == "Record every deviation within 5 days (see section 4.2)."
    heading = next(r for r in trace.rows if r.kind == "heading")
    assert heading.unit_id is None and heading.source_section_id == "SRC-6.1"
    rows_back = list(csv.DictReader(io.StringIO(traceability_csv(trace))))
    assert len(rows_back) == len(trace.rows) and rows_back[0]["claim_id"] == trace.rows[0].claim_id
    json.dumps(trace.to_clean_dict())


def test_api_export_and_edit_diff(tmp_path):
    from fastapi.testclient import TestClient

    from app.api.auth import require_user
    from app.config.settings import get_settings
    from app.main import app
    from tests.test_v2_migration_jobs import Env, _wait_for

    env = Env(tmp_path)
    app.dependency_overrides[get_settings] = lambda: env.settings
    app.dependency_overrides[require_user] = lambda: {"userId": "reviewer-1"}
    try:
        with TestClient(app) as client:
            previous = app.state.migration_orchestrator
            orch = env.orchestrator()
            app.state.migration_orchestrator = orch
            try:
                base = "/api/v1/migrations"
                job_id = client.post(base, json={"sop_record_id": env.sop["id"], "template_id": "TPL", "mode": "auto"}).json()["job_id"]
                _wait_for(client, job_id, {"HUMAN_REVIEW_REQUIRED"})
                refused = client.get(f"{base}/{job_id}/export/word")
                assert refused.status_code == 409 and refused.json()["error"]["code"] == "EXPORT_BLOCKED"

                draft = client.get(f"{base}/{job_id}/sections/TGT-3/draft").json()
                slot = next(s for s in draft["slots"] if s["claims"] and not s["claims"][0].get("is_gap_marker"))
                edited = [{**slot["claims"][0], "text": slot["claims"][0]["text"] + " Reviewed."}] + slot["claims"][1:]
                assert client.patch(f"{base}/{job_id}/sections/TGT-3/slots/{slot['slot_id']}",
                                     json={"claims": edited}).status_code == 200
                event = next(e for e in client.get(f"{base}/{job_id}/audit").json() if e["event"] == "slot_edited")
                assert event["actor"] == "reviewer-1" and event["detail"]["slot"] == slot["slot_id"]
                change = event["detail"]["changed"][0]
                assert change["after"] == change["before"] + " Reviewed." and change["fields"] == ["text"]
                assert event["detail"]["added"] == [] and event["detail"]["removed"] == []
            finally:
                client.portal.call(orch.stop)
                app.state.migration_orchestrator = previous
    finally:
        app.dependency_overrides.pop(get_settings, None)
        app.dependency_overrides.pop(require_user, None)


async def test_the_reconciling_stage_validates_and_assembles_the_patched_draft(tmp_path):
    from app.schemas.v2 import SectionDraft
    from tests.test_v2_quality import gwp_env

    env, orch, _ = gwp_env(tmp_path, lambda prompt, n: __import__("app.services.migration_v2.quality.critic",
                                                                  fromlist=["CriticOutput"]).CriticOutput())

    def respond(prompt, n):
        claim = next(line for line in prompt.splitlines() if "(paragraph, reworded)" in line)
        claim_id, text = claim[1:claim.index("]")], claim.split(") ", 1)[1]
        return ReconcileOutput(findings=[ReconcileFinding(
            problem="terminology", claim_ids=[claim_id], note="one name",
            patch=ReconcilePatch(claim_id=claim_id, text=text.rstrip(".") + " (reconciled).", reason="one name"))])

    orch.chain_factory.chains["ReconcileOutput"] = FakeChain(respond)
    job = await orch.wait(env.create(orch, gwp_id="GWP").job_id)
    draft = orch.artifacts.latest(job, ArtifactKind.SECTION_DRAFT, SectionDraft, "TGT-3")
    assert draft.origin.value == "reconcile" and any(c.text.endswith("(reconciled).") for c in draft.iter_claims())
    detail = next(e["detail"] for e in env.store.list_events(job.job_id) if e["event"] == "reconcile")
    assert len(detail["patches_applied"]) == 1 and detail["assembled_version"] == 2  # assembled again after the patch
    assembled = [a for a in job.artifacts if a.kind == ArtifactKind.ASSEMBLED_DRAFT and a.scope is None]
    assert len(assembled) == 2 and any(e["detail"]["task"] == "reconcile" for e in env.store.list_events(job.job_id)
                                       if e["event"] == "llm_call")
