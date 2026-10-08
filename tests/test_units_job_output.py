"""The extraction job writes the v2 SourceDocument next to the v3.1 output (Phase 2b)."""

import json
import sqlite3

import pytest
from docx import Document

from app.config.settings import Settings
from app.schemas.jobs import JobStatus
from app.schemas.v2 import SourceDocument
from app.services.job_manager import JobManager
from app.stores.sop_store import SopStore


@pytest.fixture
def settings(tmp_path):
    s = Settings(project_root=tmp_path)
    s.resolve_paths(tmp_path)
    s.ensure_directories()
    return s


@pytest.mark.asyncio
async def test_job_writes_versioned_units_and_records_path(settings, tmp_path):
    doc = Document()
    doc.add_heading("1 PURPOSE", level=1)
    doc.add_paragraph("This SOP describes sample handling.")
    doc.add_heading("2 PROCESS", level=1)
    doc.add_paragraph("Submit the record as described in chapter 1.")
    doc.save(settings.upload_dir / "units_job_doc.docx")

    store = SopStore(tmp_path / "sop_records.db", settings)
    manager = JobManager(settings, concurrency=1, sop_store=store)
    job = manager.create_batch_job(["units_job_doc"])
    await manager._process_document(job.job_id, "units_job_doc")

    assert job.documents[0].status == JobStatus.COMPLETED, job.documents[0].error
    (record,) = store.get_records()
    assert record["units_path"]
    units_file = settings.output_dir / f"units_job_doc_v{record['gpdat_version']}_units.json"
    assert record["units_path"] == str(units_file)
    source = SourceDocument.model_validate(json.loads(units_file.read_text(encoding="utf-8")))
    assert [s.heading for s in source.sections] == ["PURPOSE", "PROCESS"]
    assert source.cross_references[0].target == "SRC-1"
    assert (settings.output_dir / "units_job_doc_v2.json").exists()  # v3.1 output unchanged


def test_units_path_column_is_added_to_an_existing_database(settings, tmp_path):
    db = tmp_path / "old.db"
    with sqlite3.connect(db) as conn:
        conn.execute(
            "CREATE TABLE sop_records (id INTEGER PRIMARY KEY AUTOINCREMENT, job_id TEXT NOT NULL, "
            "document_uid TEXT NOT NULL, document_number TEXT, document_name TEXT, document_title TEXT, "
            "document_version TEXT, document_type TEXT, file_type TEXT, language TEXT DEFAULT 'en', "
            "page_count INTEGER DEFAULT 0, gpdat_version INTEGER DEFAULT 0, status TEXT NOT NULL DEFAULT 'in_review', "
            "source_filename TEXT, output_path TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL, "
            "UNIQUE(document_uid, gpdat_version))"
        )
    store = SopStore(db, settings)
    record = store.upsert_record(job_id="j", document_uid="d", units_path="out/d_v1_units.json")
    assert record["units_path"] == "out/d_v1_units.json"
