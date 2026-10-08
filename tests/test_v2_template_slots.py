"""Phase 3: template slot detection, normalization, overrides and readiness, on synthetic templates."""

import hashlib
import json
import sqlite3
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, RGBColor
from fastapi.testclient import TestClient
from PIL import Image

from app.config.settings import Settings, get_settings
from app.schemas.v2 import (
    AnchorKind,
    CalloutKind,
    ContentType,
    FormattingProfile,
    InstructionBehavior,
    ReadinessStatus,
)
from app.services.migration_v2.template.normalizer import normalize_template
from app.services.migration_v2.template.overrides import (
    TemplateConfig,
    apply_slot_overrides,
    load_config,
    save_config,
)
from app.services.migration_v2.template.readiness import assess
from app.services.migration_v2.template.service import (
    build_template_model,
    link_instruction_ids,
    normalize_and_save,
)
from app.services.migration_v2.template.slot_detector import detect_template_model
from app.services.parser.ooxml import element_text, iter_body_blocks
from app.stores.template_store import TemplateStore

BLUE = RGBColor(0x00, 0x75, 0xFF)


# ── Builders ───────────────────────────────────────────────────────────

def _png(path: Path, color) -> Path:
    Image.new("RGB", (16, 16), color).save(path)
    return path


def _blue(paragraph, text):
    paragraph.add_run(text).font.color.rgb = BLUE
    return paragraph


def _shade(cell, fill):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _icon_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    for row, (icon, text) in zip(table.rows, rows):
        row.cells[0].paragraphs[0].add_run().add_picture(str(icon), width=Inches(0.3))
        if text:
            _blue(row.cells[1].paragraphs[0], text)
    return table


def _box(doc, rows_spec, header=None):
    table = doc.add_table(rows=len(rows_spec) + (1 if header else 0), cols=2)
    rows = list(table.rows)
    if header:
        rows[0].cells[0].text, rows[0].cells[1].text = header
        rows = rows[1:]
    for row, (fill, icon, text, blue) in zip(rows, rows_spec):
        row.cells[0].paragraphs[0].add_run().add_picture(str(icon), width=Inches(0.22))
        run = row.cells[1].paragraphs[0].add_run(text)
        if blue:
            run.font.color.rgb = BLUE
        for cell in row.cells:
            _shade(cell, fill)
    return table


def _bookmark(paragraph, name):
    p = paragraph._p
    p.insert(0 if p.pPr is None else 1, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="7" w:name="{name}"/>'))
    p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="7"/>'))


def _template(path: Path, tmp: Path, *, icon_text=True) -> Path:
    what = _png(tmp / "what.png", (30, 90, 200))
    why = _png(tmp / "why.png", (200, 40, 40))
    geo = _png(tmp / "geo.png", (40, 160, 60))
    attn = _png(tmp / "attn.png", (220, 120, 40))
    doc = Document()
    _blue(doc.add_paragraph(), "Follow the blue text and delete it before finalization.")  # preamble

    doc.add_heading("PURPOSE", level=1)
    _blue(doc.add_paragraph(), "Describe the document in the rows below.")
    _icon_table(doc, [
        (what, "Brief description of what the document is about?"),
        (why, "Brief description of the intention. What do you want to achieve?" if icon_text else ""),
    ])

    doc.add_heading("APPLICABILITY", level=1)
    mixed = doc.add_paragraph("This ")
    _blue(mixed, "SOP/Work Instruction")
    mixed.add_run(" is applicable:")
    _icon_table(doc, [(geo, "The applicable geography, e.g., world-wide, country or site(s).")])
    _blue(doc.add_paragraph(), "Where possible identify topics not covered by this document.")

    doc.add_heading("DEFINITIONS & ABBREVIATIONS", level=1)
    _blue(doc.add_paragraph(), "If no terms are defined, please insert (None)")
    terms = doc.add_table(rows=3, cols=2)
    terms.rows[0].cells[0].text, terms.rows[0].cells[1].text = "Term", "Definition"
    terms.cell(2, 0).merge(terms.cell(2, 1))
    doc.add_paragraph()
    _box(doc, [("F5CDB9", attn, "Attention – notice taken of something important.", False)],
         header=("Infographics", "Description"))

    doc.add_heading("ROLES & RESPONSIBILITIES", level=1)
    _blue(doc.add_paragraph(), "List all roles that have responsibilities in the process.")
    competence = doc.add_table(rows=2, cols=2)
    for row, (a, b) in zip(competence.rows, (("Competence", "Description"), ("Execution", "Ability to execute."))):
        _blue(row.cells[0].paragraphs[0], a)
        _blue(row.cells[1].paragraphs[0], b)
    _blue(doc.add_paragraph(), "Table A: list roles and their responsibilities.")
    roles = doc.add_table(rows=2, cols=2)
    roles.rows[0].cells[0].text, roles.rows[0].cells[1].text = "Role", "Responsibility"
    _blue(roles.rows[1].cells[0].paragraphs[0], "Role 1")
    _blue(roles.rows[1].cells[1].paragraphs[0], "…")

    doc.add_heading("PROCESS", level=1)
    _blue(doc.add_paragraph(), "Describe the process steps.")
    _blue(doc.add_paragraph(), "Use the Attention infographic as needed:")
    _box(doc, [("F5CDB9", attn, "Attention section (what you should know). It is pink.", True)])

    doc.add_heading("OWNER", level=1)
    doc.add_paragraph("Owner: [Insert owner name]")
    _bookmark(doc.add_paragraph("Deputy"), "SLOT_OWNER_DEPUTY")
    doc.save(path)
    return path


@pytest.fixture
def template(tmp_path):
    return _template(tmp_path / "template.docx", tmp_path)


def _region(detection, kind):
    return next(a.region_id for a in detection.ambiguous if a.kind == kind)


def _decide_all(detection, decision_by_kind=None) -> TemplateConfig:
    decision_by_kind = decision_by_kind or {}
    return TemplateConfig(regions={
        a.region_id: decision_by_kind.get(a.kind, a.suggestion.value) for a in detection.ambiguous
    })


def _slot_summary(model):
    return [
        (s.slot_id, s.key, s.content_type, s.required, s.instruction_behavior, s.formatting_profile,
         s.callout_kind, s.instruction, s.icon)
        for s in model.iter_slots()
    ]


# ── Detection ──────────────────────────────────────────────────────────

def test_sections_get_keys_and_numbers(template):
    model = detect_template_model(template, "tpl").model
    assert [s.key for s in model.sections] == [
        "PURPOSE", "APPLICABILITY", "DEFINITIONS", "ROLES", "PROCESS", "OWNER"
    ]
    assert [s.section_id for s in model.sections] == [f"TGT-{n}" for n in range(1, 7)]
    assert all(s.required for s in model.sections)
    assert "Follow the blue text" not in json.dumps(model.to_clean_dict())  # preamble is not a slot


def test_icon_rows_become_cell_slots_with_their_icon(template):
    model = detect_template_model(template, "tpl").model
    what, intention = model.slot_by_ref("PURPOSE.what"), model.slot_by_ref("PURPOSE.intention")
    for row, slot in enumerate((what, intention)):
        assert slot.anchor.kind == AnchorKind.TABLE_CELL
        assert (slot.anchor.row_index, slot.anchor.column_index) == (row, 1)
        assert slot.instruction_behavior == InstructionBehavior.REPLACE
        assert slot.content_type == ContentType.PURPOSE
    assert what.icon.icon_key != intention.icon.icon_key
    assert model.sections[0].purpose_instruction == "Describe the document in the rows below."
    geography = model.slot_by_ref("APPLICABILITY.geography")
    assert geography.content_type == ContentType.SCOPE and geography.icon is not None


def test_trailing_instruction_is_its_own_paragraph_slot(template):
    slot = detect_template_model(template, "tpl").model.slot_by_ref("APPLICABILITY.not_covered")
    assert slot.anchor.kind == AnchorKind.PARAGRAPH
    assert slot.icon is None
    assert slot.required is False  # "Where possible ..."


def test_instruction_group_leads_into_the_table_after_it(template):
    model = detect_template_model(template, "tpl").model
    terms = model.slot_by_ref("DEFINITIONS.terms")
    assert terms.instruction == "If no terms are defined, please insert (None)"
    assert terms.instruction_behavior == InstructionBehavior.HIDE_AFTER_POPULATION
    assert terms.formatting_profile == FormattingProfile.TABLE
    assert terms.required is True
    assert (terms.anchor.kind, terms.anchor.row_index) == (AnchorKind.TABLE_CELL, 1)
    assert [s.key for s in model.sections[2].slots] == ["terms"]  # the legend is not a slot


def test_callout_prototype_is_a_fixed_callout_slot(template):
    model = detect_template_model(template, "tpl").model
    callout = model.slot_by_ref("PROCESS.callout_attention")
    assert callout.callout_kind == CalloutKind.ATTENTION
    assert callout.formatting_profile == FormattingProfile.CALLOUT
    assert callout.content_type == ContentType.WARNING
    assert callout.required is False
    assert callout.instruction.startswith("Use the Attention infographic")
    content = model.slot_by_ref("PROCESS.content")
    assert content.instruction == "Describe the process steps."  # the lead-in went to the callout
    assert content.content_type == ContentType.ORDERED_PROCEDURE


def test_placeholder_and_bookmark_anchors(template):
    model = detect_template_model(template, "tpl").model
    owner = model.sections[-1].slots
    placeholder = next(s for s in owner if s.anchor.kind == AnchorKind.PLACEHOLDER)
    assert placeholder.anchor.ref == "[Insert owner name]"
    deputy = model.slot_by_ref("OWNER.deputy")
    assert deputy.anchor.kind == AnchorKind.BOOKMARK and deputy.anchor.ref == "SLOT_OWNER_DEPUTY"


def test_ambiguous_regions_are_reported_and_config_decides_them(template):
    detection = detect_template_model(template, "tpl")
    kinds = sorted(a.kind for a in detection.ambiguous)
    assert kinds == ["inline_instruction", "instruction_table"]
    assert all(a.decision is None for a in detection.ambiguous)
    roles = detection.model.slot_by_ref("ROLES.roles")
    assert "Ability to execute" not in roles.instruction

    config = TemplateConfig(regions={
        _region(detection, "instruction_table"): "instruction",
        _region(detection, "inline_instruction"): "slot",
    })
    decided = detect_template_model(template, "tpl", config=config)
    assert decided.unresolved == []
    assert "Ability to execute" in decided.model.slot_by_ref("ROLES.roles").instruction
    applicability = [s.key for s in decided.model.sections[1].slots]
    assert applicability == ["content", "geography", "not_covered"]


def test_content_controls_are_detected_first(tmp_path):
    doc = Document()
    doc.add_heading("PURPOSE", level=1)
    paragraph = doc.add_paragraph("Describe what the SOP covers.")
    sdt = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:tag w:val="CC_PURPOSE_WHAT"/><w:id w:val="5"/></w:sdtPr>'
        "<w:sdtContent/></w:sdt>"
    )
    paragraph._p.addprevious(sdt)
    sdt.find(qn("w:sdtContent")).append(paragraph._p)
    doc.add_heading("SCOPE", level=1)
    other = doc.add_paragraph("Sites.")
    sdt2 = parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtPr><w:tag w:val="CC_PURPOSE_SITES"/></w:sdtPr><w:sdtContent/></w:sdt>'
    )
    other._p.addprevious(sdt2)
    sdt2.find(qn("w:sdtContent")).append(other._p)
    doc.save(tmp_path / "cc.docx")

    detection = detect_template_model(tmp_path / "cc.docx", "tpl")
    what = detection.model.slot_by_ref("PURPOSE.what")
    assert (what.anchor.kind, what.anchor.ref) == (AnchorKind.CONTENT_CONTROL, "CC_PURPOSE_WHAT")
    assert what.instruction == "Describe what the SOP covers."
    assert detection.model.slot_by_ref("SCOPE.purpose_sites") is not None
    assert any("CC_PURPOSE_SITES" in issue for issue in detection.issues)


# ── Normalization ──────────────────────────────────────────────────────

def test_normalization_wraps_every_slot_and_round_trips(template, tmp_path):
    original = hashlib.md5(template.read_bytes()).hexdigest()
    detection = detect_template_model(template, "tpl")
    before = _slot_summary(detection.model)
    result = normalize_template(detection, tmp_path / "normalized.docx")

    assert hashlib.md5(template.read_bytes()).hexdigest() == original
    assert result.skipped == []
    assert all(s.anchor.kind in (AnchorKind.CONTENT_CONTROL, AnchorKind.BOOKMARK) for s in detection.model.iter_slots())

    again = detect_template_model(tmp_path / "normalized.docx", "tpl")
    assert _slot_summary(again.model) == before
    assert [s.anchor.ref for s in again.model.iter_slots()] == [s.anchor.ref for s in detection.model.iter_slots()]
    assert detection.model.slot_by_ref("APPLICABILITY.geography").anchor.ref == "CC_APPLICABILITY_GEOGRAPHY"
    assert sorted(a.region_id for a in again.ambiguous) == sorted(a.region_id for a in detection.ambiguous)

    text = lambda p: [element_text(b) for b in iter_body_blocks(Document(str(p)).element.body)]
    assert text(tmp_path / "normalized.docx") == text(template)


def test_normalized_data_rows_each_sit_in_a_row_level_control(template, tmp_path):
    detection = detect_template_model(template, "tpl")
    normalize_template(detection, tmp_path / "normalized.docx")
    body = Document(str(tmp_path / "normalized.docx")).element.body
    by_tag = {}
    for sdt in body.iter(qn("w:sdt")):
        by_tag.setdefault(sdt.find(f"{qn('w:sdtPr')}/{qn('w:tag')}").get(qn("w:val")), []).append(sdt)
    rows = by_tag["CC_DEFINITIONS_TERMS"]
    # One control per row (the empty row and the merged one): Word splits a multi-row control into its own table.
    assert len(rows) == 2
    assert all(r.getparent().tag == qn("w:tbl") and len(r.find(qn("w:sdtContent")).findall(qn("w:tr"))) == 1
               for r in rows)
    assert by_tag["CC_PURPOSE_WHAT"][0].getparent().tag == qn("w:tc")
    ids = [s.find(f"{qn('w:sdtPr')}/{qn('w:id')}").get(qn("w:val")) for sdts in by_tag.values() for s in sdts]
    assert len(ids) == len(set(ids))


# ── Readiness and overrides ────────────────────────────────────────────

def test_readiness_moves_from_review_to_ready(template, tmp_path):
    undecided = build_template_model(template, "tpl")
    assert undecided.report.status == ReadinessStatus.NEEDS_REVIEW
    assert {i.code for i in undecided.report.blocking} == {"ambiguous_region", "paragraph_anchor"}

    config = _decide_all(undecided.detection)
    decided = build_template_model(template, "tpl", config=config)
    assert decided.report.status == ReadinessStatus.NEEDS_NORMALIZATION

    out = normalize_and_save(template, "tpl", 1, tmp_path / "out", config)
    assert out.build.report.status == ReadinessStatus.READY
    assert out.build.model.readiness == ReadinessStatus.READY
    assert out.build.model.source_file == str(template)
    assert out.build.model.normalized_file == str(out.normalized_path)
    saved = json.loads(out.model_path.read_text(encoding="utf-8"))
    assert saved["readiness"] == "ready"
    assert json.loads(out.report_path.read_text(encoding="utf-8"))["status"] == "ready"


def test_icon_row_without_instruction_needs_review(tmp_path):
    path = _template(tmp_path / "t.docx", tmp_path, icon_text=False)
    build = build_template_model(path, "tpl")
    assert any(i.code == "icon_without_instruction" for i in build.report.blocking)
    assert build.detection.model.slot_by_ref("PURPOSE.intention") is None


def test_slot_overrides_and_unknown_refs(template):
    detection = detect_template_model(template, "tpl")
    config = TemplateConfig(slots={
        "APPLICABILITY.not_covered": {"required": True, "content_type": "restriction"},
        "PURPOSE.nonexistent": {"required": False},
    })
    unknown = apply_slot_overrides(detection.model, config)
    slot = detection.model.slot_by_ref("APPLICABILITY.not_covered")
    assert slot.required is True and slot.content_type == ContentType.RESTRICTION
    assert unknown == ["PURPOSE.nonexistent"]
    report = assess(detection, unknown)
    assert any(i.code == "unknown_config_ref" for i in report.blocking)


def test_section_override_and_config_file_round_trip(template, tmp_path):
    config = TemplateConfig(sections={"OWNER": {"required": False}}, regions={"p:3": "fixed"})
    save_config(tmp_path / "cfg" / "tpl.json", config)
    loaded = load_config(tmp_path / "cfg" / "tpl.json")
    assert loaded == config
    assert load_config(tmp_path / "missing.json") == TemplateConfig()
    model = detect_template_model(template, "tpl", config=loaded).model
    owner = next(s for s in model.sections if s.key == "OWNER")
    assert owner.required is False and not any(s.required for s in owner.slots)


def test_instruction_ids_link_to_the_v1_extraction(template):
    model = detect_template_model(template, "tpl").model
    v1 = {"sections": [{"title": "PROCESS", "authoring_instructions": [
        {"instruction_id": "INS-1", "text": "Describe the process steps."},
        {"instruction_id": "INS-2", "text": "Use the Attention infographic as needed:"},
    ]}]}
    link_instruction_ids(model, v1)
    assert model.slot_by_ref("PROCESS.content").instruction_ids == ["INS-1"]
    assert model.slot_by_ref("PROCESS.callout_attention").instruction_ids == ["INS-2"]


# ── Store and API ──────────────────────────────────────────────────────

def test_template_model_columns_are_added_to_an_existing_database(tmp_path):
    db = tmp_path / "old.db"
    with sqlite3.connect(db) as conn:
        conn.execute(
            "CREATE TABLE template_records (id INTEGER PRIMARY KEY AUTOINCREMENT, template_uid TEXT NOT NULL, "
            "template_name TEXT NOT NULL, template_version INTEGER DEFAULT 1, file_type TEXT DEFAULT 'docx', "
            "file_size_bytes INTEGER DEFAULT 0, source_filename TEXT, upload_path TEXT, output_path TEXT, "
            "total_sections INTEGER DEFAULT 0, total_elements INTEGER DEFAULT 0, total_instructions INTEGER DEFAULT 0, "
            "total_icons INTEGER DEFAULT 0, global_rules_json TEXT, status TEXT NOT NULL DEFAULT 'uploaded', "
            "created_at TEXT NOT NULL, updated_at TEXT NOT NULL, UNIQUE(template_uid, template_version))"
        )
    store = TemplateStore(db, Settings(project_root=tmp_path))
    record = store.upsert_record(template_uid="t", template_name="T")
    updated = store.update_template_model(record["id"], "m.json", "n.docx", "ready")
    assert (updated["model_path"], updated["normalized_path"], updated["readiness_status"]) == ("m.json", "n.docx", "ready")


@pytest.fixture
def api(tmp_path):
    from app.main import app

    settings = Settings(project_root=tmp_path)
    settings.resolve_paths(tmp_path)
    settings.ensure_directories()
    store = TemplateStore(tmp_path / "templates.db", settings)
    app.dependency_overrides[get_settings] = lambda: settings
    with TestClient(app) as client:
        previous = app.state.template_store
        app.state.template_store = store
        try:
            yield client, store, settings
        finally:
            app.state.template_store = previous
            app.dependency_overrides.pop(get_settings, None)


def test_api_normalize_slots_readiness_and_config(api, tmp_path):
    client, store, settings = api
    upload = _template(settings.template_upload_dir / "tpl_v1.docx", tmp_path)
    store.upsert_record(template_uid="tpl", template_name="Template", upload_path=str(upload))

    slots = client.get("/templates/tpl/slots")
    assert slots.status_code == 200, slots.text
    assert "normalized_file" not in slots.json()  # detected on the fly, not normalized yet
    assert client.get("/templates/tpl/readiness").json()["status"] == "needs_review"

    ambiguous = [a["region_id"] for a in client.get("/templates/tpl/readiness").json()["ambiguous_regions"]]
    put = client.put("/templates/tpl/slot-config", json={"regions": {r: "fixed" for r in ambiguous}})
    assert put.status_code == 200, put.text
    assert put.json()["status"] == "needs_normalization"
    assert client.get("/templates/tpl/slot-config").json()["regions"] == {r: "fixed" for r in ambiguous}

    normalized = client.post("/templates/tpl/normalize")
    assert normalized.status_code == 200, normalized.text
    assert normalized.json()["readiness_status"] == "ready"
    record = store.get_record_by_uid("tpl")
    assert record["readiness_status"] == "ready"
    assert Path(record["normalized_path"]).exists() and Path(record["model_path"]).exists()
    assert client.get("/templates/tpl/readiness").json()["status"] == "ready"
    model = client.get("/templates/tpl/slots").json()
    assert all(s["anchor"]["kind"] in ("content_control", "bookmark") for sec in model["sections"] for s in sec["slots"])

    # A config change after normalization re-assesses the saved model.
    put = client.put("/templates/tpl/slot-config", json={"slots": {"PURPOSE.nothing": {"required": True}}})
    assert put.json()["status"] == "needs_review"
    assert store.get_record_by_uid("tpl")["readiness_status"] == "needs_review"

    assert client.put("/templates/tpl/slot-config", json={"regions": {"p:1": "maybe"}}).status_code == 422


# ── Conditional regions and one-of groups ──────────────────────────────

def test_conditional_regions_are_wrapped_and_round_trip(template, tmp_path):
    detection = detect_template_model(template, "tpl")
    config = _decide_all(detection, {"inline_instruction": "conditional", "instruction_table": "conditional"})
    decided = detect_template_model(template, "tpl", config=config)
    regions = [(s.key, r.kind.value, r.anchor.kind) for s in decided.model.sections for r in s.conditional_regions]
    assert regions == [("APPLICABILITY", "inline_choice", AnchorKind.PARAGRAPH),
                       ("ROLES", "block", AnchorKind.TABLE_CELL)]
    assert "Ability to execute" not in decided.model.slot_by_ref("ROLES.roles").instruction
    assert build_template_model(template, "tpl", config=config).report.status == ReadinessStatus.NEEDS_NORMALIZATION

    out = normalize_and_save(template, "tpl", 1, tmp_path / "out", config)
    assert out.build.report.status == ReadinessStatus.READY, out.build.report.blocking
    anchors = [r.anchor.ref for s in out.build.model.sections for r in s.conditional_regions]
    assert anchors == ["COND_APPLICABILITY_1", "COND_ROLES_1"]
    assert _slot_summary(out.build.model) == _slot_summary(decided.model)  # COND_ controls are not slots
    text = lambda p: [element_text(b) for b in iter_body_blocks(Document(str(p)).element.body)]
    assert text(out.normalized_path) == text(template)


def test_one_of_group_makes_each_table_optional(template):
    detection = detect_template_model(template, "tpl")
    config = TemplateConfig(sections={"DEFINITIONS": {"one_of": [["terms", "glossary"]]},
                                      "PURPOSE": {"one_of": [["what", "intention"]]}})
    unknown = apply_slot_overrides(detection.model, config)
    assert unknown == ["DEFINITIONS.glossary"]
    purpose = detection.model.sections[0]
    assert purpose.one_of == [["what", "intention"]]
    assert not any(s.required for s in purpose.slots)
    assert detection.model.sections[2].one_of == []
