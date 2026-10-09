"""Phase 11: anchor-based Word renderer — content at its anchors, real numbering, callout clones, tables (template
table or source shape), gaps (review marker, final removal), optional slots and sections, conditional regions,
instruction removal, the post-render check, the stage and API, and the three samples.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.schemas.v2 import (
    ArtifactKind,
    AssetKind,
    CalloutKind,
    Claim,
    ClaimKind,
    CrossReference,
    RefKind,
    RefResolution,
    RegionChoice,
    RegionOutcome,
    RenderMode,
    SectionDraft,
    SectionSlotPlan,
    SlotDraft,
    SlotOutcome,
    SlotPlan,
    SourceAsset,
    SourceDocument,
    SourceSection,
    SourceUnit,
    TableCell,
    TableRef,
    UnitType,
    compute_content_hash,
)
from app.services.migration_v2.render import GAP_MARKER_TEXT, RenderBlocked, gap_slots, render_document
from app.services.migration_v2.render.verify import verify_rendered
from app.services.migration_v2.template.colour import TextColour
from app.services.migration_v2.template.service import build_template_model, normalize_and_save
from tests.test_v2_template_slots import _decide_all, _png, _template

ROOT = Path(__file__).resolve().parent.parent


# ── Fixtures: the synthetic template, normalized, and a matching SOP ──


@pytest.fixture
def tpl(tmp_path):
    source = _template(tmp_path / "template.docx", tmp_path)
    config = _decide_all(build_template_model(source, "tpl").detection,
                         {"inline_instruction": "conditional", "instruction_table": "conditional"})
    return normalize_and_save(source, "tpl", 1, tmp_path / "out", config).build.model


class _Src:
    def __init__(self):
        self.sections: list[SourceSection] = []
        self.seq = 0

    def section(self, sid: str, heading: str, level: int = 1, parent=None) -> SourceSection:
        s = SourceSection(section_id=sid, number=sid[4:], heading=heading, level=level, section_order=len(self.sections),
                          parent_id=parent)
        self.sections.append(s)
        return s

    def unit(self, section: SourceSection, text: str, unit_type=UnitType.PARAGRAPH, **kw) -> SourceUnit:
        uid = f"{section.section_id}-U{len(section.units) + 1:03d}"
        u = SourceUnit(unit_id=uid, content_hash=compute_content_hash(text, uid), section_id=section.section_id,
                       seq=self.seq, unit_type=unit_type, text=text, **kw)
        self.seq += 1
        section.units.append(u)
        return u

    def row(self, section: SourceSection, table_id: str, header: list[str], cells: list[str], row: int, fills=None):
        fills = fills or {}
        ref = TableRef(table_id=table_id, row_index=row, header_cells=header,
                       cells=[TableCell(col=i, text=t, fill_hex=fills.get(i)) for i, t in enumerate(cells)])
        return self.unit(section, " | ".join(cells), UnitType.TABLE_ROW, table_ref=ref)

    def doc(self, refs=()) -> SourceDocument:
        return SourceDocument(document_id="SOP-T", source_file="sop.docx", file_type="docx", sections=self.sections,
                              cross_references=list(refs))


def _claim(cid: str, unit: SourceUnit, kind=ClaimKind.PARAGRAPH, **kw) -> Claim:
    return Claim(claim_id=cid, text=kw.pop("text", unit.text), kind=kind, source_unit_ids=[unit.unit_id], **kw)


def _gap(cid: str) -> Claim:
    return Claim(claim_id=cid, text="Source content not found", is_gap_marker=True)


def scenario(tmp_path, *, gap_geography=False, fig=True):
    """A source document and drafts for every slot of the synthetic template."""
    s = _Src()
    p = s.section("SRC-1", "PURPOSE")
    what = s.unit(p, "This SOP describes how deviations are recorded.")
    why = s.unit(p, "The aim is to keep deviations under control.\nAnd to learn from them.")
    a = s.section("SRC-2", "SCOPE")
    lead = s.unit(a, "This SOP is applicable:")
    geo = s.unit(a, "World-wide")
    nc = s.unit(a, "Animal health products are not covered.")
    d = s.section("SRC-3", "DEFINITIONS")
    t1 = s.row(d, "SRC-3-T1", ["Term", "Definition"], ["CAPA", "Corrective and preventive action"], 1)
    t2 = s.row(d, "SRC-3-T1", ["Term", "Definition"], ["QA", "Quality Assurance"], 2, fills={1: "FFFF00"})
    r = s.section("SRC-5", "ROLES")
    intro = s.unit(r, "The following roles take part:")
    role = s.row(r, "SRC-5-T1", ["Role", "Responsibility"], ["Process Owner", "Owns the process"], 1)
    wide = s.row(r, "SRC-5-T2", ["Role", "Task", "Deputy"], ["QA Head", "Approves", "QA Lead"], 1)
    pr = s.section("SRC-6", "PROCEDURE")
    p1 = s.unit(pr, "Record every deviation within 5 days (see section 4.2).")
    st1 = s.unit(pr, "Open the record.", UnitType.PROCEDURE_STEP, list_level=0)
    st2 = s.unit(pr, "Describe the event.", UnitType.PROCEDURE_STEP, list_level=0)
    bl = s.unit(pr, "including the batch", UnitType.BULLET, list_level=1)
    st3 = s.unit(pr, "Submit it.", UnitType.PROCEDURE_STEP, list_level=0)
    warn = s.unit(pr, "Never close a deviation without QA approval.", UnitType.WARNING)
    sub = s.section("SRC-6.1", "Assessment", level=2, parent="SRC-6")
    p2 = s.unit(sub, "QA assesses each deviation.")
    figure = None
    if fig:
        png = _png(tmp_path / "figure.png", (10, 120, 200))
        figure = s.unit(sub, "Figure 1", UnitType.FIGURE,
                        assets=[SourceAsset(kind=AssetKind.FIGURE, asset_key="icon_fig", path=str(png), width_px=40, height_px=20)])
    cap = s.unit(sub, "Figure 1: Assessment flow", UnitType.CAPTION)
    o = s.section("SRC-7", "OWNER")
    owner = s.unit(o, "Head of Quality")
    deputy = s.unit(o, "Deputy Head of Quality")
    source = s.doc([CrossReference(ref_id="R1", from_unit_id=p1.unit_id, raw_text="see section 4.2",
                                   ref_kind=RefKind.SECTION, target="SRC-4.2", resolution=RefResolution.RESOLVED)])

    process = [_claim("C-5-001", p1, text="Record every deviation within 5 days ({{ref:SRC-4.2}})."),
               _claim("C-5-002", st1, ClaimKind.STEP), _claim("C-5-003", st2, ClaimKind.STEP),
               _claim("C-5-004", bl, ClaimKind.BULLET, list_level=1), _claim("C-5-005", st3, ClaimKind.STEP),
               _claim("C-5-006", warn, callout_kind=CalloutKind.ATTENTION),
               Claim(claim_id="C-5-007", text="Assessment", kind=ClaimKind.HEADING, source_section_id="SRC-6.1"),
               _claim("C-5-008", p2)]
    if figure is not None:
        process.append(_claim("C-5-009", figure, ClaimKind.FIGURE))
    process.append(_claim("C-5-010", cap, ClaimKind.CAPTION))
    drafts = [
        SectionDraft(target_section_id="TGT-1", version=1, slots=[
            SlotDraft(slot_id="TGT-1-WHAT", claims=[_claim("C-1-001", what)]),
            SlotDraft(slot_id="TGT-1-INTENTION", claims=[_claim("C-1-002", why)])]),
        SectionDraft(target_section_id="TGT-2", version=1, slots=[
            SlotDraft(slot_id="TGT-2-GEOGRAPHY", claims=[_gap("C-2-001")] if gap_geography else [_claim("C-2-001", geo)]),
            SlotDraft(slot_id="TGT-2-NOT_COVERED", claims=[_claim("C-2-002", nc)])]),
        SectionDraft(target_section_id="TGT-3", version=1, slots=[
            SlotDraft(slot_id="TGT-3-TERMS", claims=[_claim("C-3-001", t1, ClaimKind.TABLE_ROW),
                                                     _claim("C-3-002", t2, ClaimKind.TABLE_ROW)])]),
        SectionDraft(target_section_id="TGT-4", version=1, slots=[
            SlotDraft(slot_id="TGT-4-ROLES", claims=[_claim("C-4-001", intro), _claim("C-4-002", role, ClaimKind.TABLE_ROW),
                                                     _claim("C-4-003", wide, ClaimKind.TABLE_ROW)])]),
        SectionDraft(target_section_id="TGT-5", version=1, slots=[
            SlotDraft(slot_id="TGT-5-CONTENT", claims=process),
            SlotDraft(slot_id="TGT-5-CALLOUT_ATTENTION", claims=[])]),
        SectionDraft(target_section_id="TGT-6", version=1, slots=[
            SlotDraft(slot_id="TGT-6-CONTENT", claims=[_claim("C-6-001", owner)]),
            SlotDraft(slot_id="TGT-6-DEPUTY", claims=[_claim("C-6-002", deputy)])]),
    ]
    plan = SlotPlan(job_id="J", version=1, sections=[
        SectionSlotPlan(target_section_id="TGT-2", region_choices=[RegionChoice(region_id="p:4", unit_ids=[lead.unit_id],
                                                                                choice="SOP")])])
    return source, drafts, plan


def _paragraphs(doc):
    """Every body paragraph, including those inside content controls (``doc.paragraphs`` skips them)."""
    from docx.text.paragraph import Paragraph

    return [Paragraph(p, doc._body) for p in doc.element.body.iter(qn("w:p"))]


def _texts(doc) -> str:
    """The body text, without the table of contents (Word refreshes it when the file is opened)."""
    from app.services.parser.ooxml import is_toc_sdt

    toc = {p for sdt in doc.element.body.iter(qn("w:sdt")) if is_toc_sdt(sdt) for p in sdt.iter(qn("w:p"))}
    return " ".join("".join(t.text or "" for t in p.iter(qn("w:t")))
                    for p in doc.element.body.iter(qn("w:p")) if p not in toc)


def _tables(doc):
    """Every table, including those inside content controls and cells (``doc.tables`` skips them)."""
    from docx.table import Table

    return [Table(t, doc._body) for t in doc.element.body.iter(qn("w:tbl"))]


def _cell_text(cell) -> str:
    return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in cell._tc.iter(qn("w:p")))


def _cell_text_of(el) -> str:
    return " ".join(t.text or "" for t in el.iter(qn("w:t")))


def _render(tpl, tmp_path, **kw):
    source, drafts, plan = scenario(tmp_path, **{k: kw.pop(k) for k in ("gap_geography", "fig") if k in kw})
    out = tmp_path / "rendered.docx"
    report = render_document(tpl, drafts, source, out, job_id="J", slot_plan=plan, **kw)
    return report, Document(str(out)), out


# ── Content at its anchors ────────────────────────────────────────────


def test_review_render_puts_every_claim_at_its_anchor(tpl, tmp_path):
    report, doc, _ = _render(tpl, tmp_path)
    assert report.problems == [], report.problems
    outcomes = {s.slot_id: s.outcome for s in report.slots}
    assert outcomes["TGT-1-WHAT"] == SlotOutcome.FILLED and outcomes["TGT-6-DEPUTY"] == SlotOutcome.FILLED
    assert outcomes["TGT-5-CALLOUT_ATTENTION"] == SlotOutcome.REMOVED_EMPTY

    controls = {}
    for sdt in doc.element.body.iter(qn("w:sdt")):
        tag = sdt.find(f"{qn('w:sdtPr')}/{qn('w:tag')}")
        if tag is not None:
            controls[tag.get(qn("w:val"))] = "".join(t.text or "" for t in sdt.iter(qn("w:t")))
    # Review draft: slot controls stay, filled with the claims instead of the instructions.
    assert controls["CC_PURPOSE_WHAT"] == "This SOP describes how deviations are recorded."
    assert controls["CC_APPLICABILITY_GEOGRAPHY"] == "World-wide"
    assert "Head of Quality" in controls["CC_OWNER_CONTENT"]
    assert not any(t.startswith("COND_") for t in controls)

    # The icon row keeps its icon; a multi-line claim becomes line breaks in one paragraph.
    purpose = _tables(doc)[0]
    assert purpose.rows[0].cells[0]._tc.find(f".//{qn('w:drawing')}") is not None
    paragraph = purpose.rows[1].cells[1]._tc.find(f".//{qn('w:p')}")
    assert [n.tag.split("}")[1] for n in paragraph.iter(qn("w:t"), qn("w:br"))] == ["t", "br", "t"]
    assert _cell_text(purpose.rows[1].cells[1]) == "The aim is to keep deviations under control.And to learn from them."

    text = _texts(doc)
    assert "Deputy Head of Quality" in text
    body = [p for p in _paragraphs(doc) if p._p.getparent() is doc.element.body]
    assert not any(p.text == "Deputy" for p in body)  # the bookmarked template line was replaced ...
    deputy = next(p for p in body if p.text == "Deputy Head of Quality")
    assert [b.get(qn("w:name")) for b in deputy._p.iter(qn("w:bookmarkStart"))] == ["SLOT_OWNER_DEPUTY"]  # ... keeping its bookmark
    assert "(see section 4.2)" in text  # the reference shows the source's wording until Phase 12
    assert "This SOP is applicable:" in text
    for instruction in ("Follow the blue text", "Describe the process steps", "Brief description of what",
                        "Where possible identify", "Infographics"):
        assert instruction not in text


def test_no_blue_text_left_and_headers_untouched(tpl, tmp_path):
    report, doc, out = _render(tpl, tmp_path)
    colour = TextColour(doc)
    leftover = [p for p in _paragraphs(doc) if colour.classify(p._p) in ("blue", "mixed")
                and not p.style.name.startswith(("Heading", "Caption"))]  # the default styles' own blue
    assert leftover == [] and report.problems == []
    ids = [d.get("id") for d in doc.element.body.iter(qn("wp:docPr"))]
    assert len(ids) == len(set(ids))


def test_lists_use_real_word_numbering(tpl, tmp_path):
    _, doc, _ = _render(tpl, tmp_path)
    numbering = doc.part.numbering_part.element
    nums = {n.get(qn("w:numId")): n.find(qn("w:abstractNumId")).get(qn("w:val")) for n in numbering.findall(qn("w:num"))}
    formats = {a.get(qn("w:abstractNumId")): a.find(qn("w:lvl")).find(qn("w:numFmt")).get(qn("w:val"))
               for a in numbering.findall(qn("w:abstractNum"))}
    items = {}
    for p in _paragraphs(doc):
        num = p._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if num is not None and p.text in ("Open the record.", "Describe the event.", "including the batch", "Submit it."):
            items[p.text] = (num.find(qn("w:numId")).get(qn("w:val")), num.find(qn("w:ilvl")).get(qn("w:val")))
    assert set(items) == {"Open the record.", "Describe the event.", "including the batch", "Submit it."}
    steps = {items[t][0] for t in ("Open the record.", "Describe the event.", "Submit it.")}
    assert len(steps) == 1  # one list, so the bullet between steps does not restart the numbering
    assert formats[nums[steps.pop()]] == "decimal"
    assert formats[nums[items["including the batch"][0]]] == "bullet" and items["including the batch"][1] == "1"
    assert not any(p.text.startswith(("1.", "1\t")) for p in doc.paragraphs)  # no literal numbers


def test_subheading_figure_and_caption(tpl, tmp_path):
    report, doc, _ = _render(tpl, tmp_path)
    heading = next(p for p in _paragraphs(doc) if p.text == "Assessment")
    assert heading.style.name == "Heading 2"
    assert heading._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None  # the synthetic template's headings are unnumbered
    pictures = [p for p in _paragraphs(doc) if p._p.find(f".//{qn('w:drawing')}") is not None and not p.text]
    assert pictures, "the figure is in the document"
    assert any(p.text == "Figure 1: Assessment flow" and p.style.name == "Caption" for p in _paragraphs(doc))
    assert [p.text for p in _paragraphs(doc) if p.style.name == "Heading 1"] == [
        "PURPOSE", "APPLICABILITY", "DEFINITIONS & ABBREVIATIONS", "ROLES & RESPONSIBILITIES", "PROCESS", "OWNER"]


def test_missing_figure_file_is_a_visible_placeholder(tpl, tmp_path):
    source, drafts, plan = scenario(tmp_path)
    Path(tmp_path / "figure.png").unlink()
    report = render_document(tpl, drafts, source, tmp_path / "r.docx", job_id="J", slot_plan=plan)
    assert any("figure file not found" in w for w in report.warnings)
    assert report.problems == []
    assert "[Figure not available" in _texts(Document(str(tmp_path / "r.docx")))


# ── Callouts ──────────────────────────────────────────────────────────


def test_callout_assignment_clones_the_template_box(tpl, tmp_path):
    report, doc, _ = _render(tpl, tmp_path)
    assert report.slot("TGT-5-CONTENT").callout_boxes == 1
    boxes = [t for t in _tables(doc) if "Never close a deviation" in "".join(_cell_text(c) for c in t._cells)]
    assert len(boxes) == 1
    box = boxes[0]
    fills = {shd.get(qn("w:fill")).upper() for shd in box._tbl.iter(qn("w:shd"))}
    assert "F5CDB9" in fills  # the template's attention colour
    assert box._tbl.find(f".//{qn('w:drawing')}") is not None  # its icon
    assert _cell_text(box.rows[0].cells[-1]) == "Never close a deviation without QA approval."
    # The empty fixed attention box and the legend table are gone.
    assert "Attention section" not in _texts(doc) and "Infographics" not in _texts(doc)


# ── Tables ────────────────────────────────────────────────────────────


def test_table_slot_fills_the_template_table(tpl, tmp_path):
    _, doc, _ = _render(tpl, tmp_path)
    terms = next(t for t in _tables(doc) if _cell_text(t.rows[0].cells[0]) == "Term")
    assert [[_cell_text(c) for c in r.cells] for r in terms.rows] == [
        ["Term", "Definition"], ["CAPA", "Corrective and preventive action"], ["QA", "Quality Assurance"]]
    shd = terms.rows[2].cells[1]._tc.find(f"{qn('w:tcPr')}/{qn('w:shd')}")
    assert shd is not None and shd.get(qn("w:fill")) == "FFFF00"  # the source cell's shading


def test_a_table_that_does_not_fit_keeps_its_own_columns(tpl, tmp_path):
    report, doc, _ = _render(tpl, tmp_path)
    tables = [[[_cell_text(c) for c in r.cells] for r in t.rows] for t in _tables(doc)]
    assert [["Role", "Responsibility"], ["Process Owner", "Owns the process"]] in tables
    assert [["Role", "Task", "Deputy"], ["QA Head", "Approves", "QA Lead"]] in tables
    assert any("SRC-5-T2" in w and "own columns" in w for w in report.warnings)
    # The lead-in line stays before the tables, in claim order.
    body = _texts(doc)
    assert body.index("The following roles take part:") < body.index("Process Owner")


def test_multi_paragraph_source_cells_keep_their_lines(tpl, tmp_path):
    source, drafts, plan = scenario(tmp_path)
    cell = source.sections[2].units[0].table_ref.cells[1]
    cell.text, cell.paragraphs = "Corrective and preventive action", ["Corrective and", "preventive action"]
    render_document(tpl, drafts, source, tmp_path / "r.docx", job_id="J", slot_plan=plan)
    terms = next(t for t in _tables(Document(str(tmp_path / "r.docx"))) if _cell_text(t.rows[0].cells[0]) == "Term")
    assert [p.text for p in terms.rows[1].cells[1].paragraphs] == ["Corrective and", "preventive action"]


# ── Gaps, optional content, conditional regions ───────────────────────


def test_gap_marker_in_review_removed_in_final_and_blocks_until_accepted(tpl, tmp_path):
    report, doc, _ = _render(tpl, tmp_path, gap_geography=True)
    assert report.slot("TGT-2-GEOGRAPHY").outcome == SlotOutcome.GAP_MARKER
    assert _texts(doc).count(GAP_MARKER_TEXT) == 1 and report.problems == []

    source, drafts, plan = scenario(tmp_path, gap_geography=True)
    assert gap_slots(drafts) == ["TGT-2-GEOGRAPHY"]
    with pytest.raises(RenderBlocked) as blocked:
        render_document(tpl, drafts, source, tmp_path / "final.docx", job_id="J", slot_plan=plan, mode=RenderMode.FINAL)
    assert blocked.value.slot_ids == ["TGT-2-GEOGRAPHY"]

    final = render_document(tpl, drafts, source, tmp_path / "final.docx", job_id="J", slot_plan=plan,
                            mode=RenderMode.FINAL, accepted_gaps=["TGT-2-GEOGRAPHY"])
    assert final.slot("TGT-2-GEOGRAPHY").outcome == SlotOutcome.REMOVED_ACCEPTED_GAP and final.problems == []
    doc = Document(str(tmp_path / "final.docx"))
    assert GAP_MARKER_TEXT not in _texts(doc)
    assert not any("geography" in t.lower() for t in [_texts(doc)])  # the icon row and its instruction went
    assert not any((s.find(f"{qn('w:sdtPr')}/{qn('w:tag')}") is not None) for s in doc.element.body.iter(qn("w:sdt"))
                   if (s.find(f"{qn('w:sdtPr')}/{qn('w:tag')}").get(qn("w:val")) or "").startswith("CC_"))


def test_optional_empty_slot_is_removed_with_its_instruction(tpl, tmp_path):
    source, drafts, plan = scenario(tmp_path)
    drafts[1].slots[1].claims = []  # APPLICABILITY.not_covered (optional)
    report = render_document(tpl, drafts, source, tmp_path / "r.docx", job_id="J", slot_plan=plan)
    assert report.slot("TGT-2-NOT_COVERED").outcome == SlotOutcome.REMOVED_EMPTY
    assert "Where possible" not in _texts(Document(str(tmp_path / "r.docx")))


def test_conditional_regions(tpl, tmp_path):
    report, doc, _ = _render(tpl, tmp_path)
    regions = {r.region_id: r for r in report.regions}
    assert regions["p:4"].outcome == RegionOutcome.CHOICE_APPLIED and regions["p:4"].choice == "SOP"
    assert regions["t:4"].outcome == RegionOutcome.REMOVED  # the competence table: no decision to keep it
    assert "Competence" not in _texts(doc)

    source, drafts, plan = scenario(tmp_path)
    plan.sections[0].region_choices.append(RegionChoice(region_id="t:4", unit_ids=[source.sections[3].units[0].unit_id]))
    kept = render_document(tpl, drafts, source, tmp_path / "kept.docx", job_id="J", slot_plan=plan)
    assert {r.region_id: r.outcome for r in kept.regions}["t:4"] == RegionOutcome.KEPT
    doc = Document(str(tmp_path / "kept.docx"))
    competence = next(t for t in _tables(doc) if _cell_text(t.rows[0].cells[0]) == "Competence")
    assert TextColour(doc).classify(competence._tbl) == "plain" and kept.problems == []

    plan = SlotPlan(job_id="J", version=1)  # no choice at all: the inline region goes, with a note
    none = render_document(tpl, drafts, source, tmp_path / "none.docx", job_id="J", slot_plan=plan)
    assert {r.region_id: r.outcome for r in none.regions}["p:4"] == RegionOutcome.REMOVED
    assert "is applicable" not in _texts(Document(str(tmp_path / "none.docx")))


def test_unnormalized_anchors_table_cell_and_placeholder(tmp_path):
    """Before normalization the slots point at table cells and a placeholder; the renderer handles those too."""
    source_file = _template(tmp_path / "template.docx", tmp_path)
    model = build_template_model(source_file, "tpl").model
    model.normalized_file = str(source_file)
    source, drafts, plan = scenario(tmp_path)
    owner = model.slot_by_ref("OWNER.content")
    drafts[5].slots[0].slot_id = owner.slot_id
    report = render_document(model, drafts, source, tmp_path / "r.docx", job_id="J", slot_plan=plan)
    assert report.slot("TGT-1-WHAT").outcome == SlotOutcome.FILLED
    assert report.slot(owner.slot_id).outcome == SlotOutcome.FILLED
    doc = Document(str(tmp_path / "r.docx"))
    assert _cell_text(_tables(doc)[0].rows[0].cells[1]) == "This SOP describes how deviations are recorded."
    assert "Head of Quality" in _texts(doc) and "[Insert owner name]" not in _texts(doc)


# ── Post-render check ─────────────────────────────────────────────────


def test_post_render_check_reports_a_missing_claim(tpl, tmp_path):
    report, _, out = _render(tpl, tmp_path)
    source, drafts, plan = scenario(tmp_path)
    drafts[0].slots[0].claims.append(Claim(claim_id="C-1-099", text="A sentence that was never rendered.",
                                           source_unit_ids=[source.sections[0].units[0].unit_id]))
    problems = verify_rendered(out, Path(tpl.normalized_file), tpl, drafts, source, report,
                               lambda target, claim: target)
    assert any("C-1-099" in p for p in problems)

    doc = Document(str(out))  # a figure that went missing is caught by its image, not by a picture count
    figure = next(p for p in doc.element.body.iter(qn("w:p")) if p.find(f".//{qn('w:drawing')}") is not None
                  and p.find(f"{qn('w:pPr')}/{qn('w:jc')}") is not None)
    figure.getparent().remove(figure)
    doc.save(str(out))
    problems = verify_rendered(out, Path(tpl.normalized_file), tpl, scenario(tmp_path)[1], source, report,
                               lambda target, claim: target)
    assert any("figures are not in the document" in p and "C-5-009" in p for p in problems)


# ── Stage and API ─────────────────────────────────────────────────────


async def test_assembling_and_rendering_stages_write_the_model_document_and_report(tpl, tmp_path):
    from app.config.settings import Settings
    from app.schemas.v2 import MigrationJob
    from app.services.migration_v2.artifacts import ArtifactWriter
    from app.services.migration_v2.stages import StageContext, assembling, document_rendered, rendering
    from app.stores.migration_store import MigrationStore

    settings = Settings(project_root=tmp_path)
    settings.resolve_paths(tmp_path)
    settings.ensure_directories()
    store = MigrationStore(tmp_path / "m.db", settings)
    writer = ArtifactWriter(tmp_path / "artifacts", store)
    job = store.create_job(MigrationJob(job_id="MIG-R", sop_record_id=1, template_id="tpl", template_version=1))
    source, drafts, plan = scenario(tmp_path)
    writer.write(job.job_id, ArtifactKind.SOURCE_MODEL, source)
    writer.write(job.job_id, ArtifactKind.TEMPLATE_MODEL, tpl)
    writer.write(job.job_id, ArtifactKind.SLOT_PLAN, plan)
    for d in drafts:
        writer.write(job.job_id, ArtifactKind.SECTION_DRAFT, d, scope=d.target_section_id)
    ctx = StageContext(job=store.get_job(job.job_id), store=store, artifacts=writer, inputs=None, settings=settings)
    await assembling(ctx)
    ctx.job = store.get_job(job.job_id)
    assert ctx.job.latest_artifact(ArtifactKind.NUMBER_MAP) and ctx.job.latest_artifact(ArtifactKind.ASSEMBLED_DRAFT)
    assert next(e for e in store.list_events(job.job_id) if e["event"] == "assembly")["detail"]["sections_removed"] == []
    await rendering(ctx)
    job = store.get_job(job.job_id)
    assert len([a for a in job.artifacts if a.kind == ArtifactKind.ASSEMBLED_DRAFT]) == 1  # unchanged drafts: reused
    ref = job.latest_artifact(ArtifactKind.DOCX)
    assert ref is not None and Path(ref.path).name == "docx_v1.docx"
    assert _tables(Document(io.BytesIO(writer.read_bytes(ref))))
    assert document_rendered(job, writer)
    event = next(e for e in store.list_events(job.job_id) if e["event"] == "render")
    assert event["detail"]["slots"]["filled"] >= 9 and event["detail"]["problems"] == []
    assert not list((tmp_path / "artifacts" / job.job_id).glob(".render-*"))  # the work file is gone

    tpl_missing = tpl.model_copy(update={"normalized_file": str(tmp_path / "nope.docx")})
    writer.write(job.job_id, ArtifactKind.TEMPLATE_MODEL, tpl_missing)
    ctx.job = store.get_job(job.job_id)
    await rendering(ctx)
    assert any(e["event"] == "render_failed" for e in store.list_events(job.job_id))


@pytest.fixture
def api_env(tmp_path):
    from tests.test_v2_migration_jobs import Env

    return Env(tmp_path)


def test_document_endpoints(api_env, tpl, tmp_path):
    from fastapi.testclient import TestClient

    from app.api.auth import require_user
    from app.config.settings import get_settings
    from app.main import app
    from app.schemas.v2 import MigrationJob

    env = api_env
    app.dependency_overrides[get_settings] = lambda: env.settings
    app.dependency_overrides[require_user] = lambda: {"userId": "reviewer-1"}
    try:
        with TestClient(app) as client:
            previous = app.state.migration_orchestrator
            orch = env.orchestrator()
            app.state.migration_orchestrator = orch
            try:
                job = env.store.create_job(MigrationJob(job_id="MIG-D", sop_record_id=1, template_id="TPL", template_version=1))
                assert client.get(f"/api/v1/migrations/{job.job_id}/document").status_code == 404
                source, drafts, plan = scenario(tmp_path)
                out = tmp_path / "doc.docx"
                report = render_document(tpl, drafts, source, out, job_id=job.job_id, slot_plan=plan)
                orch.artifacts.write_file(job.job_id, ArtifactKind.DOCX, out.read_bytes(), ".docx")
                orch.artifacts.write(job.job_id, ArtifactKind.RENDER_REPORT, report)

                response = client.get(f"/api/v1/migrations/{job.job_id}/document")
                assert response.status_code == 200
                assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
                assert "MIG-D_review_v1.docx" in response.headers["content-disposition"]
                assert _tables(Document(io.BytesIO(response.content)))
                assert client.get(f"/api/v1/migrations/{job.job_id}/artifacts/docx").content == response.content
                body = client.get(f"/api/v1/migrations/{job.job_id}/document/report").json()
                assert body["ok"] is True and body["mode"] == "review"
            finally:
                client.portal.call(orch.stop)
                app.state.migration_orchestrator = previous
    finally:
        app.dependency_overrides.pop(get_settings, None)
        app.dependency_overrides.pop(require_user, None)


# ── The three samples, placement mode ─────────────────────────────────


SOP_DIR = ROOT / "documents" / "SOPs"
TEMPLATE = ROOT / "documents" / "Templates" / "Template Main GP Docs.docx"
TEMPLATE_CONFIG = ROOT / "documents" / "template_config" / "Template_Main_GP_Docs.json"


async def test_samples_render_without_problems(tmp_path):
    if not (SOP_DIR.exists() and TEMPLATE.exists()):
        pytest.skip("sample SOPs and template not available")
    from app.services.migration_v2.inspection.runner import inspect_sop, prepare_template

    template = prepare_template(TEMPLATE, TEMPLATE_CONFIG, tmp_path / "_template")
    for sop in sorted(SOP_DIR.glob("*.docx")):
        insp = await inspect_sop(sop, template, tmp_path / "reports")
        report = insp.render
        assert report is not None and report.problems == [], (sop.name, report and report.problems)
        assert insp.document is not None and insp.document.exists()
        outcomes = {s.slot_id: s.outcome for s in report.slots}
        assert SlotOutcome.NO_ANCHOR not in outcomes.values(), sop.name
        assert "TGT-9" in report.sections_removed  # the optional distribution chapter has no content
        doc = Document(str(insp.document))
        assert "controlled prints" not in _texts(doc)
        toc = [text for _, text, _ in _toc(doc)]  # the post-render check matched it to the headings
        assert report.toc_entries == len(toc) >= 9 and "PROCESS" in " ".join(toc), sop.name
        assert not any("controlled prints" in t or "TEMPLATE DOCUMENT HISTORY" in t for t in toc), sop.name
        assert insp.document.with_suffix("").name == "docx_v1"
        gaps = [s for s, o in outcomes.items() if o == SlotOutcome.GAP_MARKER]
        assert _texts(doc).count(GAP_MARKER_TEXT) == len(gaps)
        if "10505" in sop.name:
            assert {r.region_id: r.choice for r in report.regions}["p:26"] == "SOP"
            assert "This SOP is applicable:" in _texts(doc)
        if "00535" in sop.name:
            assert sorted(gaps) == ["TGT-2-GEOGRAPHY", "TGT-2-UNITS"]
        fields = [i.text for i in doc.element.body.iter(qn("w:instrText")) if (i.text or "").strip().startswith("REF ")]
        assert len(fields) == sum(bool(r.bookmark) for r in insp.assembled.refs), sop.name  # each number is a REF field
        if "00493" in sop.name:  # DEFINITIONS: the narrative and Image 1 go below the abbreviations table
            blocks = list(doc.element.body.iter(qn("w:tbl"), qn("w:p")))
            abbr = next(i for i, b in enumerate(blocks) if b.tag == qn("w:tbl") and "Animal Health" in _cell_text_of(b))
            caption = next(i for i, b in enumerate(blocks) if b.tag == qn("w:p") and
                           "Image 1: Dimensions of the GBS unITed Program" in "".join(t.text or "" for t in b.iter(qn("w:t"))))
            assert caption > abbr and next(blocks[caption - 1].iter(qn("a:blip")), None) is not None
    json.dumps(report.to_clean_dict())


def test_reference_text_does_not_repeat_the_words_before_the_token():
    from app.services.migration_v2.render.content import show_refs

    phrase = lambda target, claim: "described in 6.2"  # noqa: E731  (the source wording of the reference)
    rewritten = Claim(claim_id="C", text="Special processes are described in {{ref:SRC-6.2}}.", source_unit_ids=["U"])
    copied = Claim(claim_id="C", text="Special processes are {{ref:SRC-6.2}}.", source_unit_ids=["U"])
    assert show_refs(rewritten, phrase) == show_refs(copied, phrase) == "Special processes are described in 6.2."


def test_callout_box_stays_on_one_page(tpl, tmp_path):
    _, doc, _ = _render(tpl, tmp_path)
    box = next(t for t in _tables(doc) if "Never close a deviation" in "".join(_cell_text(c) for c in t._cells))
    assert all(tr.find(f"{qn('w:trPr')}/{qn('w:cantSplit')}") is not None for tr in box._tbl.findall(qn("w:tr")))


# ── Table of contents ─────────────────────────────────────────────────

_STALE_TOC = (
    '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:sdtPr><w:docPartObj>'
    '<w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent>'
    '<w:p><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="9000"/></w:tabs></w:pPr>'
    r'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \o "1-2" \h \z \u </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r><w:r><w:tab/></w:r>'
    '<w:r><w:t>OLD TEMPLATE CHAPTER</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>9</w:t></w:r></w:p>'
    '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p></w:sdtContent></w:sdt>'
)


def _with_toc(tpl):
    """The synthetic template with a stale table of contents (the template's own entries) before its first heading."""
    from docx.oxml import parse_xml

    doc = Document(tpl.normalized_file)
    first_heading = next(p for p in doc.paragraphs if p.style.name.startswith("Heading"))
    first_heading._p.addprevious(parse_xml(_STALE_TOC))
    doc.save(tpl.normalized_file)
    return tpl


def _toc(doc) -> list[tuple[str, str, str]]:
    """(anchor, text, page) of each TOC entry."""
    from app.services.parser.ooxml import is_toc_sdt

    sdt = next(s for s in doc.element.body.iter(qn("w:sdt")) if is_toc_sdt(s))
    out = []
    for link in sdt.iter(qn("w:hyperlink")):
        texts = [t.text or "" for t in link.iter(qn("w:t"))]
        out.append((link.get(qn("w:anchor")), "".join(texts[:-1]), texts[-1]))
    return out


def _headings(doc) -> list[str]:
    """Heading 1 and 2 texts in document order, inside content controls too (the TOC has none)."""
    from docx.text.paragraph import Paragraph

    paragraphs = [Paragraph(p, doc._body) for p in doc.element.body.iter(qn("w:p"))]
    return [p.text for p in paragraphs if p.style.name in ("Heading 1", "Heading 2") and p.text.strip()]


def test_toc_levels_from_the_field_switches():
    from app.services.migration_v2.render.toc import DEFAULT_LEVELS, toc_levels

    assert toc_levels(r' TOC \h \u \z \t "Heading 1,1,Heading 2,2,Heading 3,3," ') == \
        {"heading 1": 1, "heading 2": 2, "heading 3": 3}
    assert toc_levels(r' TOC \o "1-2" \h ') == {"heading 1": 1, "heading 2": 2}
    assert toc_levels(r" TOC \h ") == DEFAULT_LEVELS


def test_toc_lists_the_rendered_headings(tpl, tmp_path):
    report, doc, _ = _render(_with_toc(tpl), tmp_path)
    assert report.problems == [], report.problems
    entries = _toc(doc)
    assert [text for _, text, _ in entries] == _headings(doc)  # the template's stale entry is gone
    assert "Assessment" in _headings(doc)  # a source sub-heading made it into the TOC
    marks = {m.get(qn("w:name")) for m in doc.element.body.iter(qn("w:bookmarkStart"))}
    assert all(anchor in marks for anchor, _, _ in entries)
    assert report.toc_entries == len(entries) and not report.toc_page_numbers
    assert all(page == "" for _, _, page in entries)  # no layout engine: Word fills them on opening ...
    assert doc.settings.element.find(qn("w:updateFields")) is not None  # ... when it updates the fields
    instr = "".join(i.text for i in doc.element.body.iter(qn("w:instrText")) if "TOC" in i.text)
    assert instr == r' TOC \o "1-2" \h \z \u '  # the field stays, so Word can still refresh it


def test_toc_page_numbers_from_the_page_counter(tpl, tmp_path):
    calls = []

    def counter(path, bookmarks):
        calls.append((Path(path).exists(), list(bookmarks)))
        return {b: i + 2 for i, b in enumerate(bookmarks)}

    report, doc, _ = _render(_with_toc(tpl), tmp_path, page_counter=counter)
    entries = _toc(doc)
    assert calls == [(True, [anchor for anchor, _, _ in entries])]
    assert [page for _, _, page in entries] == [str(i + 2) for i in range(len(entries))]
    assert report.toc_page_numbers and report.problems == []
    assert doc.settings.element.find(qn("w:updateFields")) is None  # nothing left for Word to update


def test_toc_without_word_keeps_asking_word_to_update(tpl, tmp_path):
    report, doc, _ = _render(_with_toc(tpl), tmp_path, page_counter=lambda path, bookmarks: None)
    assert not report.toc_page_numbers and any("table of contents" in w for w in report.warnings)
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_post_render_check_reports_a_stale_toc(tpl, tmp_path):
    from app.services.migration_v2.render.toc import toc_problems

    _, doc, _ = _render(_with_toc(tpl), tmp_path)
    assert toc_problems(doc) == []
    link = next(doc.element.body.iter(qn("w:hyperlink")))
    next(link.iter(qn("w:t"))).text = "OLD TEMPLATE CHAPTER"
    link.set(qn("w:anchor"), "_Toc_missing")
    problems = toc_problems(doc)
    assert any("missing bookmark '_Toc_missing'" in p for p in problems)
    assert any("does not match the headings" in p for p in problems)
