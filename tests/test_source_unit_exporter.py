"""SourceUnitExporter: AST → v2 SourceDocument (Phase 2b), on synthetic documents."""

import json
import re
from collections import Counter
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches
from PIL import Image

from app.config.settings import Settings
from app.schemas.v2 import AssetKind, RefKind, RefResolution, RelationType, SourceDocument, UnitType
from app.services.export.source_refs import find_reference_phrases
from app.services.export.source_unit_exporter import SourceUnitExporter
from app.services.extraction.captions import CaptionExtractor
from app.services.extraction.cross_page_stitcher import CrossPageTableStitcher
from app.services.extraction.icons import IconExtractor
from app.services.extraction.tables import TableExtractor
from app.services.hierarchy.ast_builder import ASTBuilder
from app.services.parser.docx_parser import DocxParser
from app.services.parser.ooxml import element_text, iter_body_blocks


@pytest.fixture
def settings(tmp_path):
    s = Settings(project_root=tmp_path)
    s.resolve_paths(tmp_path)
    s.ensure_directories()
    return s


def _png(path: Path, color) -> Path:
    Image.new("RGB", (64, 64), color).save(path)
    return path


def _shade(cell, fill):
    cell._element.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>')
    )


def _build(path: Path, tmp: Path) -> Path:
    icon = _png(tmp / "icon.png", (200, 40, 40))
    icon2 = _png(tmp / "icon2.png", (40, 160, 60))
    figure = _png(tmp / "figure.png", (30, 90, 200))
    doc = Document()

    cover = doc.add_table(rows=2, cols=2)
    cover.rows[0].cells[0].text, cover.rows[0].cells[1].text = "Title:", "Sample SOP"
    cover.rows[1].cells[0].text, cover.rows[1].cells[1].text = "Version:", "1.0"

    doc.add_heading("1 PURPOSE", level=1)
    purpose = doc.add_table(rows=1, cols=2)
    purpose.rows[0].cells[0].paragraphs[0].add_run().add_picture(str(icon), width=Inches(0.5))
    purpose.rows[0].cells[1].paragraphs[0].text = "This SOP describes sample handling."
    purpose.rows[0].cells[1].add_paragraph("Its intention is consistent handling.")

    doc.add_heading("2 PROCESS", level=1)
    doc.add_paragraph("The process has two steps.")
    doc.add_heading("2.1 Steps", level=2)
    doc.add_paragraph("Warning: Never skip the second check.")
    for text in ("1. Create the record.", "2. Submit the record within five business days."):
        doc.add_paragraph(text, style="List Number")
    doc.add_paragraph("Note: The record is kept for ten years.")
    doc.add_paragraph("Complete the form as described in chapter 2.1 and refer to step 2.")
    doc.add_paragraph("The validated system is defined in the SOP (see Chapter 3, no. 2), as described above.")
    doc.add_paragraph("This follows BI-VQD-12345-S and chapter 9.9.")
    marker = doc.add_paragraph()
    marker.add_run().add_picture(str(icon2), width=Inches(0.4))
    marker.add_run("Remember to sign.")
    doc.add_paragraph().add_run().add_picture(str(figure), width=Inches(4))
    doc.add_paragraph("Figure 1: Overview of the process")

    data = doc.add_table(rows=3, cols=2)
    for col, header in enumerate(("Level", "Testing")):
        data.rows[0].cells[col].text = header
        _shade(data.rows[0].cells[col], "D9D9D9")
    data.rows[1].cells[0].text, data.rows[1].cells[1].text = "1", "Full test"
    data.rows[2].cells[0].text, data.rows[2].cells[1].text = "2", "Reduced test"
    _shade(data.rows[1].cells[1], "FF0000")

    doc.add_heading("3 REFERENCES", level=1)
    refs = doc.add_table(rows=3, cols=2)
    for r, (no, title) in enumerate((("No.", "Title"), ("1", "Quality Manual"), ("2", "Validation SOP"))):
        refs.rows[r].cells[0].text, refs.rows[r].cells[1].text = no, title
    for cell in refs.rows[0].cells:
        cell.paragraphs[0].runs[0].bold = True

    doc.add_heading("4 DEFINITIONS", level=1)
    terms = doc.add_table(rows=2, cols=2)
    terms.rows[0].cells[0].text, terms.rows[0].cells[1].text = "Term", "Definition"
    for cell in terms.rows[0].cells:
        cell.paragraphs[0].runs[0].bold = True
    terms.rows[1].cells[0].text, terms.rows[1].cells[1].text = "SOP", "Standard Operating Procedure"

    doc.add_heading("5 DOCUMENT HISTORY", level=1)
    doc.add_paragraph("Version 1.0: first issue.")
    doc.save(path)
    return path


def _export(settings, path) -> SourceDocument:
    raw = DocxParser(settings=settings).parse(str(path), document_id="unit_test")
    raw = TableExtractor(settings=settings).extract(raw)
    raw = CrossPageTableStitcher(settings=settings).extract(raw)
    raw = IconExtractor(settings=settings).extract(raw, document_id="unit_test")
    raw = CaptionExtractor(settings=settings).extract(raw)
    return SourceUnitExporter.export("unit_test", ASTBuilder(settings=settings).build(raw), source_file=str(path))


@pytest.fixture
def doc(settings, tmp_path):
    return _export(settings, _build(tmp_path / "sample.docx", tmp_path))


def _units(doc, **match):
    return [u for u in doc.iter_units() if all(getattr(u, k) == v for k, v in match.items())]


def _by_text(doc, fragment):
    return next(u for u in doc.iter_units() if fragment in u.text)


# ── Sections ───────────────────────────────────────────────────────────

def test_sections_are_nested_and_ordered(doc):
    ids = [s.section_id for s in doc.sections]
    assert ids == ["SRC-0", "SRC-1", "SRC-2", "SRC-2.1", "SRC-3", "SRC-4", "SRC-5"]
    by_id = {s.section_id: s for s in doc.sections}
    assert by_id["SRC-2.1"].parent_id == "SRC-2"
    assert by_id["SRC-2.1"].heading == "Steps"  # number kept separately
    assert by_id["SRC-2.1"].number == "2.1"


def test_preamble_and_history_are_boilerplate(doc):
    assert all(u.is_boilerplate for u in _units(doc, section_id="SRC-0"))
    assert all(u.is_boilerplate for u in _units(doc, section_id="SRC-5"))
    assert not any(u.is_boilerplate for u in _units(doc, section_id="SRC-2"))


def test_reading_order_ids_and_hashes(settings, tmp_path, doc):
    seqs = [u.seq for u in doc.iter_units()]
    assert seqs == sorted(seqs) and len(set(seqs)) == len(seqs)
    again = _export(settings, tmp_path / "sample.docx")
    assert [u.content_hash for u in again.iter_units()] == [u.content_hash for u in doc.iter_units()]
    assert [u.unit_id for u in again.iter_units()] == [u.unit_id for u in doc.iter_units()]


# ── Unit types ─────────────────────────────────────────────────────────

def test_warning_and_note_are_typed_and_linked(doc):
    warning = _by_text(doc, "Never skip")
    assert warning.unit_type == UnitType.WARNING
    note = _by_text(doc, "kept for ten years")
    assert note.unit_type == UnitType.NOTE
    step1 = _by_text(doc, "Create the record")
    step2 = _by_text(doc, "within five business days")
    assert note.relations[0].type == RelationType.NOTE_FOR
    assert note.relations[0].target_unit_id == step2.unit_id
    # First unit of its section: it qualifies what follows.
    assert warning.relations[0].type == RelationType.WARNING_FOR
    assert warning.relations[0].target_unit_id == step1.unit_id


def test_numbered_list_items_are_procedure_steps(doc):
    steps = _units(doc, section_id="SRC-2.1", unit_type=UnitType.PROCEDURE_STEP)
    assert [s.text for s in steps] == ["Create the record.", "Submit the record within five business days."]
    assert [s.list_number for s in steps] == ["1", "2"]


def test_data_table_rows_headers_and_fills(doc):
    rows = [u for u in _units(doc, section_id="SRC-2.1") if u.unit_type == UnitType.TABLE_ROW]
    assert [r.text for r in rows] == ["1 | Full test", "2 | Reduced test"]
    assert rows[0].table_ref.header_cells == ["Level", "Testing"]
    fills = {c.text: c.fill_hex for c in rows[0].table_ref.cells}
    assert fills == {"1": None, "Full test": "FF0000"}  # grey header fill is dropped, red kept


def test_definition_and_reference_rows(doc):
    assert [u.text for u in _units(doc, unit_type=UnitType.DEFINITION)] == ["SOP | Standard Operating Procedure"]
    assert [u.text for u in _units(doc, unit_type=UnitType.REFERENCE)] == ["1 | Quality Manual", "2 | Validation SOP"]


# ── Icons and figures ──────────────────────────────────────────────────

def test_icon_table_becomes_paragraphs_with_icon_metadata(doc):
    units = _units(doc, section_id="SRC-1")
    assert [u.text for u in units] == ["This SOP describes sample handling.", "Its intention is consistent handling."]
    assert all(u.unit_type == UnitType.PARAGRAPH for u in units)
    assert [a.kind for a in units[0].assets] == [AssetKind.ICON]
    assert units[1].assets == []
    assert units[0].table_ref.row_index == units[1].table_ref.row_index == 0


def test_inline_icon_attaches_to_its_text(doc):
    unit = _by_text(doc, "Remember to sign.")
    assert [a.kind for a in unit.assets] == [AssetKind.ICON]
    assert unit.assets[0].asset_key.startswith("icon_")


def test_figure_stays_a_figure_with_its_caption(doc):
    (figure,) = _units(doc, unit_type=UnitType.FIGURE)
    assert [a.kind for a in figure.assets] == [AssetKind.FIGURE]
    assert figure.assets[0].width_px > 96
    caption = _by_text(doc, "Overview of the process")
    assert caption.unit_type == UnitType.CAPTION
    assert caption.relations[0].type == RelationType.CAPTION_OF
    assert caption.relations[0].target_unit_id == figure.unit_id
    assert all(a.kind == AssetKind.ICON for u in doc.iter_units() if u.unit_type != UnitType.FIGURE for a in u.assets)


# ── Cross-references ───────────────────────────────────────────────────

def test_cross_references_resolve(doc):
    refs = {r.raw_text.lower(): r for r in doc.cross_references}
    assert refs["chapter 2.1"].target == "SRC-2.1"
    assert refs["chapter 2.1"].resolution == RefResolution.RESOLVED
    step = refs["refer to step 2"]
    assert step.ref_kind == RefKind.STEP
    assert step.target == _by_text(doc, "within five business days").unit_id
    item = refs["chapter 3, no. 2"]
    assert item.target == _by_text(doc, "Validation SOP").unit_id
    assert refs["as described above"].resolution == RefResolution.UNRESOLVED
    assert refs["as described above"].direction == "before"
    assert refs["bi-vqd-12345-s"].resolution == RefResolution.EXTERNAL
    assert refs["chapter 9.9"].resolution == RefResolution.UNRESOLVED


def test_reference_phrases_ignore_plain_numbers():
    assert find_reference_phrases("In 2024 we had 3.5 deviations per site.") == []
    assert [m.raw for m in find_reference_phrases("Special processes are described in 6.2")] == ["described in 6.2"]


# ── No lost text ───────────────────────────────────────────────────────

def test_every_word_of_the_document_is_in_a_unit(tmp_path, doc):
    body = Document(str(tmp_path / "sample.docx")).element.body
    source = Counter(" ".join(element_text(b) for b in iter_body_blocks(body)).lower().split())
    output = Counter(json.dumps(doc.to_clean_dict(), ensure_ascii=False).lower().replace('"', " ").split())
    # A typed "1." before a list item moves to list_number ("1"), so trailing dots are ignored.
    strip = lambda c: Counter({re.sub(r"[^\w.-]", "", k).rstrip("."): v for k, v in c.items()})
    missing = strip(source) - strip(output)
    del missing[""]
    assert sum(missing.values()) == 0, missing
