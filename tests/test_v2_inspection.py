"""Inspection tool: checks, HTML report and index (scripts/inspect_sop.py)."""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from docx import Document

from app.services.migration_v2.inspection.checks import FAIL, PASS, overall
from app.services.migration_v2.inspection.completeness import check_completeness, words
from app.services.migration_v2.inspection.report import highlight, render, render_index
from app.services.migration_v2.inspection.runner import inspect_sop, prepare_template
from tests.test_v2_template_slots import _template

ROOT = Path(__file__).resolve().parent.parent


def _sop_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("PURPOSE", level=1)
    doc.add_paragraph("This SOP describes how deviations are handled.")
    doc.add_heading("APPLICABILITY", level=1)
    doc.add_paragraph("This SOP applies to all sites world-wide.")
    doc.add_heading("DEFINITIONS & ABBREVIATIONS", level=1)
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Abbreviation", "Description"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "QA", "Quality Assurance"
    doc.add_heading("ROLES & RESPONSIBILITIES", level=1)
    doc.add_paragraph("The Process Owner is responsible for the process.")
    doc.add_heading("PROCESS", level=1)
    doc.add_paragraph("QA must approve the deviation within 5 business days.", style="List Number")
    doc.add_paragraph("Do not close the deviation before approval.", style="List Number")
    doc.add_paragraph("Note <script>alert(1)</script> & keep the form BI-VQD-10095-S.")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


@pytest.fixture
def inspected(tmp_path):
    (tmp_path / "tpl").mkdir()
    template = prepare_template(_template(tmp_path / "tpl" / "template.docx", tmp_path / "tpl"), None, tmp_path / "work_tpl")
    sop = _sop_docx(tmp_path / "in" / "Deviation SOP.docx")
    return asyncio.run(inspect_sop(sop, template, tmp_path / "out")), tmp_path


def test_inspection_runs_the_stages_and_checks(inspected):
    insp, tmp = inspected
    checks = {c.name: c for c in insp.checks}
    assert checks["Text completeness"].status == PASS and insp.completeness.recall == 1.0
    assert checks["Protected facts"].status == PASS  # the verbatim self-check is clean
    assert {v.normalized for v in insp.facts.values} >= {"<=5 business_day", "BI-VQD-10095-S"}
    # The synthetic template has undecided regions: a real job would refuse it, and the report says so.
    assert checks["Template"].status == FAIL and any("decide" in d for d in checks["Template"].details)
    assert checks["Golden comparison"].status == "info"
    assert overall(insp.checks) == FAIL
    keys = {t.section_id: t.key for t in insp.template.model.sections}
    placed = {keys[m.target_section_id]: m.source_section_ids for m in insp.plan.mappings if m.target_section_id}
    assert all(placed[k] for k in ("PURPOSE", "APPLICABILITY", "DEFINITIONS", "ROLES", "PROCESS"))
    # Artifacts are copied next to the report.
    assert {"source_model_v1.json", "section_plan_v1.json", "protected_facts_v1.json"} <= set(insp.artifacts)
    assert all(p.exists() for p in insp.artifacts.values())


def test_report_escapes_and_highlights(inspected):
    insp, _ = inspected
    html = render(insp)
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert '<span class="m-mandatory">must</span>' in html
    assert '<mark class="v">BI-VQD-10095-S</mark>' in html
    for anchor in ('id="checks"', 'id="plan"', 'id="outline"', 'id="facts"', 'id="template"', 'id="golden"'):
        assert anchor in html
    index = render_index([("Deviation SOP.docx", "x/report.html", insp.checks)], [("broken.pdf", "ValueError: bad file")])
    assert "Deviation SOP.docx" in index and "broken.pdf" in index and "ValueError: bad file" in index


def test_completeness_reports_lost_text(inspected):
    insp, tmp = inspected
    source = insp.source
    trimmed = source.model_copy(update={"sections": [
        s.model_copy(update={"units": [u for u in s.units if "world-wide" not in u.text]}) for s in source.sections
    ]})
    result = check_completeness(insp.sop_path, trimmed)
    assert result.recall < 1.0 and "world-wide" in result.missing_words
    assert any("world-wide" in ctx for ctx in result.missing_context)


def test_highlight_handles_overlaps_and_escaping():
    html = highlight("QA must not close <b> within 5 days.", values=["within 5 days", "5"], roles=["QA"])
    assert html.startswith('<mark class="r">QA</mark> <span class="m-prohibition">must not</span>')
    assert "&lt;b&gt;" in html and '<mark class="v">within 5 days</mark>' in html
    assert words("Hello, World! • — x") == ["hello", "world", "x"]


# ── Real samples ──────────────────────────────────────────────────────

SOPS = ROOT / "documents" / "SOPs"
TEMPLATE = ROOT / "documents" / "Templates" / "Template Main GP Docs.docx"
CONFIG = ROOT / "documents" / "template_config" / "Template_Main_GP_Docs.json"


@pytest.mark.skipif(not (SOPS.exists() and TEMPLATE.exists()), reason="sample SOPs not available")
def test_samples_pass_inspection(tmp_path):
    template = prepare_template(TEMPLATE, CONFIG, tmp_path / "tpl")
    for sop in sorted(SOPS.glob("*.docx")):
        insp = asyncio.run(inspect_sop(sop, template, tmp_path / "out", golden_dir=ROOT / "documents" / "golden"))
        statuses = {c.name: c.status for c in insp.checks}
        assert statuses["Text completeness"] == PASS, sop.name
        assert statuses["Template"] == PASS and statuses["Protected facts"] == PASS, sop.name
        assert statuses["Golden comparison"] == PASS, (sop.name, insp.golden_problems, insp.missing_preserve)
        assert statuses["Section plan"] in (PASS, "warn"), sop.name
        assert "<html" in render(insp)
