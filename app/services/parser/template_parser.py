"""Concrete DOCX parser for Template documents with blue-font instruction detection.

Extends DocxParser to identify blue-colored instructions, extract embedded icons/images,
and preserve section/paragraph structures for template management.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Optional

from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.config.settings import Settings
from app.schemas.document import (
    DocumentMetadata,
    ElementType,
    ExtractedElement,
    ExtractedHeading,
    ExtractedImage,
    RawDocument,
)
from .docx_parser import DocxParser


class TemplateDocxParser(DocxParser):
    """Parse template .docx files into a structured RawDocument with instruction annotations."""

    # Numbered blue instructions are instructions, not content lists.
    _NUMPR_LISTS = False

    # 23 known blue hex codes used across templates and SOPs
    BLUE_HEX_VALUES = {
        "0000FF", "0070C0", "4472C4", "2E74B5", "5B9BD5", "00B0F0",
        "1F4E78", "2F5597", "41719C", "002060", "1B365D", "0072CE",
        "0075FF", "00A2E8", "004B87", "0080FF", "0099FF", "0066CC",
        "1B75BB", "005696", "3399FF", "003366", "2E75B6",
    }

    def __init__(self, settings: Settings, logger=None):
        super().__init__(settings=settings, logger=logger)
        self.template_id: Optional[str] = None

    def parse(self, file_path: str, template_id: str | None = None) -> RawDocument:
        """Open template *file_path*, extract elements with instruction tags, return RawDocument."""
        self._validate_file(file_path)
        self.logger.info(f"TemplateDocxParser: opening {file_path}")
        self.template_id = template_id
        self.current_document_id = template_id

        doc = DocxDocument(file_path)

        metadata = self._extract_metadata(doc, file_path, document_id=template_id)
        elements: list[ExtractedElement] = []
        sequence = 0
        self._current_page = 1

        # --- Iterate body in document order (paragraphs + tables interleaved) ---
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn

        extracted_image_hashes: set[str] = set()  # track inline-extracted images to avoid duplicates
        self._extracted_hash_paths: dict[str, Path] = {}

        for block in doc.element.body:
            tag = block.tag
            if tag == qn('w:p'):
                para = Paragraph(block, doc)
                if self._paragraph_starts_new_page(para):
                    self._current_page += 1

                element = self._classify_paragraph(para, sequence)
                if element is not None:
                    elements.append(element)
                    sequence += 1

                # Extract inline images from this paragraph in document order
                inline_images = self._extract_inline_images(block, doc, sequence, extracted_image_hashes)
                for img_el in inline_images:
                    elements.append(img_el)
                    sequence += 1

            elif tag == qn('w:tbl'):
                table = DocxTable(block, doc)
                table_element = self._extract_table(table, sequence, doc=doc, extracted_hashes=extracted_image_hashes)
                if table_element is not None:
                    elements.append(table_element)
                    sequence += 1

        # --- Images fallback (relationship-based, catches images not found inline) ---
        image_elements = self._extract_images(doc, sequence, skip_hashes=extracted_image_hashes)
        elements.extend(image_elements)

        # Group elements into per-page PageContent objects
        pages = self._group_elements_into_pages(elements)

        self.logger.info(
            f"TemplateDocxParser: extracted {len(elements)} element(s) across "
            f"{len(pages)} page(s) from {file_path}"
        )

        return RawDocument(
            source=file_path,
            metadata=metadata,
            pages=pages,
        )

    def _extract_metadata(
        self, doc: DocxDocument, file_path: str, document_id: str | None = None
    ) -> DocumentMetadata:
        """Extract simplified template metadata from docx properties."""
        props = doc.core_properties
        file_stat = Path(file_path).stat()
        tpl_name = Path(file_path).stem

        return DocumentMetadata(
            title=props.title or tpl_name,
            author=props.author or "",
            subject=props.subject or "",
            creator=props.last_modified_by or "",
            creation_date=props.created.isoformat() if props.created else "",
            modification_date=props.modified.isoformat() if props.modified else "",
            page_count=0,
            file_type="docx",
            file_size_bytes=file_stat.st_size,
            document_title=props.title or tpl_name,
            document_name=tpl_name,
            document_number="",
            document_version="",
            document_type="template",
            gpdat_version=1,
        )

    def _classify_paragraph(self, para, sequence: int) -> Optional[ExtractedElement]:
        """Classify paragraph, tagging blue instruction paragraphs."""
        base_element = super()._classify_paragraph(para, sequence)
        if base_element is None:
            return None

        # Check for blue font color indicating template instruction
        is_blue, color_hex, detection_method = self._detect_color(para)
        if is_blue:
            # Extract specifically blue run text
            blue_parts: list[str] = []
            for run in para.runs:
                run_blue, _, _ = self._detect_run_color(run)
                if run_blue and run.text.strip():
                    blue_parts.append(run.text.strip())

            instruction_text = " ".join(blue_parts) if blue_parts else para.text.strip()

            base_element.font_color_hex = color_hex
            base_element.color_detection_method = detection_method
            if not hasattr(base_element, "metadata") or base_element.metadata is None:
                base_element.metadata = {}
            base_element.metadata["is_instruction"] = True
            base_element.metadata["instruction_text"] = instruction_text
            base_element.metadata["font_color_hex"] = color_hex
            base_element.metadata["color_detection_method"] = detection_method

        return base_element

    def _detect_color(self, para) -> tuple[bool, Optional[str], Optional[str]]:
        """Detect blue instruction text on a paragraph.

        Returns ``(is_blue, font_color_hex, detection_method)``. The hex is only
        populated when an actual colour value was found — style- and theme-based
        matches report how they were detected instead.
        """
        if para.style and para.style.name:
            st_name = para.style.name.lower()
            if "instruction" in st_name:
                return True, None, "style_name"

        WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

        # Check individual runs
        for run in para.runs:
            run_is_blue, run_color, run_method = self._detect_run_color(run)
            if run_is_blue:
                return True, run_color, run_method

        # Check paragraph-level rPr
        p_pr = para._element.find(f"{{{WNS}}}pPr")
        if p_pr is not None:
            r_pr = p_pr.find(f"{{{WNS}}}rPr")
            if r_pr is not None:
                color_el = r_pr.find(f"{{{WNS}}}color")
                if color_el is not None:
                    val = (
                        color_el.get(f"{{{WNS}}}val", "")
                        or color_el.get("w:val", "")
                        or color_el.get("val", "")
                    ).upper()
                    if self._check_hex_or_rgb(val):
                        return True, val, "paragraph_color"

        return False, None, None

    def _detect_run_color(self, run) -> tuple[bool, Optional[str], Optional[str]]:
        """Detect blue text on a single run.

        Returns ``(is_blue, font_color_hex, detection_method)``.
        """
        WNS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

        # 1. python-docx RGBColor
        if run.font and run.font.color and run.font.color.rgb:
            rgb_str = str(run.font.color.rgb).upper()
            if self._check_hex_or_rgb(rgb_str):
                return True, rgb_str, "run_color"

        # 2. OXML w:rPr/w:color
        r_pr = run._element.find(f"{{{WNS}}}rPr")
        if r_pr is not None:
            color_el = r_pr.find(f"{{{WNS}}}color")
            if color_el is not None:
                val = (
                    color_el.get(f"{{{WNS}}}val", "")
                    or color_el.get("w:val", "")
                    or color_el.get("val", "")
                ).upper()
                if self._check_hex_or_rgb(val):
                    return True, val, "run_color"

                theme_val = (
                    color_el.get(f"{{{WNS}}}themeColor", "")
                    or color_el.get("w:themeColor", "")
                    or color_el.get("themeColor", "")
                ).lower()
                if theme_val in {"accent1", "accent2", "accent5", "hyperlink"}:
                    return True, None, "theme_color"

        return False, None, None

    def _check_hex_or_rgb(self, hex_str: str) -> bool:
        """Validate if hex or RGB string represents a blue tone."""
        return self.is_blue_hex(hex_str)

    @classmethod
    def is_blue_hex(cls, hex_str: str) -> bool:
        """True when *hex_str* (``RRGGBB``, optional ``#``) is an instruction blue."""
        clean = (hex_str or "").upper().lstrip("#")
        if clean in cls.BLUE_HEX_VALUES:
            return True
        if len(clean) == 6:
            try:
                r = int(clean[0:2], 16)
                g = int(clean[2:4], 16)
                b = int(clean[4:6], 16)
                if (b > 130 and b > r + 30) or (b > 140 and r < 100):
                    return True
            except ValueError:
                pass
        return False

    def _extract_table(
        self, table, sequence: int, doc: Optional[DocxDocument] = None,
        extracted_hashes: set[str] | None = None,
    ):
        """Convert a table and tag cells with blue instruction annotations."""
        ext_table = super()._extract_table(table, sequence, doc=doc, extracted_hashes=extracted_hashes)
        if ext_table is None:
            return None

        for row_idx, row in enumerate(table.rows):
            is_header_row = (row_idx == 0 and bool(ext_table.headers))
            data_idx = row_idx - (1 if ext_table.headers else 0)
            if not is_header_row and (data_idx < 0 or data_idx >= len(ext_table.rows)):
                continue
            target_cells = ext_table.headers if is_header_row else ext_table.rows[data_idx]

            for col_idx, cell in enumerate(row.cells):
                if col_idx >= len(target_cells):
                    continue
                ext_cell = target_cells[col_idx]

                # Check for blue font color in any paragraph of the cell
                for para in cell.paragraphs:
                    is_blue, color_hex, detection_method = self._detect_color(para)
                    if is_blue:
                        blue_parts = []
                        for run in para.runs:
                            run_blue, _, _ = self._detect_run_color(run)
                            if run_blue and run.text.strip():
                                blue_parts.append(run.text.strip())

                        inst_text = " ".join(blue_parts) if blue_parts else ext_cell.content_text
                        if not ext_cell.metadata:
                            ext_cell.metadata = {}
                        ext_cell.metadata["is_instruction"] = True
                        ext_cell.metadata["font_color_hex"] = color_hex
                        ext_cell.metadata["color_detection_method"] = detection_method
                        ext_cell.metadata["instruction_text"] = inst_text
                        break

        return ext_table

    def _extract_inline_images(
        self, paragraph_element, doc: DocxDocument, start_sequence: int,
        extracted_hashes: set[str],
    ) -> list[ExtractedImage]:
        """Extract images embedded inline within a paragraph XML element.

        Template override — saves to template image directory.
        """
        from docx.oxml.ns import qn

        images: list[ExtractedImage] = []
        sequence = start_sequence
        tpl_id = self.template_id or "default"

        r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

        rids: list[str] = []

        for blip in paragraph_element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            embed = blip.get(f'{{{r_ns}}}embed')
            if embed:
                rids.append(embed)
            for s in blip.xpath('.//*[local-name()="svgBlip"]'):
                s_embed = s.get(f'{{{r_ns}}}embed')
                if s_embed and s_embed in doc.part.rels:
                    try:
                        s_bytes = doc.part.rels[s_embed].target_part.blob
                        extracted_hashes.add(hashlib.md5(s_bytes).hexdigest()[:10])
                    except Exception:
                        pass

        for imagedata in paragraph_element.iter('{urn:schemas-microsoft-com:vml}imagedata'):
            rid = imagedata.get(f'{{{r_ns}}}id')
            if rid:
                rids.append(rid)

        for rid in rids:
            try:
                rel = doc.part.rels.get(rid)
                if rel is None or "image" not in rel.reltype:
                    continue

                image_part = rel.target_part
                image_bytes = image_part.blob
                img_hash = hashlib.md5(image_bytes).hexdigest()[:10]

                # If already extracted, reuse the saved path so repeated icons retain the image
                if hasattr(self, "_extracted_hash_paths") and img_hash in self._extracted_hash_paths:
                    save_path = self._extracted_hash_paths[img_hash]
                    images.append(
                        ExtractedImage(
                            content=f"[Image: {save_path.name}]",
                            page=getattr(self, '_current_page', 1),
                            sequence=sequence,
                            image_path=str(save_path),
                            content_hash=img_hash,
                        )
                    )
                    sequence += 1
                    continue

                extracted_hashes.add(img_hash)

                content_type = image_part.content_type or "image/png"
                ext = content_type.split("/")[-1]
                if ext == "jpeg":
                    ext = "jpg"

                filename = f"tpl_img{sequence}_{img_hash}.{ext}"
                save_path = self.settings.get_template_image_dir(tpl_id) / filename

                save_path.write_bytes(image_bytes)
                if hasattr(self, "_extracted_hash_paths"):
                    self._extracted_hash_paths[img_hash] = save_path

                images.append(
                    ExtractedImage(
                        content=f"[Image: {filename}]",
                        page=getattr(self, '_current_page', 1),
                        sequence=sequence,
                        image_path=str(save_path),
                        content_hash=img_hash,
                    )
                )
                sequence += 1

            except Exception:
                self.logger.warning(
                    f"TemplateDocxParser: failed to extract inline image from rId {rid}"
                )
                continue

        return images

    def _extract_images(
        self, doc: DocxDocument, start_sequence: int,
        skip_hashes: set[str] | None = None,
    ) -> list[ExtractedImage]:
        """Extract embedded images into the template's image directory (fallback).

        Images already extracted inline (tracked by *skip_hashes*) are skipped.
        """
        images: list[ExtractedImage] = []
        sequence = start_sequence
        tpl_id = self.template_id or "default"

        for rel in doc.part.rels.values():
            if "image" not in rel.reltype:
                continue

            try:
                image_part = rel.target_part
                image_bytes = image_part.blob
                content_type = image_part.content_type or "image/png"
                ext = content_type.split("/")[-1]
                if ext == "jpeg":
                    ext = "jpg"

                img_hash = hashlib.md5(image_bytes).hexdigest()[:10]

                # Skip images already extracted inline
                if skip_hashes and img_hash in skip_hashes:
                    continue

                filename = f"tpl_img{sequence}_{img_hash}.{ext}"
                save_path = self.settings.get_template_image_dir(tpl_id) / filename

                save_path.write_bytes(image_bytes)

                images.append(
                    ExtractedImage(
                        content=f"[Image: {filename}]",
                        page=getattr(self, '_current_page', 1),
                        sequence=sequence,
                        image_path=str(save_path),
                        content_hash=img_hash,
                    )
                )
                sequence += 1

            except Exception:
                self.logger.warning(
                    f"TemplateDocxParser: failed to extract image from relationship "
                    f"{rel.rId}"
                )
                continue

        return images

