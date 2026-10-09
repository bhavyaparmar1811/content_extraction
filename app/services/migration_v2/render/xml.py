"""Small WordprocessingML helpers for the renderer: paragraphs, runs, cells and element bookkeeping."""

from __future__ import annotations

import copy
import re
from typing import Iterable, Optional

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W_P, W_TBL, W_TR, W_TC, W_SDT, W_R = (qn(f"w:{t}") for t in ("p", "tbl", "tr", "tc", "sdt", "r"))
W_SDT_CONTENT, W_PPR, W_RPR, W_TCPR = qn("w:sdtContent"), qn("w:pPr"), qn("w:rPr"), qn("w:tcPr")

# pPr children in schema order (CT_PPrBase, then rPr / sectPr / pPrChange).
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
              "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
              "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
# tcPr children in schema order.
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
               "tcPrChange"]


def w_el(tag: str, **attrs) -> OxmlElement:
    el = OxmlElement(f"w:{tag}")
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), str(value))
    return el


def text_of(el) -> str:
    """All text under *el*; paragraphs joined by a space."""
    paragraphs = [el] if el.tag == W_P else list(el.iter(W_P))
    return " ".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in paragraphs).strip()


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().lower()


def remove(el) -> None:
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def insert_child(parent, child, order: list[str]) -> None:
    """Insert *child* into *parent* (a pPr or tcPr) at its schema position."""
    name = child.tag.split("}")[1]
    rank = order.index(name) if name in order else len(order)
    for existing in parent:
        other = existing.tag.split("}")[1]
        if other in order and order.index(other) > rank:
            existing.addprevious(child)
            return
    parent.append(child)


def set_ppr_child(ppr, child) -> None:
    for old in ppr.findall(child.tag):
        ppr.remove(old)
    insert_child(ppr, child, _PPR_ORDER)


def set_tcpr_child(tcpr, child) -> None:
    for old in tcpr.findall(child.tag):
        tcpr.remove(old)
    insert_child(tcpr, child, _TCPR_ORDER)


def base_ppr(p) -> Optional[object]:
    """A paragraph's own layout (spacing, alignment...) to reuse for new text: numbering, mark formatting and
    section breaks are dropped, and so is a heading style."""
    ppr = p.find(W_PPR) if p is not None else None
    if ppr is None:
        return None
    out = copy.deepcopy(ppr)
    for name in ("numPr", "rPr", "sectPr", "pPrChange"):
        for child in out.findall(qn(f"w:{name}")):
            out.remove(child)
    style = out.find(qn("w:pStyle"))
    if style is not None and (style.get(qn("w:val")) or "").lower().startswith(("heading", "title", "toc")):
        out.remove(style)
    return out


def base_rpr(el) -> Optional[object]:
    """Font face and size of the first text run under *el*; colour, emphasis and highlight are not kept."""
    for r in el.iter(W_R) if el is not None else ():
        if not "".join(t.text or "" for t in r.iter(qn("w:t"))).strip():
            continue
        rpr = r.find(W_RPR)
        if rpr is None:
            return None
        out = w_el("rPr")
        for name in ("rFonts", "sz", "szCs", "lang"):
            node = rpr.find(qn(f"w:{name}"))
            if node is not None:
                out.append(copy.deepcopy(node))
        return out if len(out) else None
    return None


def paragraph(text: str = "", ppr=None, rpr=None, *, bold: bool = False, italic: bool = False,
              colour: Optional[str] = None, highlight: Optional[str] = None):
    p = w_el("p")
    if ppr is not None:
        p.append(copy.deepcopy(ppr))
    if text:
        add_runs(p, text, rpr, bold=bold, italic=italic, colour=colour, highlight=highlight)
    return p


# A REF field inside run text: FIELD_START bookmark FIELD_MID shown number FIELD_END (``ref_field``).
FIELD_START, FIELD_MID, FIELD_END = "\ue000", "\ue001", "\ue002"  # private-use characters, never in source text
_FIELD = re.compile(f"{FIELD_START}([^{FIELD_MID}]*){FIELD_MID}([^{FIELD_END}]*){FIELD_END}")


def ref_field(bookmark: str, shown: str) -> str:
    """Marks *shown* as a Word REF field to *bookmark* (full paragraph number) for ``add_runs``."""
    return f"{FIELD_START}{bookmark}{FIELD_MID}{shown}{FIELD_END}"


def plain_text(text: str) -> str:
    """The text as the reader sees it: REF field marks replaced by their shown number."""
    return _FIELD.sub(lambda m: m.group(2), text or "")


def add_runs(p, text: str, rpr=None, *, bold: bool = False, italic: bool = False,
             colour: Optional[str] = None, highlight: Optional[str] = None) -> None:
    """Runs for *text*; line breaks become ``w:br``, tabs ``w:tab``, and ``ref_field`` marks REF fields."""
    props = copy.deepcopy(rpr) if rpr is not None else w_el("rPr")
    # rPr children in schema order: rFonts, b, bCs, i, iCs, ..., color, ..., sz, szCs, highlight, ...
    extras = []
    if bold:
        extras += [w_el("b"), w_el("bCs")]
    if italic:
        extras += [w_el("i"), w_el("iCs")]
    fonts = props.find(qn("w:rFonts"))
    anchor = fonts if fonts is not None else None
    for extra in extras:
        if anchor is None:
            props.insert(0, extra)
        else:
            anchor.addnext(extra)
        anchor = extra
    if colour:
        node = w_el("color", val=colour)
        size = props.find(qn("w:sz"))
        size.addprevious(node) if size is not None else props.append(node)
    if highlight:
        node = w_el("highlight", val=highlight)
        lang = props.find(qn("w:lang"))
        lang.addprevious(node) if lang is not None else props.append(node)

    def run():
        r = w_el("r")
        if len(props):
            r.append(copy.deepcopy(props))
        p.append(r)
        return r

    last = 0
    for m in _FIELD.finditer(text):
        _text_into(run(), text[last:m.start()]) if m.start() > last else None
        for part in ("begin", "instr", "separate", "shown", "end"):
            r = run()
            if part == "instr":
                instr = w_el("instrText")
                instr.text = f" REF {m.group(1)} \\w \\h "
                instr.set(qn("xml:space"), "preserve")
                r.append(instr)
            elif part == "shown":
                _text_into(r, m.group(2))
            else:
                r.append(w_el("fldChar", fldCharType=part))
        last = m.end()
    if last < len(text) or not text:
        _text_into(run(), text[last:])


def _text_into(r, text: str) -> None:
    for i, line in enumerate(text.split("\n")):
        if i:
            r.append(w_el("br"))
        for j, chunk in enumerate(line.split("\t")):
            if j:
                r.append(w_el("tab"))
            if chunk:
                t = w_el("t")
                t.text = chunk
                t.set(qn("xml:space"), "preserve")
                r.append(t)


def add_bookmark(p, name: str, bookmark_id: int) -> None:
    """A bookmark around paragraph *p*'s content (after its pPr), unless *p* already has one of that name."""
    if any(b.get(qn("w:name")) == name for b in p.iter(qn("w:bookmarkStart"))):
        return
    start = w_el("bookmarkStart", id=bookmark_id, name=name)
    ppr = p.find(W_PPR)
    (ppr.addnext(start) if ppr is not None else p.insert(0, start))
    p.append(w_el("bookmarkEnd", id=bookmark_id))


def next_bookmark_id(root) -> int:
    ids = [int(b.get(qn("w:id"))) for b in root.iter(qn("w:bookmarkStart")) if (b.get(qn("w:id")) or "").lstrip("-").isdigit()]
    return max(ids, default=0) + 1


def make_plain(el) -> None:
    """Drop colour and highlight from every run (and paragraph mark) under *el*: kept template text turns black."""
    for node in list(el.iter(qn("w:color"), qn("w:highlight"))):
        parent = node.getparent()
        if parent is not None and parent.tag == W_RPR:
            parent.remove(node)


def ensure_ends_with_paragraph(container) -> None:
    """A table cell (or a content control inside one) must end with a paragraph."""
    blocks = [c for c in container if c.tag in (W_P, W_TBL, W_SDT)]
    if not blocks or blocks[-1].tag != W_P:
        container.append(w_el("p"))


def cell_container(tc):
    """Where a cell's blocks live: its single block content control if it has one, else the cell."""
    sdts = [c for c in tc if c.tag == W_SDT]
    if len(sdts) == 1 and sdts[0].find(W_SDT_CONTENT) is not None:
        return sdts[0].find(W_SDT_CONTENT)
    return tc


def clear_blocks(container) -> list:
    """Remove a container's block children (paragraphs, tables, controls); returns them."""
    removed = [c for c in container if c.tag in (W_P, W_TBL, W_SDT)]
    for c in removed:
        container.remove(c)
    return removed


def unwrap(sdt) -> list:
    """Replace a content control by its content; returns the moved children."""
    content = sdt.find(W_SDT_CONTENT)
    children = list(content) if content is not None else []
    for child in children:
        sdt.addprevious(child)
    remove(sdt)
    return children


def sdt_tag(sdt) -> Optional[str]:
    tag = sdt.find(f"{qn('w:sdtPr')}/{qn('w:tag')}")
    return tag.get(qn("w:val")) if tag is not None else None


def ancestor(el, tag):
    parent = el.getparent()
    while parent is not None and parent.tag != tag:
        parent = parent.getparent()
    return parent


def iter_paragraphs(blocks: Iterable) -> Iterable:
    for block in blocks:
        if block.tag == W_P:
            yield block
        else:
            yield from block.iter(W_P)
