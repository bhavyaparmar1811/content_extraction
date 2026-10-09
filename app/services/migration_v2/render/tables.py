"""Source table rows into Word tables.

A table slot fills the template's own table: the header rows stay, each source
row becomes a copy of the first example row, and the example rows go. Columns
are matched by position; a column whose source and template headers share no
word is filled anyway and reported for the reviewer.

A source table the template table cannot hold (more columns than the template,
or a template header with placeholders such as "[Role 1]") is written in its own
shape instead, with its own header row, in the template table's formatting
(table properties, header and body cell properties). No source cell is ever
dropped. Each further source table in the same slot gets its own table.

Source rows in a free-text slot (e.g. a table inside PROCESS) are always written
in their own shape, formatted like the template's first data table.
"""

from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from typing import Optional

from docx.oxml.ns import qn

from app.schemas.v2 import Claim, SourceUnit

from .xml import (
    W_P,
    W_SDT,
    W_SDT_CONTENT,
    W_TC,
    W_TCPR,
    W_TR,
    base_ppr,
    base_rpr,
    clear_blocks,
    paragraph,
    set_tcpr_child,
    text_of,
    unwrap,
    w_el,
)

_PLACEHOLDER_HEADER = re.compile(r"\[[^\]]+\]")
_STOP = {"the", "and", "for", "of", "to", "in", "or", "a", "an", "no"}


@dataclass
class RowCell:
    text: str
    fill_hex: Optional[str] = None


@dataclass
class SourceRow:
    table_id: str
    header: list[str]
    cells: list[RowCell]
    claim: Claim


def source_row(claim: Claim, unit: SourceUnit, warnings: list[str]) -> SourceRow:
    """The claim's row in logical columns: header cells left empty by a merge belong to the column before."""
    ref = unit.table_ref
    starts: list[int] = []
    header: list[str] = []
    for col, text in enumerate(ref.header_cells):
        if text.strip() or not starts:
            starts.append(col)
            header.append(text.strip())
    if not starts:
        width = max((c.col + c.col_span for c in ref.cells), default=1)
        starts, header = list(range(width)), [""] * width
    cells = [RowCell("") for _ in starts]
    if len(ref.cells) == len(starts):
        for i, c in enumerate(ref.cells):
            cells[i] = RowCell(_cell_text(c), c.fill_hex)
    else:
        for c in ref.cells:
            i = max((k for k, s in enumerate(starts) if s <= c.col), default=0)
            joined = "\n".join(t for t in (cells[i].text, _cell_text(c)) if t)
            cells[i] = RowCell(joined, cells[i].fill_hex or c.fill_hex)
    if claim.text.strip() != unit.text.strip():  # a reviewer edited the row
        parts = [p.strip() for p in claim.text.split("|")]
        if len(parts) == len(cells):
            cells = [RowCell(t, c.fill_hex) for t, c in zip(parts, cells)]
        else:
            warnings.append(f"{claim.claim_id}: edited row has {len(parts)} cells, the source table "
                            f"{ref.table_id} has {len(cells)}; the source cells were used")
    return SourceRow(ref.table_id, header, cells, claim)


def _cell_text(cell) -> str:
    """A source cell's text with its paragraphs on separate lines."""
    return "\n".join(cell.paragraphs) if cell.paragraphs else cell.text


def group_by_table(rows: list[SourceRow]) -> list[list[SourceRow]]:
    """Consecutive rows of the same source table."""
    groups: list[list[SourceRow]] = []
    for row in rows:
        if groups and groups[-1][0].table_id == row.table_id:
            groups[-1].append(row)
        else:
            groups.append([row])
    return groups


# ── Template tables ───────────────────────────────────────────────────


def row_cells(tr) -> list:
    """A row's cells, including cells inside cell-level content controls."""
    out = []
    for child in tr:
        if child.tag == W_TC:
            out.append(child)
        elif child.tag == W_SDT:
            content = child.find(W_SDT_CONTENT)
            out.extend(content.findall(qn("w:tc")) if content is not None else [])
    return out


def _span(tc) -> int:
    span = tc.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}")
    return int(span.get(qn("w:val"))) if span is not None else 1


def _grid_positions(tr) -> list[tuple[int, int, object]]:
    out, col = [], 0
    for tc in row_cells(tr):
        out.append((col, _span(tc), tc))
        col += _span(tc)
    return out


def table_rows(tbl) -> list:
    rows = []
    for child in tbl:
        if child.tag == W_TR:
            rows.append(child)
        elif child.tag == W_SDT:
            content = child.find(W_SDT_CONTENT)
            rows.extend(content.findall(qn("w:tr")) if content is not None else [])
    return rows


@dataclass
class TemplateTable:
    """A template table split into header rows and the example rows a slot replaces."""

    tbl: object
    header_rows: list
    example_rows: list
    prototype: object = field(init=False)
    columns: int = field(init=False)
    header_texts: list[str] = field(init=False)

    def __post_init__(self):
        self.prototype = copy.deepcopy(self.example_rows[0] if self.example_rows else self.header_rows[-1])
        for sdt in list(self.prototype.iter(W_SDT)):
            unwrap(sdt)
        positions = _grid_positions(self.prototype)
        self.columns = len(positions)
        self.header_texts = []
        for start, _, _ in positions:
            texts = []
            for hr in self.header_rows:
                cover = next((tc for s, n, tc in _grid_positions(hr) if s <= start < s + n), None)
                texts.append(text_of(cover) if cover is not None else "")
            self.header_texts.append(next((t for t in reversed(texts) if t), ""))

    @property
    def has_placeholder_header(self) -> bool:
        return any(_PLACEHOLDER_HEADER.search(text_of(hr)) for hr in self.header_rows)

    def fits(self, header: list[str]) -> bool:
        return not self.has_placeholder_header and len(header) <= self.columns


def _words(text: str) -> set[str]:
    """Word stems: "Responsibilities" and "Responsibility" share one."""
    return {w[:6] for w in re.findall(r"[a-z]{3,}", text.lower()) if w not in _STOP}


def _plain(text: str) -> str:
    return re.sub(r"[\W_]+", "", text.lower())


def header_mismatches(template: TemplateTable, header: list[str]) -> list[str]:
    out = []
    for i, src in enumerate(header):
        tpl = template.header_texts[i] if i < len(template.header_texts) else ""
        if src and tpl and _plain(src) != _plain(tpl) and not (_words(src) & _words(tpl)):
            out.append(f"source column '{src}' is under template column '{tpl}'")
    return out


def fill_cell(tc, text: str, fill_hex: Optional[str] = None, bold: bool = False) -> None:
    """Replace a cell's blocks with *text*, keeping the cell's paragraph layout and font size."""
    first = tc.find(qn("w:p"))
    ppr, rpr = base_ppr(first), base_rpr(tc)
    for sdt in [c for c in tc if c.tag == W_SDT]:
        unwrap(sdt)
    clear_blocks(tc)
    lines = text.split("\n") if text else [""]
    for line in lines:
        tc.append(paragraph(line, ppr, rpr, bold=bold))
    if fill_hex:
        tcpr = tc.find(W_TCPR)
        if tcpr is None:
            tcpr = w_el("tcPr")
            tc.insert(0, tcpr)
        set_tcpr_child(tcpr, w_el("shd", val="clear", color="auto", fill=fill_hex))


def build_row(prototype, cells: list[RowCell]):
    tr = copy.deepcopy(prototype)
    positions = _grid_positions(tr)
    for i, (_, _, tc) in enumerate(positions):
        cell = cells[i] if i < len(cells) else RowCell("")
        fill_cell(tc, cell.text, cell.fill_hex)
    return tr


def gap_row(prototype, marker) -> object:
    """One row spanning the table, holding the gap marker paragraph."""
    tr = copy.deepcopy(prototype)
    positions = _grid_positions(tr)
    first = positions[0][2]
    total = sum(n for _, n, _ in positions)
    for _, _, tc in positions[1:]:
        tc.getparent().remove(tc)
    fill_cell(first, "")
    for p in first.findall(W_P):
        first.remove(p)
    first.append(marker)
    tcpr = first.find(W_TCPR)
    if tcpr is None:
        tcpr = w_el("tcPr")
        first.insert(0, tcpr)
    if total > 1:
        set_tcpr_child(tcpr, w_el("gridSpan", val=total))
    return tr


# ── Tables in the source's shape ──────────────────────────────────────


@dataclass
class TableStyle:
    """Formatting taken from a template table: table properties, a header cell and a body cell."""

    tbl_pr: object
    width: int
    header_tr: object
    header_tc: object
    body_tr: object
    body_tc: object
    header_bold: bool = False

    @classmethod
    def from_template(cls, template: TemplateTable) -> "TableStyle":
        tbl_pr = copy.deepcopy(template.tbl.find(qn("w:tblPr")))
        grid = template.tbl.find(qn("w:tblGrid"))
        width = sum(int(float(g.get(qn("w:w")) or 0)) for g in grid.findall(qn("w:gridCol"))) if grid is not None else 0
        header_tr = copy.deepcopy(template.header_rows[0]) if template.header_rows else copy.deepcopy(template.prototype)
        header_tc = row_cells(header_tr)[0]
        bold = any(b.get(qn("w:val")) not in ("0", "false") for b in header_tc.iter(qn("w:b")))
        return cls(tbl_pr, width or 9000, header_tr, header_tc, template.prototype, row_cells(template.prototype)[0], bold)

    @classmethod
    def plain(cls) -> "TableStyle":
        tbl_pr = w_el("tblPr")
        tbl_pr.append(w_el("tblW", w=0, type="auto"))
        borders = w_el("tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(w_el(side, val="single", sz=4, space=0, color="000000"))
        tbl_pr.append(borders)
        tr = w_el("tr")
        tc = w_el("tc")
        tc.append(w_el("p"))
        tr.append(tc)
        return cls(tbl_pr, 9000, tr, tc, copy.deepcopy(tr), copy.deepcopy(tc))


def _row_like(tr_proto, tc_proto, texts: list[RowCell], widths: list[int], bold: bool, header: bool):
    tr = w_el("tr")
    tr_pr = tr_proto.find(qn("w:trPr"))
    if tr_pr is not None:
        tr_pr = copy.deepcopy(tr_pr)
        if not header:  # a long source row may break across pages instead of leaving a blank half page
            for node in tr_pr.findall(qn("w:cantSplit")) + tr_pr.findall(qn("w:tblHeader")):
                tr_pr.remove(node)
        if len(tr_pr):
            tr.append(tr_pr)
    for cell, width in zip(texts, widths):
        tc = copy.deepcopy(tc_proto)
        for sdt in [c for c in tc if c.tag == W_SDT]:
            unwrap(sdt)
        tcpr = tc.find(W_TCPR)
        if tcpr is None:
            tcpr = w_el("tcPr")
            tc.insert(0, tcpr)
        for name in ("gridSpan", "vMerge", "hMerge"):
            for node in tcpr.findall(qn(f"w:{name}")):
                tcpr.remove(node)
        set_tcpr_child(tcpr, w_el("tcW", w=width, type="dxa"))
        fill_cell(tc, cell.text, cell.fill_hex, bold=bold)
        tr.append(tc)
    return tr


def column_widths(total: int, header: list[str], rows: list[SourceRow], columns: int) -> list[int]:
    """Widths that follow the content: the longest word must fit, longer text gets more room."""
    weights = []
    for j in range(columns):
        texts = [r.cells[j].text for r in rows if j < len(r.cells)] + ([header[j]] if j < len(header) else [])
        longest = max((len(w) for t in texts for w in t.split()), default=1)
        average = sum(len(t) for t in texts) / max(len(texts), 1)
        weights.append(max(longest, 4) + 2 * min(average, 400) ** 0.5)
    scale = total / sum(weights)
    return [int(w * scale) for w in weights]


def _header_repeats_rows(header: list[str], rows: list[SourceRow]) -> bool:
    """A "header" that is only the rows' own text joined (a table whose every row was marked as a header)."""
    if len(rows) < 2:
        return False
    for j, label in enumerate(header):
        column = [r.cells[j].text for r in rows if j < len(r.cells) and r.cells[j].text]
        if label != " / ".join(dict.fromkeys(column)):
            return False
    return True


def source_shaped_table(style: TableStyle, rows: list[SourceRow]):
    columns = max(len(rows[0].header), max(len(r.cells) for r in rows))
    header = rows[0].header + [""] * (columns - len(rows[0].header))
    if _header_repeats_rows(rows[0].header, rows):
        header = [""] * columns
    widths = column_widths(style.width, header, rows, columns)
    tbl = w_el("tbl")
    tbl_pr = copy.deepcopy(style.tbl_pr)
    for node in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(node)
    tbl_w = w_el("tblW", w=sum(widths), type="dxa")
    anchor = tbl_pr.find(qn("w:tblStyle"))
    anchor.addnext(tbl_w) if anchor is not None else tbl_pr.insert(0, tbl_w)
    tbl.append(tbl_pr)
    grid = w_el("tblGrid")
    for width in widths:
        grid.append(w_el("gridCol", w=width))
    tbl.append(grid)
    if any(header):
        tbl.append(_row_like(style.header_tr, style.header_tc, [RowCell(h) for h in header], widths,
                             bold=style.header_bold, header=True))
    for row in rows:
        cells = row.cells + [RowCell("")] * (columns - len(row.cells))
        tbl.append(_row_like(style.body_tr, style.body_tc, cells, widths, bold=False, header=False))
    return tbl
