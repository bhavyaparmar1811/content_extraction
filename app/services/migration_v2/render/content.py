"""A slot's claims as Word blocks: paragraphs, real lists, sub-headings, figures, captions, tables, callout boxes.

Grouping, in claim order:
- consecutive claims with a callout kind (other than the slot's own fixed kind)
  go into one cloned callout box;
- consecutive table rows of the same source table become one table in the
  source's shape (table slots fill the template table instead, in the renderer);
- consecutive bullets and steps form one list: steps share one numbered list
  that restarts at 1, bullets the template's bullet list, nesting by list level.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from docx.shared import Emu
from docx.text.paragraph import Paragraph

from app.schemas.v2 import CalloutKind, Claim, ClaimKind, SourceUnit
from app.schemas.v2.refs import REF_TOKEN_PATTERN

from .callouts import Callouts
from .numbering import ListNumbering
from .tables import TableStyle, group_by_table, source_row, source_shaped_table
from .xml import W_PPR, W_TBL, add_bookmark, paragraph, set_ppr_child, w_el

GAP_MARKER_TEXT = "Source content not found. Human review required."
_EMU_PER_PX = 9525  # 96 dpi

RefText = Callable[[str, Claim], str]  # (ref target, claim) → text to show for a {{ref:...}} token


@dataclass
class ContentContext:
    doc: object
    numbering: ListNumbering
    callouts: Callouts
    units: dict[str, SourceUnit]
    ref_text: RefText
    table_style: TableStyle
    max_image_emu: int
    section_level: int = 1
    bookmarks: dict[str, str] = field(default_factory=dict)  # heading claim_id → bookmark (REF field target)
    bookmark_id: int = 1
    warnings: list[str] = field(default_factory=list)
    callout_boxes: int = 0
    tables: int = 0

    def text(self, claim: Claim) -> str:
        return show_refs(claim, self.ref_text)


def show_refs(claim: Claim, ref_text: RefText) -> str:
    """The claim's text with each ``{{ref:...}}`` token shown by *ref_text*. Words the shown phrase starts with are
    dropped when the text already has them just before the token ("described in" + "described in 6.2")."""
    text = claim.text or ""
    out, last = [], 0
    for match in REF_TOKEN_PATTERN.finditer(text):
        before = text[last:match.start()]
        phrase = ref_text(match.group(1), claim)
        words = phrase.split()
        for k in range(len(words) - 1, 0, -1):
            lead = " ".join(words[:k])
            if re.search(rf"(?:^|\W){re.escape(lead)}\s*$", "".join(out) + before, re.I):
                phrase = " ".join(words[k:])
                break
        out.append(before + phrase)
        last = match.end()
    out.append(text[last:])
    return "".join(out)


def gap_marker(ppr=None):
    return paragraph(GAP_MARKER_TEXT, ppr, bold=True, colour="C00000", highlight="yellow")


def build_blocks(claims: list[Claim], ctx: ContentContext, ppr=None, rpr=None,
                 own_callout: Optional[CalloutKind] = None) -> list:
    """Blocks for *claims*. *ppr* / *rpr* are the anchor's paragraph layout and font, reused for body text."""
    blocks: list = []
    i = 0
    while i < len(claims):
        claim = claims[i]
        kind = claim.callout_kind if claim.callout_kind != own_callout else None
        if kind is not None:
            j = i
            while j < len(claims) and claims[j].callout_kind == kind:
                j += 1
            inner = build_blocks(claims[i:j], ctx, ppr, rpr, own_callout=kind)
            box = ctx.callouts.box(kind, inner)
            if box is None:
                ctx.warnings.append(f"no callout box for kind '{kind.value}'; {claims[i].claim_id}.. rendered as text")
                blocks += inner
            else:
                ctx.callout_boxes += 1
                blocks.append(box)
            i = j
            continue
        if claim.kind == ClaimKind.TABLE_ROW and not claim.is_gap_marker:
            j = i
            while j < len(claims) and claims[j].kind == ClaimKind.TABLE_ROW and claims[j].callout_kind == claim.callout_kind:
                j += 1
            blocks += table_blocks(claims[i:j], ctx, ppr, rpr)
            i = j
            continue
        if claim.kind in (ClaimKind.BULLET, ClaimKind.STEP) and not claim.is_gap_marker:
            j = i
            while j < len(claims) and claims[j].kind in (ClaimKind.BULLET, ClaimKind.STEP) and claims[j].callout_kind == claim.callout_kind:
                j += 1
            blocks += list_blocks(claims[i:j], ctx, ppr, rpr)
            i = j
            continue
        blocks += single_blocks(claim, ctx, ppr, rpr)
        i += 1
    return _separate_tables(blocks)


def single_blocks(claim: Claim, ctx: ContentContext, ppr=None, rpr=None) -> list:
    if claim.is_gap_marker:
        return [gap_marker(ppr)]
    if claim.kind == ClaimKind.HEADING:
        return [heading_block(claim, ctx)]
    if claim.kind == ClaimKind.FIGURE:
        return figure_blocks(claim, ctx)
    if claim.kind == ClaimKind.CAPTION:
        return [caption_block(claim, ctx, ppr)]
    return [paragraph(ctx.text(claim), ppr, rpr)]


def heading_block(claim: Claim, ctx: ContentContext):
    level = ctx.section_level + 1 + claim.list_level
    ppr = w_el("pPr")
    style = ctx.numbering.heading_style_id(level)
    if style:
        set_ppr_child(ppr, w_el("pStyle", val=style))
    if style and ctx.numbering.heading_num_id is not None:
        num_pr = w_el("numPr")
        num_pr.append(w_el("ilvl", val=min(level - 1, 8)))
        num_pr.append(w_el("numId", val=ctx.numbering.heading_num_id))
        set_ppr_child(ppr, num_pr)
    p = paragraph(ctx.text(claim), ppr, bold=style is None)
    name = ctx.bookmarks.get(claim.claim_id)
    if name:
        add_bookmark(p, name, ctx.bookmark_id)
        ctx.bookmark_id += 1
    return p


def list_blocks(claims: list[Claim], ctx: ContentContext, ppr=None, rpr=None) -> list:
    ordered_num = bullet_num = None
    out = []
    for claim in claims:
        if claim.kind == ClaimKind.STEP:
            ordered_num = ordered_num or ctx.numbering.ordered_list()
            num = ordered_num
        else:
            bullet_num = bullet_num or ctx.numbering.bullet_list()
            num = bullet_num
        p = paragraph(ctx.text(claim), ppr, rpr)
        p_ppr = p.find(W_PPR)
        if p_ppr is None:
            p_ppr = w_el("pPr")
            p.insert(0, p_ppr)
        ListNumbering.apply(p_ppr, num, claim.list_level)
        out.append(p)
    return out


def table_blocks(claims: list[Claim], ctx: ContentContext, ppr=None, rpr=None) -> list:
    rows = []
    for claim in claims:
        unit = ctx.units.get(claim.source_unit_ids[0]) if claim.source_unit_ids else None
        if unit is None or unit.table_ref is None:
            ctx.warnings.append(f"{claim.claim_id}: table row without a source table; rendered as text")
            rows.append(None)
            continue
        rows.append(source_row(claim, unit, ctx.warnings))
    out = []
    pending = [r for r in rows if r is not None]
    for group in group_by_table(pending):
        out.append(source_shaped_table(ctx.table_style, group))
        ctx.tables += 1
    out += [paragraph(ctx.text(c), ppr, rpr) for c, r in zip(claims, rows) if r is None]
    return out


def figure_blocks(claim: Claim, ctx: ContentContext) -> list:
    unit = ctx.units.get(claim.source_unit_ids[0]) if claim.source_unit_ids else None
    asset = next((a for a in (unit.assets if unit else []) if a.kind.value == "figure"), None)
    ppr = w_el("pPr")
    set_ppr_child(ppr, w_el("jc", val="center"))
    p = paragraph("", ppr)
    if asset is None or not asset.path or not Path(asset.path).exists():
        ctx.warnings.append(f"{claim.claim_id}: figure file not found ({asset.path if asset else 'no asset'}); "
                            "a placeholder was written")
        return [paragraph(f"[Figure not available: {claim.text or claim.source_unit_ids[0]}]", ppr,
                          italic=True, colour="C00000")]
    width = height = None
    if asset.width_px and asset.height_px:
        width = min(asset.width_px * _EMU_PER_PX, ctx.max_image_emu)
        height = int(width * asset.height_px / asset.width_px)
    try:
        Paragraph(p, ctx.doc._body).add_run().add_picture(
            asset.path, width=Emu(width) if width else None, height=Emu(height) if height else None)
    except Exception as exc:  # an unreadable image must not stop the document
        ctx.warnings.append(f"{claim.claim_id}: figure could not be inserted ({exc}); a placeholder was written")
        return [paragraph(f"[Figure not available: {claim.source_unit_ids[0]}]", ppr, italic=True, colour="C00000")]
    return [p]


def caption_block(claim: Claim, ctx: ContentContext, ppr=None):
    try:
        style = ctx.doc.styles["Caption"].style_id
    except KeyError:
        style = None
    cap_ppr = w_el("pPr")
    if style:
        set_ppr_child(cap_ppr, w_el("pStyle", val=style))
    set_ppr_child(cap_ppr, w_el("jc", val="center"))
    return paragraph(ctx.text(claim), cap_ppr, italic=style is None)


def _separate_tables(blocks: list) -> list:
    """Word merges two tables that touch; keep an empty paragraph between them."""
    out = []
    for block in blocks:
        if out and out[-1].tag == W_TBL and block.tag == W_TBL:
            out.append(w_el("p"))
        out.append(block)
    return out

