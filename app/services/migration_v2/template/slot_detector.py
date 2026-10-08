"""Detect the target slots of a Word template and build its ``TemplateModel`` (Phase 3a).

A slot is a place the migrated content goes. Detection follows this priority:

1. Content controls tagged ``CC_<SECTION>_<KEY>`` (what the normalizer writes).
2. Bookmarks named ``SLOT_<SECTION>_<KEY>``.
3. Placeholders such as ``[Insert title]`` or ``<<owner>>`` in black text.
4. Table pairs: an image-only icon cell beside a blue instruction cell (one
   slot per row), a data table (header plus example rows), and the callout
   prototype boxes found by ``callout_palette``.
5. Blue instruction paragraphs. A group of them that leads into a table is
   that table's instruction; otherwise the group itself is the destination.
6. Icons are metadata on their row's slot, never a slot of their own.

Regions the rules can't classify, such as black text with an inline blue
choice ("This Directive/SOP/...:") or a table written entirely in blue, are
reported as ``AmbiguousRegion``s. A ``TemplateConfig`` decides them by region ID:
fixed, instruction, slot, or conditional. A conditional region becomes a
``ConditionalRegion`` on its section, settled per migration from the SOP's
content (for example "This SOP:", or keeping the competence table only when
the SOP assigns competences).

Region IDs are stable for one template version, and normalization keeps them:
``p:<n>`` is the n-th body paragraph and ``t:<n>`` the n-th body table, both
counted by ``iter_body_blocks`` (block content controls are flattened and the
TOC is skipped).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from app.schemas.v2 import (
    AnchorKind,
    CalloutKind,
    ConditionalKind,
    ConditionalRegion,
    ContentType,
    FormattingProfile,
    InstructionBehavior,
    SlotAnchor,
    SlotIcon,
    TargetSection,
    TargetSlot,
    TemplateModel,
)
from app.services.parser.ooxml import NumberingResolver, iter_body_blocks
from app.services.parser.template_parser import TemplateDocxParser

from .callout_palette import (
    PaletteResult,
    block_text,
    icon_key_for,
    is_heading,
    is_image_only,
    palette_from_document,
    style_name,
)
from .overrides import RegionDecision, TemplateConfig

TAG_PREFIX = "CC_"
COND_PREFIX = "COND_"  # conditional regions; never confused with a slot's CC_ tag
BOOKMARK_PREFIX = "SLOT_"

_PLACEHOLDER = re.compile(r"\[Insert\s+[^\]]+\]|<<[^>]+>>", re.I)

# Icon-row instructions of the GP templates → slot key. First match wins.
_ICON_ROW_KEYS: list[tuple[str, re.Pattern]] = [
    ("what", re.compile(r"what the document is about|what process", re.I)),
    ("intention", re.compile(r"intention|want to achieve", re.I)),
    ("roles", re.compile(r"target roles|which roles", re.I)),
    ("units", re.compile(r"business units|group functions|departments", re.I)),
    ("geography", re.compile(r"geograph|world-?wide", re.I)),
    ("processes", re.compile(r"processes\s*/\s*systems|affected processes|materials", re.I)),
]

_PARAGRAPH_KEYS: list[tuple[str, re.Pattern]] = [
    ("not_covered", re.compile(r"not covered|out of scope|not in scope", re.I)),
]

# Header of a data table's first column → slot key.
_TABLE_KEYS: list[tuple[str, re.Pattern]] = [
    ("terms", re.compile(r"^term", re.I)),
    ("abbreviations", re.compile(r"^abbreviation", re.I)),
    ("roles", re.compile(r"^role", re.I)),
    ("raci", re.compile(r"^process step", re.I)),
    ("documents", re.compile(r"^no\.?$", re.I)),
    ("entries", re.compile(r"^version", re.I)),
]

# Section heading → content type of its slots.
_SECTION_TYPES: list[tuple[re.Pattern, ContentType]] = [
    (re.compile(r"purpose", re.I), ContentType.PURPOSE),
    (re.compile(r"applicab|scope", re.I), ContentType.SCOPE),
    (re.compile(r"definition|abbreviation|glossary", re.I), ContentType.DEFINITION),
    (re.compile(r"implementation|pre-?requisite", re.I), ContentType.PREREQUISITE),
    (re.compile(r"role|responsib", re.I), ContentType.RESPONSIBILITY),
    (re.compile(r"associated document|reference", re.I), ContentType.REFERENCE),
    (re.compile(r"history", re.I), ContentType.RECORD),
    (re.compile(r"process|procedure", re.I), ContentType.ORDERED_PROCEDURE),
]

# Instruction wording → content type, for sections the headings don't settle.
_INSTRUCTION_TYPES: list[tuple[re.Pattern, ContentType]] = [
    (re.compile(r"warning|caution|attention", re.I), ContentType.WARNING),
    (re.compile(r"responsib|accountab", re.I), ContentType.RESPONSIBILITY),
    (re.compile(r"\bstep|procedure|how[- ]style", re.I), ContentType.ORDERED_PROCEDURE),
    (re.compile(r"record|archiv|retain", re.I), ContentType.RECORD),
    (re.compile(r"definition|abbreviation", re.I), ContentType.DEFINITION),
]

_REQUIRED_EMPTY_VALUE = re.compile(r"insert\s*[\"“(]*\s*(?:none|n/a)", re.I)
_OPTIONAL_WORDING = re.compile(r"\bif applicable\b|\bif none\b|\boptional\b|\bif needed\b|\bwhere possible\b|\bif not used\b", re.I)


# ── Results ────────────────────────────────────────────────────────────

@dataclass
class Region:
    """Where a slot's content goes, as XML elements of ``DetectionResult.document``."""

    kind: str  # "paragraphs" | "cell" | "rows" | "tagged" (already inside a CC_ content control)
    region_id: str
    elements: list = field(default_factory=list)


@dataclass
class AmbiguousRegion:
    region_id: str
    kind: str  # "inline_instruction" | "instruction_table"
    section_key: Optional[str]
    text: str
    suggestion: RegionDecision
    decision: Optional[RegionDecision] = None


@dataclass
class DetectionResult:
    model: TemplateModel
    document: object
    regions: dict[str, Region] = field(default_factory=dict)  # slot_id → region
    palette: PaletteResult = field(default_factory=PaletteResult)
    ambiguous: list[AmbiguousRegion] = field(default_factory=list)
    icons_without_instruction: list[str] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)
    conditional_elements: dict[str, object] = field(default_factory=dict)  # region_id → paragraph or table

    @property
    def unresolved(self) -> list[AmbiguousRegion]:
        return [a for a in self.ambiguous if a.decision is None]


# ── Detection ──────────────────────────────────────────────────────────

def detect_template_model(
    docx_path: str | Path,
    template_id: str,
    template_version: int = 1,
    config: Optional[TemplateConfig] = None,
) -> DetectionResult:
    doc = Document(str(docx_path))
    config = config or TemplateConfig()
    detector = _Detector(doc, config)
    detector.run()
    model = TemplateModel(
        template_id=template_id,
        template_version=template_version,
        source_file=str(docx_path),
        sections=detector.sections,
        callout_palette=detector.palette.palette,
    )
    return DetectionResult(
        model=model,
        document=doc,
        regions=detector.regions,
        conditional_elements=detector.conditional_elements,
        palette=detector.palette,
        ambiguous=detector.ambiguous,
        icons_without_instruction=detector.icons_without_instruction,
        issues=detector.palette.issues + detector.issues,
    )


@dataclass
class _Pending:
    """Blue instruction paragraphs waiting to learn whether they lead into a table."""

    paragraphs: list = field(default_factory=list)
    first_index: Optional[int] = None
    texts: list[str] = field(default_factory=list)  # in reading order, with instruction tables folded in
    bookmark: Optional[str] = None

    @property
    def text(self) -> Optional[str]:
        return "\n".join(t for t in self.texts if t) or None

    def pop_last_paragraph(self):
        p = self.paragraphs.pop()
        text = block_text(p)
        if text in self.texts:
            del self.texts[len(self.texts) - 1 - self.texts[::-1].index(text)]
        if not self.paragraphs:
            self.first_index = None
        return p


class _Detector:
    def __init__(self, doc, config: TemplateConfig):
        self.doc = doc
        self.config = config
        self.palette = palette_from_document(doc, config.callout_palette)
        self.prototypes = {p.table_index: p for p in self.palette.prototypes}
        self.numbering = NumberingResolver(doc)
        self.sections: list[TargetSection] = []
        self.regions: dict[str, Region] = {}
        self.conditional_elements: dict[str, object] = {}
        self.ambiguous: list[AmbiguousRegion] = []
        self.icons_without_instruction: list[str] = []
        self.issues: list[str] = []
        self._section: Optional[TargetSection] = None
        self._pending = _Pending()
        self._colour_cache: dict[tuple[str, str], Optional[str]] = {}
        self._used_tags: set[str] = set()

    # ── Walk ───────────────────────────────────────────────────────────

    def run(self) -> None:
        paragraph_index = table_index = 0
        top_level = 0
        blocks = list(iter_body_blocks(self.doc.element.body))
        i = 0
        while i < len(blocks):
            el = blocks[i]
            if el.tag == qn("w:tbl"):
                self._table(el, table_index)
                table_index += 1
                i += 1
                continue

            numbering = self.numbering.advance(el)
            if is_heading(self.doc, el) and block_text(el):
                level = self._heading_level(el)
                if level == 1:
                    top_level += 1
                number = numbering.number if numbering and numbering.number else (str(top_level) if level == 1 else None)
                self._start_section(el, level, number, paragraph_index)
                paragraph_index += 1
                i += 1
                continue

            tag = _cc_tag(el)
            if tag is not None and self._section is not None:
                group = [el]
                while i + 1 < len(blocks) and blocks[i + 1].tag == qn("w:p") and _cc_tag(blocks[i + 1]) == tag:
                    i += 1
                    self.numbering.advance(blocks[i])
                    group.append(blocks[i])
                self._tagged_paragraphs(tag, group, paragraph_index)
                paragraph_index += len(group)
                i += 1
                continue

            self._paragraph(el, paragraph_index)
            paragraph_index += 1
            i += 1
        self._flush_pending()

    def _heading_level(self, p) -> int:
        match = re.match(r"heading\s*(\d)", style_name(self.doc, p), re.I)
        if match:
            return int(match.group(1))
        outline = self.numbering.outline_level(p)
        return outline + 1 if outline is not None else 1

    def _start_section(self, p, level: int, number: Optional[str], paragraph_index: int) -> None:
        self._flush_pending()
        heading = block_text(p)
        optional = heading in self.palette.optional_section_headings
        clean = re.sub(r"\s*\(optional\)\s*", " ", heading, flags=re.I).strip()
        key = _unique(_section_key(clean), {s.key for s in self.sections})
        override = self.config.sections.get(key)
        required = not optional
        if override is not None and override.required is not None:
            required = override.required
        section = TargetSection(
            section_id=f"TGT-{number}" if number else f"TGT-S{len(self.sections) + 1:02d}",
            key=key,
            number=number,
            heading=clean,
            level=level,
            display_order=len(self.sections),
            required=required,
            optional_marker=optional,
        )
        if any(s.section_id == section.section_id for s in self.sections):
            section.section_id = f"{section.section_id}-{len(self.sections) + 1}"
        self.sections.append(section)
        self._section = section

    # ── Paragraphs ─────────────────────────────────────────────────────

    def _paragraph(self, p, paragraph_index: int) -> None:
        text = block_text(p)
        if self._section is None or not text:
            return  # cover sheet and general rules (v1 global rules), or empty
        region_id = f"p:{paragraph_index}"
        bookmark = _slot_bookmark(p)
        colour = self._colour_class(p)

        if colour == "mixed":
            decision = self._ambiguous(region_id, "inline_instruction", text, RegionDecision.FIXED)
            if decision == RegionDecision.INSTRUCTION:
                colour = "blue"
            elif decision == RegionDecision.SLOT:
                self._flush_pending()
                self._add_paragraph_slot([p], paragraph_index, text, bookmark)
                return
            if decision == RegionDecision.CONDITIONAL:
                self._conditional(p, region_id, ConditionalKind.INLINE_CHOICE, text,
                                  SlotAnchor(kind=AnchorKind.PARAGRAPH, ref=region_id))
            return  # kept text; a pending group may still lead past it
        if colour == "blue":
            if bookmark and self._pending.paragraphs:
                self._flush_pending()
            if not self._pending.paragraphs:
                self._pending.first_index = paragraph_index
            self._pending.paragraphs.append(p)
            self._pending.texts.append(text)
            self._pending.bookmark = self._pending.bookmark or bookmark
            return

        placeholder = _PLACEHOLDER.search(text)
        if bookmark or placeholder:
            self._flush_pending()
            self._add_paragraph_slot([p], paragraph_index, text, bookmark, placeholder.group(0) if placeholder else None)

    def _tagged_paragraphs(self, tag: str, group: list, paragraph_index: int) -> None:
        self._flush_pending()
        text = "\n".join(t for t in (block_text(p) for p in group) if t)
        self._add_paragraph_slot(group, paragraph_index, text, tag=tag)

    def _flush_pending(self) -> None:
        pending, self._pending = self._pending, _Pending()
        if not pending.paragraphs or self._section is None:
            return
        self._add_paragraph_slot(pending.paragraphs, pending.first_index, pending.text, pending.bookmark)

    def _take_pending(self) -> Optional[str]:
        """The pending instruction text, consumed by the table it leads into."""
        pending, self._pending = self._pending, _Pending()
        return pending.text

    def _add_paragraph_slot(self, paragraphs, paragraph_index, text, bookmark=None, placeholder=None, tag=None) -> None:
        if tag:
            anchor = SlotAnchor(kind=AnchorKind.CONTENT_CONTROL, ref=tag)
        elif bookmark:
            anchor = SlotAnchor(kind=AnchorKind.BOOKMARK, ref=bookmark)
        elif placeholder:
            anchor = SlotAnchor(kind=AnchorKind.PLACEHOLDER, ref=placeholder)
        else:
            anchor = SlotAnchor(kind=AnchorKind.PARAGRAPH, ref=f"p:{paragraph_index}")
        key = (self._key_from_tag(tag) or self._key_from_bookmark(bookmark)
               or _match_key(_PARAGRAPH_KEYS, text) or "content")
        content_type = self._content_type(text)
        slot = self._new_slot(
            key=key,
            instruction=text,
            anchor=anchor,
            behavior=InstructionBehavior.REPLACE,
            profile=None if content_type == ContentType.ORDERED_PROCEDURE else FormattingProfile.PARAGRAPHS,
            content_type=content_type,
        )
        self.regions[slot.slot_id] = Region("tagged" if tag else "paragraphs", f"p:{paragraph_index}", list(paragraphs))

    # ── Tables ─────────────────────────────────────────────────────────

    def _table(self, tbl, table_index: int) -> None:
        region_id = f"t:{table_index}"
        if self._section is None or table_index == self.palette.legend_table_index:
            return
        rows = _rows(tbl)
        if table_index in self.prototypes:
            self._callout_slot(tbl, rows, table_index)
            return
        icon_rows = [(ri, cells) for ri, cells in enumerate(_cells(r) for r in rows)
                     if len(cells) >= 2 and is_image_only(cells[0])]
        if icon_rows and len(icon_rows) == len(rows):
            self._icon_table(tbl, icon_rows, table_index)
            return

        texts = [t for t in (block_text(tc) for r in rows for tc in _cells(r)) if t]
        if texts and all(self._colour_class(tc) in ("blue", None) for r in rows for tc in _cells(r)):
            decision = self._ambiguous(region_id, "instruction_table", " | ".join(texts), RegionDecision.INSTRUCTION)
            if decision == RegionDecision.INSTRUCTION:
                self._pending.texts.append(" | ".join(texts))
                return
            if decision == RegionDecision.CONDITIONAL:
                self._conditional(tbl, region_id, ConditionalKind.BLOCK, " | ".join(texts),
                                  SlotAnchor(kind=AnchorKind.TABLE_CELL, ref=region_id,
                                             table_index=table_index, row_index=0, column_index=0))
                return
            if decision != RegionDecision.SLOT:
                return
        if len(rows) >= 2:
            self._data_table(tbl, rows, table_index)

    def _icon_table(self, tbl, icon_rows, table_index: int) -> None:
        lead_in = self._take_pending()
        if lead_in:
            self._section.purpose_instruction = "\n".join(filter(None, [self._section.purpose_instruction, lead_in]))
        for ri, cells in icon_rows:
            text_cell = cells[1]
            text = block_text(text_cell)
            icon = icon_key_for(self.doc, cells[0])
            if not text:
                self.icons_without_instruction.append(f"t:{table_index}:r{ri}")
                continue
            tag = _cell_tag(text_cell)
            key = self._key_from_tag(tag) or _match_key(_ICON_ROW_KEYS, text) or f"row{ri + 1}"
            slot = self._new_slot(
                key=key,
                instruction=text,
                anchor=_cell_anchor(tag, table_index, ri, 1),
                behavior=InstructionBehavior.REPLACE,
                profile=FormattingProfile.PARAGRAPHS,
                icon=SlotIcon(icon_key=icon) if icon else None,
            )
            self.regions[slot.slot_id] = Region("cell", f"t:{table_index}:r{ri}:c1", [text_cell])

    def _callout_slot(self, tbl, rows, table_index: int) -> None:
        prototype = self.prototypes[table_index]
        lead_in = self._pending.pop_last_paragraph() if self._pending.paragraphs else None
        self._flush_pending()
        cells = _cells(rows[0])
        text_cell = cells[-1]
        tag = _cell_tag(text_cell)
        style = self.palette.style_for(prototype.kind)
        instruction = "\n".join(filter(None, [block_text(lead_in) if lead_in is not None else None, prototype.instruction]))
        slot = self._new_slot(
            key=self._key_from_tag(tag) or f"callout_{prototype.kind.value}",
            instruction=instruction,
            anchor=_cell_anchor(tag, table_index, 0, len(cells) - 1),
            behavior=InstructionBehavior.REPLACE,
            profile=FormattingProfile.CALLOUT,
            callout_kind=prototype.kind,
            content_type=ContentType.WARNING if prototype.kind == CalloutKind.ATTENTION else ContentType.SUPPORTING_INFORMATION,
            icon=SlotIcon(icon_key=style.icon_key, type=prototype.kind.value) if style and style.icon_key else None,
            required=False,
        )
        self.regions[slot.slot_id] = Region("cell", f"t:{table_index}:r0:c{len(cells) - 1}", [text_cell])

    def _data_table(self, tbl, rows, table_index: int) -> None:
        first = _first_data_row(rows, self)
        if first is None:
            return
        header = " | ".join(block_text(tc) for tc in _cells(rows[0]))
        lead_in = self._take_pending()
        tag = _rows_tag(rows[first])
        key = self._key_from_tag(tag) or _match_key(_TABLE_KEYS, block_text(_cells(rows[0])[0])) or "table"
        instruction = lead_in or f"Fill the table: {header}"
        anchor = (SlotAnchor(kind=AnchorKind.CONTENT_CONTROL, ref=tag, table_index=table_index, row_index=first, column_index=0)
                  if tag else
                  SlotAnchor(kind=AnchorKind.TABLE_CELL, ref=f"t:{table_index}:r{first}",
                             table_index=table_index, row_index=first, column_index=0))
        slot = self._new_slot(
            key=key,
            instruction=instruction,
            anchor=anchor,
            behavior=InstructionBehavior.HIDE_AFTER_POPULATION if lead_in else InstructionBehavior.REPLACE,
            profile=FormattingProfile.TABLE,
        )
        self.regions[slot.slot_id] = Region("rows", f"t:{table_index}:r{first}", rows[first:])

    # ── Slots ──────────────────────────────────────────────────────────

    def _new_slot(self, *, key, instruction, anchor, behavior, profile, callout_kind=None,
                  content_type=None, icon=None, required=None) -> TargetSlot:
        section = self._section
        key = _unique(key, {s.key for s in section.slots})
        if required is None:
            # Only the opening line counts: "If applicable, list ..." is about the slot,
            # an "optional" deep in a long guidance block is usually about something else.
            opening = instruction.split("\n", 1)[0]
            required = section.required and (
                bool(_REQUIRED_EMPTY_VALUE.search(instruction)) or not _OPTIONAL_WORDING.search(opening)
            )
        slot = TargetSlot(
            slot_id=f"{section.section_id}-{key.upper()}",
            section_id=section.section_id,
            key=key,
            instruction=instruction,
            content_type=content_type or self._content_type(instruction),
            required=required,
            display_order=len(section.slots),
            anchor=anchor,
            instruction_behavior=behavior,
            icon=icon,
            formatting_profile=profile,
            callout_kind=callout_kind,
        )
        section.slots.append(slot)
        return slot

    def _content_type(self, instruction: str) -> ContentType:
        for pattern, content_type in _SECTION_TYPES:
            if pattern.search(self._section.heading):
                return content_type
        for pattern, content_type in _INSTRUCTION_TYPES:
            if pattern.search(instruction):
                return content_type
        return ContentType.SUPPORTING_INFORMATION

    def _key_from_tag(self, tag: Optional[str]) -> Optional[str]:
        if not tag:
            return None
        prefix = f"{TAG_PREFIX}{self._section.key}_"
        if tag.startswith(prefix):
            return _slug(tag[len(prefix):]) or None
        self.issues.append(f"content control '{tag}' is in section {self._section.key}, not the one its tag names")
        return _slug(tag[len(TAG_PREFIX):]) or None

    def _key_from_bookmark(self, name: Optional[str]) -> Optional[str]:
        if not name:
            return None
        prefix = f"{BOOKMARK_PREFIX}{self._section.key}_"
        return _slug(name[len(prefix):] if name.startswith(prefix) else name[len(BOOKMARK_PREFIX):]) or None

    def _conditional(self, el, region_id: str, kind: ConditionalKind, text: str, anchor: SlotAnchor) -> None:
        tag = _cond_tag(el)
        if tag:
            anchor = SlotAnchor(kind=AnchorKind.CONTENT_CONTROL, ref=tag, table_index=anchor.table_index,
                                row_index=anchor.row_index, column_index=anchor.column_index)
        self._section.conditional_regions.append(
            ConditionalRegion(region_id=region_id, kind=kind, text=text[:500], anchor=anchor)
        )
        self.conditional_elements[region_id] = el

    def _ambiguous(self, region_id, kind, text, suggestion: RegionDecision) -> Optional[RegionDecision]:
        decision = self.config.regions.get(region_id)
        self.ambiguous.append(AmbiguousRegion(
            region_id=region_id,
            kind=kind,
            section_key=self._section.key if self._section else None,
            text=text[:200],
            suggestion=suggestion,
            decision=decision,
        ))
        return decision

    # ── Colour ─────────────────────────────────────────────────────────

    def _colour_class(self, el) -> Optional[str]:
        """"blue", "plain" or "mixed" by the text runs of *el*; None if it has no text."""
        blue = plain = False
        for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))):
            p_style = _style_id(p, "pPr", "pStyle")
            for r in p.iter(qn("w:r")):
                text = "".join(t.text or "" for t in r.iter(qn("w:t")))
                if not text.strip():
                    continue
                if self._run_is_blue(r, p_style):
                    blue = True
                else:
                    plain = True
        if blue and plain:
            return "mixed"
        return "blue" if blue else ("plain" if plain else None)

    def _run_is_blue(self, r, p_style: Optional[str]) -> bool:
        colour = r.find(f"{qn('w:rPr')}/{qn('w:color')}")
        value = colour.get(qn("w:val")) if colour is not None else None
        if value is None:
            value = self._style_colour(_style_id(r, "rPr", "rStyle"), WD_STYLE_TYPE.CHARACTER)
        if value is None:
            value = self._style_colour(p_style, WD_STYLE_TYPE.PARAGRAPH)
        return bool(value) and value.lower() != "auto" and TemplateDocxParser.is_blue_hex(value)

    def _style_colour(self, style_id: Optional[str], style_type) -> Optional[str]:
        if not style_id:
            return None
        cache_key = (style_id, str(style_type))
        if cache_key not in self._colour_cache:
            colour = None
            try:
                style = self.doc.styles.get_by_id(style_id, style_type)
                while style is not None and colour is None:
                    rgb = style.font.color.rgb if style.font.color is not None and style.font.color.type else None
                    colour = str(rgb) if rgb is not None else None
                    style = style.base_style
            except Exception:
                colour = None
            self._colour_cache[cache_key] = colour
        return self._colour_cache[cache_key]


# ── XML helpers ────────────────────────────────────────────────────────

def _rows(tbl) -> list:
    """Table rows, including rows inside row-level content controls."""
    rows = []
    for child in tbl:
        if child.tag == qn("w:tr"):
            rows.append(child)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            rows.extend(content.findall(qn("w:tr")) if content is not None else [])
    return rows


def _cells(tr) -> list:
    from app.services.parser.docx_parser import DocxParser

    return DocxParser._row_cells(tr)


def _first_data_row(rows, detector: _Detector) -> Optional[int]:
    """First row after the header: header rows are row 0 and black rows of bracket placeholders."""
    for ri in range(1, len(rows)):
        text = block_text(rows[ri])
        if detector._colour_class(rows[ri]) != "blue" and re.search(r"\[[^\]]+\]", text) and not _PLACEHOLDER.search(text):
            continue
        return ri
    return None


def _sdt_tag(sdt) -> Optional[str]:
    tag = sdt.find(f"{qn('w:sdtPr')}/{qn('w:tag')}")
    return tag.get(qn("w:val")) if tag is not None else None


def _cc_tag(el) -> Optional[str]:
    """Tag of the nearest enclosing ``CC_`` content control of a body block."""
    parent = el.getparent()
    while parent is not None and parent.tag != qn("w:body"):
        if parent.tag == qn("w:sdt"):
            tag = _sdt_tag(parent)
            if tag and tag.startswith(TAG_PREFIX):
                return tag
        parent = parent.getparent()
    return None


def _cond_tag(el) -> Optional[str]:
    """Tag of the ``COND_`` content control directly around a paragraph or table."""
    parent = el.getparent()
    if parent is not None and parent.tag == qn("w:sdtContent"):
        tag = _sdt_tag(parent.getparent())
        if tag and tag.startswith(COND_PREFIX):
            return tag
    return None


def _cell_tag(tc) -> Optional[str]:
    for sdt in tc.findall(qn("w:sdt")):
        tag = _sdt_tag(sdt)
        if tag and tag.startswith(TAG_PREFIX):
            return tag
    return None


def _rows_tag(tr) -> Optional[str]:
    parent = tr.getparent()
    if parent is not None and parent.tag == qn("w:sdtContent"):
        tag = _sdt_tag(parent.getparent())
        if tag and tag.startswith(TAG_PREFIX):
            return tag
    return None


def _cell_anchor(tag: Optional[str], table_index: int, row: int, column: int) -> SlotAnchor:
    if tag:
        return SlotAnchor(kind=AnchorKind.CONTENT_CONTROL, ref=tag, table_index=table_index, row_index=row, column_index=column)
    return SlotAnchor(kind=AnchorKind.TABLE_CELL, ref=f"t:{table_index}:r{row}:c{column}",
                      table_index=table_index, row_index=row, column_index=column)


def _slot_bookmark(p) -> Optional[str]:
    for bm in p.iter(qn("w:bookmarkStart")):
        name = bm.get(qn("w:name")) or ""
        if name.startswith(BOOKMARK_PREFIX):
            return name
    return None


def _style_id(el, pr: str, tag: str) -> Optional[str]:
    node = el.find(f"{qn('w:' + pr)}/{qn('w:' + tag)}")
    return node.get(qn("w:val")) if node is not None else None


# ── Names ──────────────────────────────────────────────────────────────

def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")


def _section_key(heading: str) -> str:
    """"IMPLEMENTATION AND/OR PRE-REQUISITES" → "IMPLEMENTATION"."""
    head = re.split(r"\s*(?:&|\(|/|\band\b|\bor\b)", heading, maxsplit=1, flags=re.I)[0]
    slug = _slug(head).upper() or "SECTION"
    return slug if slug[0].isalpha() else f"S_{slug}"


def _match_key(rules: list[tuple[str, re.Pattern]], text: str) -> Optional[str]:
    return next((key for key, pattern in rules if pattern.search(text or "")), None)


def _unique(key: str, taken: set) -> str:
    if key not in taken:
        return key
    n = 2
    while f"{key}_{n}" in taken:
        n += 1
    return f"{key}_{n}"
